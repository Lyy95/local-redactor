from __future__ import annotations

import hashlib
import io
import posixpath
import re
import uuid
import zipfile
from collections import defaultdict
from collections.abc import Sequence
from contextlib import suppress
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, cast

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree  # type: ignore[import-untyped]

from ..models import (
    DocumentKind,
    DocumentModel,
    Finding,
    ImageObject,
    ObjectDisposition,
    PackageItem,
    SourceLocation,
    TextBlock,
)
from .common import (
    cleanup_temp,
    ensure_export_ready,
    finalize_export,
    forbidden_originals,
    image_replacement_legend,
    item_disposition,
    prepare_image_for_export,
    sanitize_image,
    secure_target_path,
    temporary_export_path,
    transformed_blocks,
    validate_source_image,
    verify_rebuilt_ooxml,
)
from .security import (
    ExportSafetyError,
    OoxmlPreflight,
    OoxmlSecurityError,
    preflight_ooxml,
    safe_xml_root,
    sha256_file,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W_NS, "r": R_NS, "wp": WP_NS, "a": A_NS, "v": V_NS}

_REMOVED_DOCX_PART_PREFIXES = (
    "customXml/",
    "word/activeX/",
    "word/embeddings/",
    "word/comments",
    "word/footnotes",
    "word/endnotes",
    "word/glossary/",
)
_REMOVED_DOCX_PART_NAMES = {
    "docProps/thumbnail.jpeg",
    "word/vbaProject.bin",
}


def _q(namespace: str, local_name: str) -> str:
    return f"{{{namespace}}}{local_name}"


def _local_name(element: etree._Element) -> str:
    return cast(str, etree.QName(element).localname)


def _ancestor_has(element: etree._Element, local_names: set[str]) -> bool:
    return any(_local_name(ancestor) in local_names for ancestor in element.iterancestors())


def _run_is_hidden(element: etree._Element) -> bool:
    run = element if _local_name(element) == "r" else next(
        (ancestor for ancestor in element.iterancestors() if _local_name(ancestor) == "r"),
        None,
    )
    if run is None:
        return False
    properties = run.find(_q(W_NS, "rPr"))
    if properties is None:
        return False
    return (
        properties.find(_q(W_NS, "vanish")) is not None
        or properties.find(_q(W_NS, "webHidden")) is not None
    )


def _paragraph_text(
    paragraph: etree._Element,
    *,
    hidden: bool = False,
    deleted: bool = False,
    include_textboxes: bool = False,
) -> str:
    fragments: list[str] = []
    for node in paragraph.iter():
        local_name = _local_name(node)
        if local_name not in {"t", "delText", "tab", "br", "cr"}:
            continue
        in_textbox = _ancestor_has(node, {"txbxContent"})
        if in_textbox and not include_textboxes:
            continue
        in_deleted = _ancestor_has(node, {"del", "moveFrom"}) or local_name == "delText"
        if in_deleted != deleted:
            continue
        is_hidden = _run_is_hidden(node)
        if is_hidden != hidden:
            continue
        if local_name == "tab":
            fragments.append("\t")
        elif local_name in {"br", "cr"}:
            fragments.append("\n")
        elif node.text:
            fragments.append(node.text)
    return "".join(fragments)


def _paragraph_style(paragraph: etree._Element) -> dict[str, Any]:
    properties = paragraph.find(_q(W_NS, "pPr"))
    style: dict[str, Any] = {}
    if properties is not None:
        style_node = properties.find(_q(W_NS, "pStyle"))
        if style_node is not None:
            style["style_id"] = style_node.get(_q(W_NS, "val"), "")
        align_node = properties.find(_q(W_NS, "jc"))
        if align_node is not None:
            style["alignment"] = align_node.get(_q(W_NS, "val"), "")
        numbering = properties.find(_q(W_NS, "numPr"))
        if numbering is not None:
            num_id = numbering.find(_q(W_NS, "numId"))
            level = numbering.find(_q(W_NS, "ilvl"))
            style["numbered"] = True
            style["num_id"] = num_id.get(_q(W_NS, "val"), "") if num_id is not None else ""
            style["level"] = level.get(_q(W_NS, "val"), "0") if level is not None else "0"
        # Capture paragraph spacing
        spacing = properties.find(_q(W_NS, "spacing"))
        if spacing is not None:
            for attr in ("before", "after", "line", "lineRule"):
                val = spacing.get(_q(W_NS, attr))
                if val is not None:
                    style[f"spacing_{attr}"] = val
        # Capture first-line indent
        ind = properties.find(_q(W_NS, "ind"))
        if ind is not None:
            first_line = ind.get(_q(W_NS, "firstLine"))
            if first_line is not None:
                style["first_line_indent"] = first_line
    visible_runs = [
        run
        for run in paragraph.iter(_q(W_NS, "r"))
        if not _ancestor_has(run, {"del", "moveFrom", "txbxContent"})
    ]
    if visible_runs:
        bold_values: list[bool] = []
        italic_values: list[bool] = []
        run_formats: list[dict[str, Any]] = []
        for run in visible_runs:
            rpr = run.find(_q(W_NS, "rPr"))
            is_bold = rpr is not None and rpr.find(_q(W_NS, "b")) is not None
            is_italic = rpr is not None and rpr.find(_q(W_NS, "i")) is not None
            bold_values.append(is_bold)
            italic_values.append(is_italic)
            rf: dict[str, Any] = {
                "text": _paragraph_text(run),
                "bold": is_bold,
                "italic": is_italic,
                "underline": False,
            }
            if rpr is not None:
                rf_fonts = rpr.find(_q(W_NS, "rFonts"))
                if rf_fonts is not None:
                    ascii_font = rf_fonts.get(_q(W_NS, "ascii"), "")
                    east_asia = rf_fonts.get(_q(W_NS, "eastAsia"), "")
                    h_ansi = rf_fonts.get(_q(W_NS, "hAnsi"), "")
                    rf["font_name"] = ascii_font or east_asia or h_ansi or ""
                sz = rpr.find(_q(W_NS, "sz"))
                if sz is not None:
                    rf["font_size"] = sz.get(_q(W_NS, "val"), "")
                sz_cs = rpr.find(_q(W_NS, "szCs"))
                if sz_cs is not None and "font_size" not in rf:
                    rf["font_size"] = sz_cs.get(_q(W_NS, "val"), "")
                color = rpr.find(_q(W_NS, "color"))
                if color is not None:
                    rf["font_color"] = color.get(_q(W_NS, "val"), "")
                u = rpr.find(_q(W_NS, "u"))
                if u is not None:
                    rf["underline"] = True
            if rf["text"]:
                run_formats.append(rf)
        style["bold"] = all(bold_values)
        style["italic"] = all(italic_values)
        # Use the first run's font formatting as paragraph default
        if run_formats:
            first_rf = run_formats[0]
            style["font_name"] = first_rf.get("font_name", "")
            style["font_size"] = first_rf.get("font_size", "")
            style["font_color"] = first_rf.get("font_color", "")
            style["underline"] = first_rf.get("underline", False)
        style["run_formats"] = run_formats
    return style


def _new_block(
    model: DocumentModel,
    *,
    text: str,
    part: str,
    display: str,
    block_kind: str,
    style: dict[str, Any] | None = None,
    hidden: bool = False,
    sheet: str | None = None,
    cell: str | None = None,
) -> TextBlock:
    block_id = uuid.uuid4().hex
    block = TextBlock(
        id=block_id,
        text=text,
        location=SourceLocation(
            part=part,
            display=display,
            block_id=block_id,
            sheet=sheet,
            cell=cell,
        ),
        block_kind=block_kind,
        style=style or {},
        hidden=hidden,
    )
    model.blocks.append(block)
    return block


def _add_virtual_item(
    model: DocumentModel,
    *,
    part: str,
    kind: str,
    hidden: bool = True,
) -> PackageItem:
    existing = next((item for item in model.inventory.parts if item.part == part), None)
    if existing is not None:
        existing.requires_decision = True
        existing.hidden = existing.hidden or hidden
        return existing
    item = PackageItem(
        part=part,
        kind=kind,
        hidden=hidden,
        requires_decision=True,
    )
    model.inventory.parts.append(item)
    return item


def _load_roots(
    archive: zipfile.ZipFile,
    preflight: OoxmlPreflight,
) -> dict[str, etree._Element]:
    roots: dict[str, etree._Element] = {}
    for name in preflight.part_names:
        if (name.startswith("word/") or name.startswith("docProps/")) and name.endswith(".xml"):
            roots[name] = safe_xml_root(archive.read(name), part_name=name)
    return roots


def _scan_properties(
    model: DocumentModel,
    roots: dict[str, etree._Element],
) -> None:
    for part, root in sorted(roots.items()):
        if not part.startswith("docProps/"):
            continue
        for element in root.iter():
            if len(element) or not element.text or not element.text.strip():
                continue
            key = _local_name(element)
            value = element.text.strip()
            unique_key = key
            counter = 2
            while unique_key in model.document_properties:
                unique_key = f"{key}_{counter}"
                counter += 1
            model.document_properties[unique_key] = value
            _new_block(
                model,
                text=value,
                part=part,
                display=f"文件属性：{key}",
                block_kind="metadata",
                hidden=True,
            )


def _scan_paragraph_side_channels(
    model: DocumentModel,
    paragraph: etree._Element,
    *,
    part: str,
    display: str,
    serial: int,
) -> None:
    hidden_text = _paragraph_text(paragraph, hidden=True, deleted=False)
    if hidden_text:
        virtual_part = f"{part}#hidden:{serial}"
        _add_virtual_item(model, part=virtual_part, kind="hidden_text")
        block = _new_block(
            model,
            text=hidden_text,
            part=virtual_part,
            display=f"{display}（隐藏文字）",
            block_kind="hidden_text",
            hidden=True,
        )
        model.structure.setdefault("side_channels", []).append(
            {
                "part": virtual_part,
                "kind": "hidden_text",
                "blocks": [block.id],
            }
        )
    deleted_text = _paragraph_text(paragraph, hidden=False, deleted=True)
    if deleted_text:
        virtual_part = f"{part}#revision:{serial}"
        _add_virtual_item(model, part=virtual_part, kind="revision")
        block = _new_block(
            model,
            text=deleted_text,
            part=virtual_part,
            display=f"{display}（修订删除内容）",
            block_kind="revision_deleted",
            hidden=True,
        )
        model.structure.setdefault("side_channels", []).append(
            {
                "part": virtual_part,
                "kind": "revision",
                "blocks": [block.id],
            }
        )
    field_codes = [
        node.text
        for node in paragraph.iter(_q(W_NS, "instrText"))
        if node.text and node.text.strip()
    ]
    if field_codes:
        virtual_part = f"{part}#field:{serial}"
        _add_virtual_item(model, part=virtual_part, kind="field_code")
        block = _new_block(
            model,
            text=" ".join(field_codes),
            part=virtual_part,
            display=f"{display}（字段代码）",
            block_kind="field_code",
            hidden=True,
        )
        model.structure.setdefault("side_channels", []).append(
            {
                "part": virtual_part,
                "kind": "field_code",
                "blocks": [block.id],
            }
        )


def _scan_body(
    model: DocumentModel,
    root: etree._Element,
) -> tuple[list[dict[str, Any]], dict[int, str]]:
    body = root.find(_q(W_NS, "body"))
    if body is None:
        raise OoxmlSecurityError("missing-word-body", "Word 主文档缺少正文容器")
    structure: list[dict[str, Any]] = []
    paragraph_ids: dict[int, str] = {}
    paragraph_serial = 0
    table_serial = 0
    for child in body:
        local_name = _local_name(child)
        if local_name == "p":
            paragraph_serial += 1
            block = _new_block(
                model,
                text=_paragraph_text(child),
                part="word/document.xml",
                display=f"正文第 {paragraph_serial} 段",
                block_kind="paragraph",
                style=_paragraph_style(child),
            )
            paragraph_ids[id(child)] = block.id
            structure.append(
                {
                    "kind": "paragraph",
                    "block_id": block.id,
                    "style": dict(block.style),
                }
            )
            _scan_paragraph_side_channels(
                model,
                child,
                part="word/document.xml",
                display=block.location.display,
                serial=paragraph_serial,
            )
        elif local_name == "tbl":
            table_serial += 1
            rows_data: list[list[dict[str, Any]]] = []
            table_paragraph = 0
            for row_index, row in enumerate(child.findall(_q(W_NS, "tr")), start=1):
                cells_data: list[dict[str, Any]] = []
                grid_column = 0
                for cell_index, cell in enumerate(
                    row.findall(_q(W_NS, "tc")),
                    start=1,
                ):
                    properties = cell.find(_q(W_NS, "tcPr"))
                    grid_span = 1
                    vertical_merge = ""
                    if properties is not None:
                        grid_span_node = properties.find(_q(W_NS, "gridSpan"))
                        if grid_span_node is not None:
                            try:
                                grid_span = max(
                                    1,
                                    int(grid_span_node.get(_q(W_NS, "val"), "1")),
                                )
                            except ValueError:
                                grid_span = 1
                        vertical_merge_node = properties.find(_q(W_NS, "vMerge"))
                        if vertical_merge_node is not None:
                            vertical_merge = vertical_merge_node.get(
                                _q(W_NS, "val"),
                                "continue",
                            )
                    block_ids: list[str] = []
                    for paragraph in cell.findall(_q(W_NS, "p")):
                        table_paragraph += 1
                        block = _new_block(
                            model,
                            text=_paragraph_text(paragraph),
                            part="word/document.xml",
                            display=(f"表格 {table_serial} 第 {row_index} 行第 {cell_index} 列"),
                            block_kind="table_cell",
                            style=_paragraph_style(paragraph),
                        )
                        paragraph_ids[id(paragraph)] = block.id
                        block_ids.append(block.id)
                        _scan_paragraph_side_channels(
                            model,
                            paragraph,
                            part="word/document.xml",
                            display=block.location.display,
                            serial=10_000 * table_serial + table_paragraph,
                        )
                    cells_data.append(
                        {
                            "blocks": block_ids,
                            "column": grid_column,
                            "grid_span": grid_span,
                            "vertical_merge": vertical_merge,
                        }
                    )
                    grid_column += grid_span
                rows_data.append(cells_data)
            style_node = child.find(f"{_q(W_NS, 'tblPr')}/{_q(W_NS, 'tblStyle')}")
            structure.append(
                {
                    "kind": "table",
                    "style_id": (
                        style_node.get(_q(W_NS, "val"), "") if style_node is not None else ""
                    ),
                    "rows": rows_data,
                }
            )
    return structure, paragraph_ids


def _scan_supplemental_parts(
    model: DocumentModel,
    roots: dict[str, etree._Element],
    paragraph_ids: dict[tuple[str, int], str],
) -> list[dict[str, Any]]:
    supplemental: list[dict[str, Any]] = []
    side_groups: list[dict[str, Any]] = []
    supported_prefixes = (
        "word/header",
        "word/footer",
        "word/footnotes",
        "word/endnotes",
        "word/comments",
    )
    for part, root in sorted(roots.items()):
        if not part.startswith(supported_prefixes):
            continue
        group: dict[str, Any] = {"part": part, "kind": "supplemental", "blocks": []}
        paragraph_number = 0
        for paragraph in root.iter(_q(W_NS, "p")):
            if _ancestor_has(paragraph, {"txbxContent"}):
                continue
            paragraph_number += 1
            text = _paragraph_text(paragraph)
            hidden_text = _paragraph_text(paragraph, hidden=True)
            if text:
                block = _new_block(
                    model,
                    text=text,
                    part=part,
                    display=f"{part} 第 {paragraph_number} 段",
                    block_kind=_supplemental_kind(part),
                    style=_paragraph_style(paragraph),
                    hidden="comments" in part or "notes" in part,
                )
                group["blocks"].append(block.id)
                paragraph_ids[(part, id(paragraph))] = block.id
            if hidden_text:
                hidden_part = f"{part}#hidden:{paragraph_number}"
                _add_virtual_item(model, part=hidden_part, kind="hidden_text")
                hidden_block = _new_block(
                    model,
                    text=hidden_text,
                    part=hidden_part,
                    display=f"{part} 第 {paragraph_number} 段（隐藏文字）",
                    block_kind="hidden_text",
                    hidden=True,
                )
                side_groups.append(
                    {
                        "part": hidden_part,
                        "kind": "hidden_text",
                        "blocks": [hidden_block.id],
                    }
                )
        if group["blocks"]:
            supplemental.append(group)

    textbox_number = 0
    for part, root in sorted(roots.items()):
        if not part.startswith("word/"):
            continue
        for paragraph in root.xpath(".//w:txbxContent//w:p", namespaces=NS):
            textbox_number += 1
            virtual_part = f"{part}#textbox:{textbox_number}"
            _add_virtual_item(model, part=virtual_part, kind="textbox")
            block = _new_block(
                model,
                text=_paragraph_text(paragraph, include_textboxes=True),
                part=virtual_part,
                display=f"文本框 {textbox_number}",
                block_kind="textbox",
                style=_paragraph_style(paragraph),
                hidden=True,
            )
            paragraph_ids[(part, id(paragraph))] = block.id
            supplemental.append(
                {
                    "part": virtual_part,
                    "kind": "textbox",
                    "blocks": [block.id],
                }
            )

    watermark_number = 0
    for part, root in sorted(roots.items()):
        if not part.startswith(("word/header", "word/footer", "word/document")):
            continue
        for textpath in root.iter(_q(V_NS, "textpath")):
            value = textpath.get("string", "")
            if not value:
                continue
            watermark_number += 1
            virtual_part = f"{part}#watermark:{watermark_number}"
            _add_virtual_item(model, part=virtual_part, kind="watermark")
            block = _new_block(
                model,
                text=value,
                part=virtual_part,
                display=f"水印 {watermark_number}",
                block_kind="watermark",
                hidden=True,
            )
            supplemental.append(
                {
                    "part": virtual_part,
                    "kind": "watermark",
                    "blocks": [block.id],
                }
            )
    supplemental.extend(side_groups)
    return supplemental


def _supplemental_kind(part: str) -> str:
    lowered = part.casefold()
    if "header" in lowered:
        return "header"
    if "footer" in lowered:
        return "footer"
    if "footnotes" in lowered:
        return "footnote"
    if "endnotes" in lowered:
        return "endnote"
    if "comments" in lowered:
        return "comment"
    return "supplemental"


def _find_paragraph_block_id(
    node: etree._Element,
    *,
    part: str,
    body_ids: dict[int, str],
    supplemental_ids: dict[tuple[str, int], str],
) -> str | None:
    paragraph = next(
        (ancestor for ancestor in node.iterancestors() if _local_name(ancestor) == "p"),
        None,
    )
    if paragraph is None:
        return None
    return body_ids.get(id(paragraph)) or supplemental_ids.get((part, id(paragraph)))


def _scan_images_and_alt_text(
    model: DocumentModel,
    archive: zipfile.ZipFile,
    preflight: OoxmlPreflight,
    roots: dict[str, etree._Element],
    body_ids: dict[int, str],
    supplemental_ids: dict[tuple[str, int], str],
) -> dict[str, str]:
    relation_map = {
        (record.source_part, record.relationship_id): record
        for record in preflight.relationships
        if not record.external
        and record.resolved_target is not None
        and record.relationship_type.casefold().endswith("/image")
    }
    represented_parts: set[str] = set()
    alt_blocks: dict[str, str] = {}
    for part, root in sorted(roots.items()):
        if not part.startswith("word/"):
            continue
        for blip in root.iter(_q(A_NS, "blip")):
            relationship_id = blip.get(_q(R_NS, "embed"), "")
            relation = relation_map.get((part, relationship_id))
            if relation is None or relation.resolved_target is None:
                continue
            media_part = relation.resolved_target
            represented_parts.add(media_part)
            drawing = next(
                (
                    ancestor
                    for ancestor in blip.iterancestors()
                    if _local_name(ancestor) in {"drawing", "pict"}
                ),
                None,
            )
            doc_properties = (
                drawing.find(f".//{_q(WP_NS, 'docPr')}") if drawing is not None else None
            )
            alt_text = ""
            if doc_properties is not None:
                alt_text = " ".join(
                    value
                    for value in (
                        doc_properties.get("name", ""),
                        doc_properties.get("title", ""),
                        doc_properties.get("descr", ""),
                    )
                    if value
                )
            block_id = _find_paragraph_block_id(
                blip,
                part=part,
                body_ids=body_ids,
                supplemental_ids=supplemental_ids,
            )
            image_id = uuid.uuid4().hex
            location = SourceLocation(
                part=part,
                display=f"图片：{media_part}",
                block_id=block_id,
                image_id=image_id,
            )
            content = archive.read(media_part)
            width, height = validate_source_image(
                content,
                source_part=media_part,
            )
            image = ImageObject(
                id=image_id,
                source_part=media_part,
                location=location,
                content=content,
                media_type=preflight.content_types.get(media_part, "application/octet-stream"),
                width=width,
                height=height,
                alt_text=alt_text,
            )
            model.images.append(image)
            if alt_text:
                alt_block = _new_block(
                    model,
                    text=alt_text,
                    part=f"{part}#image-alt:{image_id}",
                    display=f"图片替代文字：{media_part}",
                    block_kind="alt_text",
                    hidden=True,
                )
                alt_block.location.image_id = image_id
                alt_blocks[image_id] = alt_block.id

    for media_part in sorted(
        name
        for name in preflight.part_names
        if preflight.content_types.get(name, "").startswith("image/")
        and "/media/" in name
        and name not in represented_parts
    ):
        image_id = uuid.uuid4().hex
        content = archive.read(media_part)
        width, height = validate_source_image(
            content,
            source_part=media_part,
        )
        model.images.append(
            ImageObject(
                id=image_id,
                source_part=media_part,
                location=SourceLocation(
                    part=media_part,
                    display=f"未引用图片：{media_part}",
                    image_id=image_id,
                ),
                content=content,
                media_type=preflight.content_types.get(media_part, "application/octet-stream"),
                width=width,
                height=height,
            )
        )
        model.warnings.append("发现未被正文关系引用的图片，仍需逐图复核")
    return alt_blocks


def _scan_external_relationship_text(
    model: DocumentModel,
    preflight: OoxmlPreflight,
) -> None:
    for relationship in preflight.relationships:
        if not relationship.external:
            continue
        _new_block(
            model,
            text=relationship.target,
            part=(f"{relationship.relationship_part}#{relationship.relationship_id}"),
            display="外部关系真实目标",
            block_kind="external_relationship",
            hidden=True,
        )


class DocxAdapter:
    kind = DocumentKind.DOCX

    def scan(self, path: Path) -> DocumentModel:
        source_path = Path(path)
        before_hash = sha256_file(source_path)
        preflight = preflight_ooxml(source_path, self.kind)
        model = DocumentModel(
            kind=self.kind,
            source_path=source_path,
            source_sha256=preflight.source_sha256,
            display_name=source_path.name,
            inventory=preflight.inventory,
            structure={"format": "docx"},
        )
        try:
            with zipfile.ZipFile(source_path, mode="r") as archive:
                roots = _load_roots(archive, preflight)
                main_root = roots.get("word/document.xml")
                if main_root is None:
                    raise OoxmlSecurityError(
                        "missing-word-document",
                        "Word 主文档无法读取",
                    )
                _scan_properties(model, roots)
                body_structure, body_paragraph_ids = _scan_body(model, main_root)
                supplemental_ids: dict[tuple[str, int], str] = {}
                supplemental = _scan_supplemental_parts(
                    model,
                    roots,
                    supplemental_ids,
                )
                supplemental.extend(model.structure.pop("side_channels", []))
                alt_blocks = _scan_images_and_alt_text(
                    model,
                    archive,
                    preflight,
                    roots,
                    body_paragraph_ids,
                    supplemental_ids,
                )
                _scan_external_relationship_text(model, preflight)
                model.structure.update(
                    {
                        "body": body_structure,
                        "supplemental": supplemental,
                        "image_alt_blocks": alt_blocks,
                    }
                )
        except (zipfile.BadZipFile, KeyError) as exc:
            raise OoxmlSecurityError(
                "changed-or-damaged-source",
                "Word 文件在扫描期间发生变化或已损坏",
            ) from exc
        after_hash = sha256_file(source_path)
        if before_hash != after_hash or after_hash != preflight.source_sha256:
            raise OoxmlSecurityError(
                "source-changed",
                "Word 原文件在扫描期间发生变化，请重新选择",
            )
        return model

    def export(
        self,
        document: DocumentModel,
        findings: Sequence[Finding],
        target_path: Path,
    ) -> None:
        if document.kind is not self.kind:
            raise ExportSafetyError("文档模型类型与 Word 适配器不一致")
        if sha256_file(document.source_path) != document.source_sha256:
            raise ExportSafetyError("原文件自扫描后发生变化，不能继续导出")
        ensure_export_ready(document, findings)
        target = secure_target_path(document.source_path, target_path, ".docx")
        temporary = temporary_export_path(target)
        texts = transformed_blocks(document, findings)
        original_image_hashes = {
            hashlib.sha256(image.content).hexdigest() for image in document.images
        }
        prepared_images = [prepare_image_for_export(image, findings) for image in document.images]
        try:
            _export_preserving_docx_layout(
                document,
                texts,
                prepared_images,
                temporary,
            )
            verify_rebuilt_ooxml(
                temporary,
                self.kind,
                forbidden_values=forbidden_originals(findings),
                original_image_hashes=original_image_hashes,
            )
            Document(str(temporary))
            finalize_export(temporary, target)
        finally:
            cleanup_temp(temporary)


def _export_preserving_docx_layout(
    document: DocumentModel,
    texts: dict[str, str],
    images: Sequence[ImageObject],
    target: Path,
) -> None:
    """Rewrite reviewed values inside the original safe layout container."""

    image_outputs: dict[str, bytes | None] = {}
    for image in images:
        sanitized = sanitize_image(image)
        image_outputs[image.source_part] = sanitized[0] if sanitized is not None else None

    with zipfile.ZipFile(document.source_path, mode="r") as incoming:
        raw_parts = {info.filename: incoming.read(info.filename) for info in incoming.infolist()}
        infos = {info.filename: info for info in incoming.infolist()}

    removed_parts = {
        name
        for name in raw_parts
        if _docx_part_must_be_removed(name)
        or (name in image_outputs and image_outputs[name] is None)
        or any(
            item.part == name
            and item.requires_decision
            and item.disposition is ObjectDisposition.REMOVE
            for item in document.inventory.parts
        )
    }
    removed_rel_ids = _removed_relationship_ids(raw_parts, removed_parts)
    rewritten: dict[str, bytes] = {}
    for name, data in raw_parts.items():
        if name in removed_parts:
            continue
        if name in image_outputs and image_outputs[name] is not None:
            rewritten[name] = cast(bytes, image_outputs[name])
            continue
        if name == "[Content_Types].xml":
            rewritten[name] = _sanitize_content_types(data, removed_parts, image_outputs)
            continue
        if name.endswith(".rels"):
            rewritten[name] = _sanitize_relationship_part(data, name, removed_parts)
            continue
        if name in {"word/document.xml"} or name.startswith(("word/header", "word/footer")):
            rewritten[name] = _sanitize_and_rewrite_word_part(
                data,
                name,
                document,
                texts,
                removed_rel_ids.get(name, set()),
            )
            continue
        if name.startswith("docProps/") and name.endswith(".xml"):
            rewritten[name] = _sanitize_document_properties(data)
            continue
        rewritten[name] = data

    with zipfile.ZipFile(target, mode="w", compression=zipfile.ZIP_DEFLATED) as outgoing:
        for name, data in rewritten.items():
            info = infos.get(name)
            if info is None:
                outgoing.writestr(name, data)
            else:
                outgoing.writestr(info, data)


def _docx_part_must_be_removed(name: str) -> bool:
    normalized = name.replace("\\", "/")
    return normalized in _REMOVED_DOCX_PART_NAMES or normalized.startswith(
        _REMOVED_DOCX_PART_PREFIXES
    )


def _relationship_owner_part(rel_name: str) -> str:
    normalized = rel_name.replace("\\", "/")
    if normalized == "_rels/.rels":
        return ""
    directory, filename = posixpath.split(normalized)
    if not directory.endswith("/_rels") or not filename.endswith(".rels"):
        return ""
    owner_directory = directory[: -len("/_rels")]
    return posixpath.join(owner_directory, filename[: -len(".rels")])


def _relationship_target_part(owner: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))


def _removed_relationship_ids(
    parts: dict[str, bytes],
    removed_parts: set[str],
) -> dict[str, set[str]]:
    result: dict[str, set[str]] = defaultdict(set)
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        owner = _relationship_owner_part(name)
        try:
            root = safe_xml_root(data, part_name=name)
        except OoxmlSecurityError:
            continue
        for relation in list(root):
            target = relation.get("Target", "")
            external = relation.get("TargetMode", "").casefold() == "external"
            target_part = _relationship_target_part(owner, target) if target else ""
            if external or target_part in removed_parts or _docx_part_must_be_removed(target_part):
                relation_id = relation.get("Id", "")
                if relation_id:
                    result[owner].add(relation_id)
    return result


def _sanitize_relationship_part(data: bytes, name: str, removed_parts: set[str]) -> bytes:
    root = safe_xml_root(data, part_name=name)
    owner = _relationship_owner_part(name)
    for relation in list(root):
        target = relation.get("Target", "")
        external = relation.get("TargetMode", "").casefold() == "external"
        target_part = _relationship_target_part(owner, target) if target else ""
        if external or target_part in removed_parts or _docx_part_must_be_removed(target_part):
            root.remove(relation)
    return cast(bytes, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))


def _sanitize_content_types(
    data: bytes,
    removed_parts: set[str],
    image_outputs: dict[str, bytes | None],
) -> bytes:
    root = safe_xml_root(data, part_name="[Content_Types].xml")
    for node in list(root):
        part_name = node.get("PartName", "").lstrip("/")
        if part_name and (part_name in removed_parts or _docx_part_must_be_removed(part_name)):
            root.remove(node)
    override_tag = f"{{{CONTENT_TYPES_NS}}}Override"
    existing_overrides = {
        node.get("PartName", "").lstrip("/"): node
        for node in root.findall(override_tag)
    }
    for part, content in image_outputs.items():
        if content is None:
            continue
        node = existing_overrides.get(part)
        if node is None:
            node = etree.SubElement(root, override_tag)
            node.set("PartName", f"/{part}")
        node.set("ContentType", "image/png")
    return cast(bytes, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))


def _sanitize_document_properties(data: bytes) -> bytes:
    root = safe_xml_root(data, part_name="docProps")
    for node in root.iter():
        if len(node) == 0 and node.text:
            node.text = ""
    return cast(bytes, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))


def _sanitize_and_rewrite_word_part(
    data: bytes,
    part: str,
    document: DocumentModel,
    texts: dict[str, str],
    removed_relation_ids: set[str],
) -> bytes:
    root = safe_xml_root(data, part_name=part)
    _remove_word_side_channels(root, removed_relation_ids)
    block_ids = _block_ids_for_word_part(document, part)
    paragraphs = _paragraphs_for_rewrite(root, part)
    if len(paragraphs) != len(block_ids):
        raise ExportSafetyError(
            f"Word 格式结构与扫描快照不一致（部件：{part}）"
        )
    for paragraph, block_id in zip(paragraphs, block_ids, strict=True):
        _rewrite_paragraph_text_in_place(paragraph, texts.get(block_id, ""))
    return cast(bytes, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))


def _block_ids_for_word_part(document: DocumentModel, part: str) -> list[str]:
    if part == "word/document.xml":
        block_ids: list[str] = []
        for item in document.structure.get("body", []):
            if item.get("kind") == "paragraph":
                block_ids.append(str(item["block_id"]))
            elif item.get("kind") == "table":
                for row in item.get("rows", []):
                    for cell in row:
                        block_ids.extend(str(value) for value in cell.get("blocks", []))
        return block_ids
    for group in document.structure.get("supplemental", []):
        if group.get("part") == part:
            return [str(value) for value in group.get("blocks", [])]
    return []


def _paragraphs_for_rewrite(root: etree._Element, part: str) -> list[etree._Element]:
    if part == "word/document.xml":
        body = root.find(_q(W_NS, "body"))
        if body is None:
            raise ExportSafetyError("Word 正文结构丢失")
        paragraphs: list[etree._Element] = []
        for child in body:
            if _local_name(child) == "p":
                paragraphs.append(child)
            elif _local_name(child) == "tbl":
                for row in child.findall(_q(W_NS, "tr")):
                    for cell in row.findall(_q(W_NS, "tc")):
                        paragraphs.extend(cell.findall(_q(W_NS, "p")))
        return paragraphs
    return [
        paragraph
        for paragraph in root.iter(_q(W_NS, "p"))
        if not _ancestor_has(paragraph, {"txbxContent"}) and _paragraph_text(paragraph)
    ]


def _remove_word_side_channels(root: etree._Element, removed_relation_ids: set[str]) -> None:
    for local_name in (
        "del",
        "moveFrom",
        "commentRangeStart",
        "commentRangeEnd",
        "commentReference",
        "footnoteReference",
        "endnoteReference",
        "object",
        "altChunk",
    ):
        for node in list(root.iter(_q(W_NS, local_name))):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for local_name in ("ins", "moveTo", "hyperlink"):
        for node in list(root.iter(_q(W_NS, local_name))):
            parent = node.getparent()
            if parent is None:
                continue
            index = parent.index(node)
            for child in list(node):
                node.remove(child)
                parent.insert(index, child)
                index += 1
            parent.remove(node)
    for run in list(root.iter(_q(W_NS, "r"))):
        if _run_is_hidden(run):
            parent = run.getparent()
            if parent is not None:
                parent.remove(run)
    for node in list(root.iter()):
        if any(
            attr in {_q(R_NS, "id"), _q(R_NS, "embed"), _q(R_NS, "link")}
            and value in removed_relation_ids
            for attr, value in node.attrib.items()
        ):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
                continue
        local_name = _local_name(node)
        if local_name in {"instrText", "fldChar"}:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
        elif local_name == "docPr":
            node.attrib.pop("descr", None)
            node.attrib.pop("title", None)
        elif local_name == "textpath":
            node.attrib.pop("string", None)
    for textbox in root.xpath(".//w:txbxContent", namespaces=NS):
        for child in list(textbox):
            textbox.remove(child)


def _rewrite_paragraph_text_in_place(paragraph: etree._Element, replacement: str) -> None:
    text_nodes = [
        node
        for node in paragraph.iter(_q(W_NS, "t"))
        if not _ancestor_has(node, {"txbxContent", "del", "moveFrom"}) and not _run_is_hidden(node)
    ]
    original = "".join(node.text or "" for node in text_nodes)
    if original == replacement:
        return
    if not text_nodes:
        if replacement:
            run = etree.SubElement(paragraph, _q(W_NS, "r"))
            node = etree.SubElement(run, _q(W_NS, "t"))
            node.text = replacement
        return
    char_nodes: list[int] = []
    for index, node in enumerate(text_nodes):
        char_nodes.extend([index] * len(node.text or ""))
    assigned = [""] * len(text_nodes)
    matcher = SequenceMatcher(a=original, b=replacement, autojunk=False)
    for tag, i1, _i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for offset, character in enumerate(replacement[j1:j2]):
                assigned[char_nodes[i1 + offset]] += character
        elif tag in {"replace", "insert"}:
            if i1 < len(char_nodes):
                target_index = char_nodes[i1]
            elif char_nodes:
                target_index = char_nodes[-1]
            else:
                target_index = 0
            assigned[target_index] += replacement[j1:j2]
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    for node, value in zip(text_nodes, assigned, strict=True):
        node.text = value
        if value.startswith(" ") or value.endswith(" "):
            node.set(xml_space, "preserve")
        else:
            node.attrib.pop(xml_space, None)


def _remove_template_custom_xml(path: Path) -> None:
    """Remove custom XML inherited from python-docx's built-in blank template."""

    clean_path = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.clean{path.suffix}")
    try:
        with (
            zipfile.ZipFile(path, mode="r") as incoming,
            zipfile.ZipFile(
                clean_path,
                mode="w",
                compression=zipfile.ZIP_DEFLATED,
            ) as outgoing,
        ):
            for info in incoming.infolist():
                name = info.filename.replace("\\", "/")
                if name.startswith("customXml/"):
                    continue
                data = incoming.read(info)
                if name == "[Content_Types].xml":
                    root = safe_xml_root(data, part_name=name)
                    for element in list(root):
                        part_name = element.get("PartName", "").lstrip("/")
                        if part_name.startswith("customXml/"):
                            root.remove(element)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                elif name.endswith(".rels"):
                    root = safe_xml_root(data, part_name=name)
                    for relationship in list(root):
                        relation_type = relationship.get("Type", "").casefold()
                        target = relationship.get("Target", "").replace("\\", "/")
                        if (
                            relation_type.endswith("/customxml")
                            or "customxml/" in target.casefold()
                        ):
                            root.remove(relationship)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                outgoing.writestr(info, data)
        clean_path.replace(path)
    finally:
        cleanup_temp(clean_path)


def _reset_core_properties(document: Any) -> None:
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = ""
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    properties.content_status = ""
    properties.identifier = ""
    properties.language = ""
    properties.version = ""
    fixed_time = datetime(2000, 1, 1)
    properties.created = fixed_time
    properties.modified = fixed_time
    properties.last_printed = fixed_time
    properties.revision = 1


def _resolve_word_style(style: dict[str, Any]) -> str | None:
    source_style = str(style.get("style_id", ""))
    lowered = source_style.casefold()
    if lowered in {"title", "subtitle", "quote", "intensequote"}:
        return {
            "title": "Title",
            "subtitle": "Subtitle",
            "quote": "Quote",
            "intensequote": "Intense Quote",
        }[lowered]
    heading = re.fullmatch(r"heading\s*([1-9])", source_style, flags=re.IGNORECASE)
    if heading:
        return f"Heading {heading.group(1)}"
    if style.get("numbered"):
        return "List Number"
    return None


def _write_word_paragraph(
    paragraph: Any,
    text: str,
    style: dict[str, Any],
) -> None:
    resolved_style = _resolve_word_style(style)
    if resolved_style:
        with suppress(KeyError):
            paragraph.style = resolved_style
    alignment = str(style.get("alignment", "")).casefold()
    alignment_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
    }
    if alignment in alignment_map:
        paragraph.alignment = alignment_map[alignment]
    pf = paragraph.paragraph_format
    spacing_before = style.get("spacing_before")
    if spacing_before is not None:
        with suppress(TypeError, ValueError):
            pf.space_before = Twips(int(spacing_before))
    spacing_after = style.get("spacing_after")
    if spacing_after is not None:
        with suppress(TypeError, ValueError):
            pf.space_after = Twips(int(spacing_after))
    first_line_indent = style.get("first_line_indent")
    if first_line_indent is not None:
        with suppress(TypeError, ValueError):
            pf.first_line_indent = Twips(int(first_line_indent))

    run_formats = style.get("run_formats", [])
    chunks = _styled_word_chunks(text, run_formats) if run_formats else [(text, style)]
    for chunk, formatting in chunks:
        if chunk:
            _apply_word_run_format(paragraph.add_run(chunk), formatting)


def _styled_word_chunks(
    transformed: str,
    run_formats: object,
) -> list[tuple[str, dict[str, Any]]]:
    if not isinstance(run_formats, list):
        return []
    formats = [item for item in run_formats if isinstance(item, dict)]
    original = "".join(str(item.get("text", "")) for item in formats)
    if not original or not formats:
        return []

    boundaries: list[tuple[int, int, dict[str, Any]]] = []
    offset = 0
    for item in formats:
        value = str(item.get("text", ""))
        boundaries.append((offset, offset + len(value), item))
        offset += len(value)

    def format_at(index: int) -> dict[str, Any]:
        probe = min(max(index, 0), max(len(original) - 1, 0))
        for start, end, item in boundaries:
            if start <= probe < end:
                return item
        return formats[-1]

    chunks: list[tuple[str, dict[str, Any]]] = []
    matcher = SequenceMatcher(a=original, b=transformed, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            cursor = i1
            target_cursor = j1
            while cursor < i2:
                formatting = format_at(cursor)
                boundary = next(
                    (
                        end
                        for start, end, item in boundaries
                        if item is formatting and start <= cursor < end
                    ),
                    i2,
                )
                length = min(boundary, i2) - cursor
                chunks.append(
                    (transformed[target_cursor : target_cursor + length], formatting)
                )
                cursor += length
                target_cursor += length
        elif tag in {"replace", "insert"} and j1 < j2:
            chunks.append((transformed[j1:j2], format_at(i1)))

    merged: list[tuple[str, dict[str, Any]]] = []
    for chunk, formatting in chunks:
        signature = {key: value for key, value in formatting.items() if key != "text"}
        previous_signature = (
            {key: value for key, value in merged[-1][1].items() if key != "text"}
            if merged
            else None
        )
        if merged and previous_signature == signature:
            merged[-1] = (merged[-1][0] + chunk, merged[-1][1])
        else:
            merged.append((chunk, formatting))
    return merged


def _apply_word_run_format(run: Any, style: dict[str, Any]) -> None:
    run.bold = bool(style.get("bold", False))
    run.italic = bool(style.get("italic", False))
    run.underline = bool(style.get("underline", False))
    font_name = str(style.get("font_name", ""))
    if font_name:
        run.font.name = font_name
    font_size = style.get("font_size", "")
    if font_size:
        with suppress(TypeError, ValueError):
            run.font.size = Pt(int(str(font_size)) / 2)
    font_color = str(style.get("font_color", ""))
    if len(font_color) == 6:
        with suppress(ValueError):
            run.font.color.rgb = RGBColor(
                int(font_color[0:2], 16),
                int(font_color[2:4], 16),
                int(font_color[4:6], 16),
            )


def _append_word_images(
    paragraph: Any,
    images: Sequence[ImageObject],
    document: DocumentModel,
    texts: dict[str, str],
    emitted: set[str],
) -> None:
    alt_blocks = document.structure.get("image_alt_blocks", {})
    for image in images:
        sanitized = sanitize_image(image)
        emitted.add(image.id)
        if sanitized is None:
            continue
        content, _media_type = sanitized
        run = paragraph.add_run()
        shape = run.add_picture(io.BytesIO(content), width=Inches(2.5))
        alt_block_id = alt_blocks.get(image.id)
        alt_text = texts.get(alt_block_id, "") if alt_block_id else ""
        legend = image_replacement_legend(image)
        description = "；".join(value for value in (alt_text, legend) if value)
        shape._inline.docPr.set("name", "经处理图片")
        shape._inline.docPr.set("title", "")
        shape._inline.docPr.set("descr", description)
        if legend:
            paragraph.add_run(f"\n图片替代图例：{legend}")


def _rebuild_table(
    output: Any,
    item: dict[str, Any],
    texts: dict[str, str],
    image_by_block: dict[str | None, list[ImageObject]],
    document: DocumentModel,
    emitted_images: set[str],
) -> None:
    rows = item.get("rows", [])
    if not rows:
        return
    column_count = max(
        (
            int(cell.get("column", 0)) + int(cell.get("grid_span", 1))
            for row in rows
            for cell in row
        ),
        default=1,
    )
    table = output.add_table(rows=len(rows), cols=max(1, column_count))
    with suppress(KeyError):
        table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for cell_data in row:
            column = int(cell_data.get("column", 0))
            grid_span = max(1, int(cell_data.get("grid_span", 1)))
            cell = table.cell(row_index, column)
            if grid_span > 1 and column + grid_span - 1 < column_count:
                cell = cell.merge(table.cell(row_index, column + grid_span - 1))
            block_ids = [str(block_id) for block_id in cell_data.get("blocks", [])]
            if not block_ids:
                continue
            cell.text = ""
            paragraph = cell.paragraphs[0]
            for block_index, block_id in enumerate(block_ids):
                target_paragraph = paragraph if block_index == 0 else cell.add_paragraph()
                block = next(
                    (candidate for candidate in document.blocks if candidate.id == block_id),
                    None,
                )
                _write_word_paragraph(
                    target_paragraph,
                    texts.get(block_id, ""),
                    block.style if block is not None else {},
                )
                _append_word_images(
                    target_paragraph,
                    image_by_block.get(block_id, []),
                    document,
                    texts,
                    emitted_images,
                )


def _rebuild_headers_and_footers(
    output: Any,
    document: DocumentModel,
    texts: dict[str, str],
) -> None:
    if not output.sections:
        return
    section = output.sections[0]
    for group in document.structure.get("supplemental", []):
        part = str(group.get("part", ""))
        kind = _supplemental_kind(part)
        if kind not in {"header", "footer"}:
            continue
        disposition = item_disposition(document, part)
        if disposition is not ObjectDisposition.VISIBLE_NOTE:
            continue
        target = section.header if kind == "header" else section.footer
        if target.paragraphs:
            target.paragraphs[0].text = ""
        for index, block_id in enumerate(group.get("blocks", [])):
            paragraph = (
                target.paragraphs[0] if index == 0 and target.paragraphs else target.add_paragraph()
            )
            paragraph.add_run(texts.get(str(block_id), ""))


def _append_visible_supplemental(
    output: Any,
    document: DocumentModel,
    texts: dict[str, str],
) -> None:
    notes: list[tuple[str, str]] = []
    for group in document.structure.get("supplemental", []):
        part = str(group.get("part", ""))
        kind = _supplemental_kind(part)
        if kind in {"header", "footer"}:
            continue
        if item_disposition(document, part) is not ObjectDisposition.VISIBLE_NOTE:
            continue
        label = {
            "footnote": "脚注",
            "endnote": "尾注",
            "comment": "原批注",
            "supplemental": "附注",
        }.get(kind, "可见附注")
        for block_id in group.get("blocks", []):
            value = texts.get(str(block_id), "")
            if value:
                notes.append((label, value))
    if not notes:
        return
    output.add_heading("可见附注", level=1)
    for label, value in notes:
        output.add_paragraph(f"【{label}】{value}")
