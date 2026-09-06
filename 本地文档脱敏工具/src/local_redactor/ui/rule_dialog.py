from __future__ import annotations

import re
from collections.abc import Callable
from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from PySide6.QtCore import Qt, Signal
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from local_redactor.builtin_rules import BUILTIN_RULES
from local_redactor.rule_library import (
    ActionKind,
    ConflictPolicy,
    MatchMode,
    RuleConflictError,
    RuleDefinition,
    RuleImportConflictError,
    RuleImportResult,
    RuleKind,
    RuleLibrary,
    RuleLibraryError,
    RuleSession,
    RuleStore,
    RuleStoreError,
    RuleValidationError,
    import_rules,
    is_preset_rule,
    restore_default_rules,
)

_MATCH_LABELS = {
    MatchMode.CONTAINS: "包含",
    MatchMode.EXACT: "等于",
    MatchMode.REGEX: "符合格式",
    MatchMode.ANY_KEYWORD: "包含任一关键词",
    MatchMode.ALL_KEYWORDS: "同时包含全部关键词",
    MatchMode.MANUAL: "仅人工选择",
}
_ACTION_LABELS = {
    ActionKind.FIXED_REPLACEMENT: "固定代号",
    ActionKind.MASK_MIDDLE: "中间星号",
    ActionKind.SEQUENCE_CODE: "顺序代号",
    ActionKind.DELETE: "删除",
}
_SCOPE_LABELS = {
    "all": "全部内容",
    "text": "正文",
    "cell": "表格单元格",
    "ocr": "图片文字",
    "metadata": "文件属性",
    "hidden": "隐藏内容",
}
_TEMPLATE_HEADERS = (
    "规则类型",
    "规则名称",
    "匹配方式",
    "关键词",
    "处理方式",
    "替换内容",
    "规则编号",
    "必须处理",
    "启用",
    "优先级",
    "正例",
    "反例",
    "适用范围",
    "区分大小写",
    "保留开头",
    "保留结尾",
    "代号位数",
    "内容类别",
)
_TEMPLATE_WIDTHS = {
    "规则类型": 14,
    "规则名称": 24,
    "匹配方式": 18,
    "关键词": 30,
    "处理方式": 20,
    "替换内容": 24,
    "规则编号": 24,
    "必须处理": 12,
    "启用": 10,
    "优先级": 10,
    "正例": 30,
    "反例": 30,
    "适用范围": 20,
    "区分大小写": 14,
    "保留开头": 12,
    "保留结尾": 12,
    "代号位数": 12,
    "内容类别": 16,
}


class RuleTemplateError(RuleLibraryError):
    """The blank import template could not be safely created."""


def export_blank_rule_template(path: Path) -> Path:
    """Create a new, blank and import-compatible XLSX rule template.

    Existing files are never overwritten. The first worksheet contains only
    the supported headers; the second worksheet explains the user-facing
    values accepted by the importer.
    """

    target = Path(path)
    if target.suffix.casefold() != ".xlsx":
        raise RuleTemplateError("空白模板必须保存为 .xlsx 文件")
    if not target.parent.is_dir():
        raise RuleTemplateError("模板保存目录不存在")
    if target.exists():
        raise RuleTemplateError("模板文件已存在，请更换名称后再导出")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "规则导入模板"
    worksheet.append(list(_TEMPLATE_HEADERS))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(_TEMPLATE_HEADERS))}1"
    worksheet.row_dimensions[1].height = 28
    for column_index, header in enumerate(_TEMPLATE_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B62")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[get_column_letter(column_index)].width = (
            _TEMPLATE_WIDTHS[header]
        )

    instructions = workbook.create_sheet("填写说明")
    instruction_rows = (
        ("项目", "填写方式"),
        ("规则类型", "填写“固定替换”或“判断标准”。"),
        (
            "匹配方式",
            "填写“包含”“精确”“受控正则”“任一关键词”“全部关键词”或“仅人工”。"
            "界面中的“等于”对应“精确”，“符合格式”对应“受控正则”。",
        ),
        (
            "处理方式",
            "填写“固定代号”“中间星号”“任务内顺序代号”或“删除”。",
        ),
        (
            "固定替换",
            "匹配方式必须填写“包含”，处理方式必须填写“固定代号”。",
        ),
        (
            "判断标准",
            "正例必填，且至少填写一个确实会命中该判断标准的正确示例。",
        ),
        ("多个值", "关键词、正例、反例和多个适用范围使用竖线 | 分隔。"),
        (
            "适用范围",
            "填写 all、text、cell、ocr、metadata 或 hidden；all 不能与其他范围混用。",
        ),
        ("布尔字段", "必须处理、启用、区分大小写填写“是”或“否”。"),
        ("安全说明", "不要在模板中填写公式；导入时不会执行公式。"),
    )
    for row in instruction_rows:
        instructions.append(row)
    instructions.freeze_panes = "A2"
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 88
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B62")
        cell.alignment = Alignment(horizontal="center")
    for row in instructions.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    created = False
    try:
        with target.open("xb") as stream:
            created = True
            workbook.save(stream)
    except FileExistsError as exc:
        raise RuleTemplateError("模板文件已存在，请更换名称后再导出") from exc
    except OSError as exc:
        if created:
            target.unlink(missing_ok=True)
        raise RuleTemplateError("无法保存空白规则模板") from exc
    finally:
        workbook.close()
    return target


class RuleEditorDialog(QDialog):
    """Small user-language form for one fixed or criterion rule."""

    def __init__(
        self,
        kind: RuleKind,
        rule: RuleDefinition | None = None,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        if rule is not None and rule.kind is not kind:
            raise ValueError("待编辑规则类型与当前页签不一致")
        self.kind = kind
        self.original_rule = rule
        self.result_rule: RuleDefinition | None = None
        self.setWindowTitle("编辑固定替换" if rule else "新增固定替换")
        if kind is RuleKind.STANDARD:
            self.setWindowTitle("编辑判断标准" if rule else "新增判断标准")
        self.setModal(True)
        self.setMinimumWidth(680)

        layout = QVBoxLayout(self)
        explanation = QLabel(
            "填写一个完整原词和固定代号；正文、表格或图片文字中出现该词时替换。"
            if kind is RuleKind.FIXED
            else "用日常语言描述何时命中、如何处理。"
            "多个关键词或示例用竖线 | 分隔。"
        )
        explanation.setWordWrap(True)
        layout.addWidget(explanation)

        form = QFormLayout()
        self.form = form
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.AllNonFixedFieldsGrow)
        self.name_edit = QLineEdit()
        self.name_edit.setObjectName("ruleNameEdit")
        self.name_edit.setPlaceholderText("例如：单位简称统一代号")
        form.addRow("规则名称", self.name_edit)

        self.match_mode_combo = QComboBox()
        self.match_mode_combo.setObjectName("matchModeCombo")
        match_choices = (
            ((MatchMode.CONTAINS, _MATCH_LABELS[MatchMode.CONTAINS]),)
            if kind is RuleKind.FIXED
            else tuple(_MATCH_LABELS.items())
        )
        for mode, label in match_choices:
            self.match_mode_combo.addItem(label, mode.value)
        if kind is RuleKind.STANDARD:
            form.addRow("判断方式", self.match_mode_combo)
        else:
            self.match_mode_combo.hide()

        self.pattern_edit = QLineEdit()
        self.pattern_edit.setObjectName("rulePatternEdit")
        self.pattern_edit.setPlaceholderText(
            "要替换的完整原词"
            if kind is RuleKind.FIXED
            else "关键词；多个关键词用 | 分隔"
        )
        form.addRow(
            "完整原词" if kind is RuleKind.FIXED else "关键词或格式",
            self.pattern_edit,
        )

        self.action_combo = QComboBox()
        self.action_combo.setObjectName("actionCombo")
        action_choices = (
            (
                (
                    ActionKind.FIXED_REPLACEMENT,
                    _ACTION_LABELS[ActionKind.FIXED_REPLACEMENT],
                ),
            )
            if kind is RuleKind.FIXED
            else tuple(_ACTION_LABELS.items())
        )
        for action, label in action_choices:
            self.action_combo.addItem(label, action.value)
        if kind is RuleKind.STANDARD:
            form.addRow("处理方式", self.action_combo)
        else:
            self.action_combo.hide()

        self.replacement_edit = QLineEdit()
        self.replacement_edit.setObjectName("replacementEdit")
        self.replacement_edit.setPlaceholderText("例如：GA、机构甲、项目")
        form.addRow("固定代号或前缀", self.replacement_edit)

        self.mask_row = QWidget()
        mask_layout = QHBoxLayout(self.mask_row)
        mask_layout.setContentsMargins(0, 0, 0, 0)
        self.keep_prefix_spin = QSpinBox()
        self.keep_prefix_spin.setRange(0, 64)
        self.keep_prefix_spin.setValue(1)
        self.keep_prefix_spin.setPrefix("开头 ")
        self.keep_prefix_spin.setSuffix(" 字")
        self.keep_suffix_spin = QSpinBox()
        self.keep_suffix_spin.setRange(0, 64)
        self.keep_suffix_spin.setValue(1)
        self.keep_suffix_spin.setPrefix("结尾 ")
        self.keep_suffix_spin.setSuffix(" 字")
        mask_layout.addWidget(self.keep_prefix_spin)
        mask_layout.addWidget(self.keep_suffix_spin)
        mask_layout.addStretch(1)
        if kind is RuleKind.STANDARD:
            form.addRow("中间星号保留", self.mask_row)
        else:
            self.mask_row.hide()

        self.code_width_spin = QSpinBox()
        self.code_width_spin.setRange(1, 8)
        self.code_width_spin.setValue(2)
        self.code_width_spin.setSuffix(" 位")
        if kind is RuleKind.STANDARD:
            form.addRow("顺序代号位数", self.code_width_spin)
        else:
            self.code_width_spin.hide()

        flags_row = QWidget()
        flags_layout = QHBoxLayout(flags_row)
        flags_layout.setContentsMargins(0, 0, 0, 0)
        self.mandatory_check = QCheckBox("必须替换，不能保留原文")
        self.mandatory_check.setObjectName("mandatoryCheck")
        self.mandatory_check.setChecked(True)
        self.enabled_check = QCheckBox("启用")
        self.enabled_check.setObjectName("enabledCheck")
        self.enabled_check.setChecked(True)
        self.case_sensitive_check = QCheckBox("英文区分大小写")
        self.case_sensitive_check.setChecked(True)
        flags_layout.addWidget(self.mandatory_check)
        flags_layout.addWidget(self.enabled_check)
        flags_layout.addWidget(self.case_sensitive_check)
        flags_layout.addStretch(1)
        form.addRow("规则状态", flags_row)

        scopes_row = QWidget()
        scopes_layout = QHBoxLayout(scopes_row)
        scopes_layout.setContentsMargins(0, 0, 0, 0)
        self.scope_checks: dict[str, QCheckBox] = {}
        for scope, label in _SCOPE_LABELS.items():
            checkbox = QCheckBox(label)
            checkbox.setObjectName(f"scope_{scope}")
            checkbox.toggled.connect(
                lambda checked, selected=scope: self._scope_toggled(selected, checked)
            )
            self.scope_checks[scope] = checkbox
            scopes_layout.addWidget(checkbox)
        scopes_layout.addStretch(1)
        form.addRow("适用范围", scopes_row)

        self.positive_examples_edit = QLineEdit()
        self.positive_examples_edit.setObjectName("positiveExamplesEdit")
        self.positive_examples_edit.setPlaceholderText(
            "必填；填写至少一个会命中的正确示例"
        )

        self.negative_examples_edit = QLineEdit()
        self.negative_examples_edit.setObjectName("negativeExamplesEdit")
        self.negative_examples_edit.setPlaceholderText("可选；多个反例用 | 分隔")
        if kind is RuleKind.STANDARD:
            form.addRow("正确示例（原内容）", self.positive_examples_edit)
            form.addRow("反例", self.negative_examples_edit)

        self.category_edit = QLineEdit()
        self.category_edit.setPlaceholderText("可选，例如：单位、项目、账号")
        form.addRow("内容类别", self.category_edit)
        layout.addLayout(form)

        preview_row = QHBoxLayout()
        self.sample_edit = QLineEdit()
        self.sample_edit.setObjectName("ruleSampleEdit")
        self.sample_edit.setPlaceholderText("输入一段虚构样例，保存前先验证结果")
        preview_button = QPushButton("验证样例")
        preview_button.setObjectName("ruleSamplePreviewButton")
        preview_button.clicked.connect(self._preview_sample)
        preview_row.addWidget(self.sample_edit, 1)
        preview_row.addWidget(preview_button)
        layout.addLayout(preview_row)
        self.sample_result = QLabel("样例只在当前窗口内验证，不会保存。")
        self.sample_result.setObjectName("mutedText")
        self.sample_result.setWordWrap(True)
        layout.addWidget(self.sample_result)

        self.form_error = QLabel()
        self.form_error.setObjectName("ruleFormError")
        self.form_error.setWordWrap(True)
        self.form_error.setStyleSheet("color: #B42318;")
        self.form_error.hide()
        layout.addWidget(self.form_error)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save
            | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.button(QDialogButtonBox.StandardButton.Save).setText("保存")
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        self.match_mode_combo.currentIndexChanged.connect(self._sync_match_fields)
        self.action_combo.currentIndexChanged.connect(self._sync_action_fields)
        self._load_rule(rule)
        self._sync_match_fields()
        self._sync_action_fields()

    def accept(self) -> None:
        try:
            self.result_rule = self.rule_definition()
        except RuleLibraryError as exc:
            self.form_error.setText(f"无法保存规则：{exc}")
            self.form_error.show()
            return
        self.form_error.clear()
        self.form_error.hide()
        super().accept()

    def _preview_sample(self) -> None:
        sample = self.sample_edit.text()
        if not sample:
            self.sample_result.setText("请先输入不含真实敏感信息的虚构样例。")
            return
        try:
            rule = self.rule_definition()
            result = RuleSession(RuleLibrary((rule,))).apply(sample)
        except RuleLibraryError as exc:
            self.sample_result.setText(f"暂时不能验证：{exc}")
            return
        if not result.matches:
            self.sample_result.setText("样例未命中，请检查关键词、格式和作用范围。")
            return
        self.sample_result.setText(f"验证结果：{result.replacement}")

    def rule_definition(self) -> RuleDefinition:
        match_mode = MatchMode(str(self.match_mode_combo.currentData()))
        action = ActionKind(str(self.action_combo.currentData()))
        patterns = () if match_mode is MatchMode.MANUAL else _split_user_values(
            self.pattern_edit.text()
        )
        if match_mode in {MatchMode.CONTAINS, MatchMode.EXACT, MatchMode.REGEX} and (
            len(patterns) != 1
        ):
            raise RuleLibraryError("“包含、等于、符合格式”每条规则只能填写一个关键词或格式")
        if action in {ActionKind.MASK_MIDDLE, ActionKind.DELETE}:
            replacement = ""
        else:
            replacement = self.replacement_edit.text()
        scopes = self._selected_scopes()
        original = self.original_rule
        if self.kind is RuleKind.FIXED:
            if len(patterns) != 1:
                raise RuleValidationError("固定替换必须填写一个完整原词")
            return RuleDefinition.fixed(
                patterns[0],
                replacement,
                name=self.name_edit.text(),
                match_mode=MatchMode.CONTAINS,
                mandatory=self.mandatory_check.isChecked(),
                enabled=self.enabled_check.isChecked(),
                priority=original.priority if original is not None else 0,
                applies_to=scopes,
                case_sensitive=self.case_sensitive_check.isChecked(),
                category=self.category_edit.text(),
                rule_id=original.id if original is not None else _new_rule_id(),
            )
        positive_examples = _split_user_values(self.positive_examples_edit.text())
        if not positive_examples:
            raise RuleValidationError("判断标准至少需要填写一个正确示例（正例）")
        return RuleDefinition(
            id=original.id if original is not None else _new_rule_id(),
            name=self.name_edit.text(),
            kind=self.kind,
            match_mode=match_mode,
            patterns=patterns,
            action=action,
            replacement=replacement,
            mandatory=self.mandatory_check.isChecked(),
            enabled=self.enabled_check.isChecked(),
            priority=original.priority if original is not None else 0,
            positive_examples=positive_examples,
            negative_examples=_split_user_values(self.negative_examples_edit.text()),
            applies_to=scopes,
            case_sensitive=self.case_sensitive_check.isChecked(),
            keep_prefix=self.keep_prefix_spin.value(),
            keep_suffix=self.keep_suffix_spin.value(),
            code_width=self.code_width_spin.value(),
            category=self.category_edit.text(),
        )

    def _load_rule(self, rule: RuleDefinition | None) -> None:
        if rule is None:
            self.scope_checks["all"].setChecked(True)
            return
        self.name_edit.setText(rule.name)
        _select_combo_value(
            self.match_mode_combo,
            (
                MatchMode.CONTAINS.value
                if self.kind is RuleKind.FIXED
                else rule.match_mode.value
            ),
        )
        self.pattern_edit.setText(" | ".join(rule.patterns))
        _select_combo_value(
            self.action_combo,
            (
                ActionKind.FIXED_REPLACEMENT.value
                if self.kind is RuleKind.FIXED
                else rule.action.value
            ),
        )
        self.replacement_edit.setText(rule.replacement)
        self.keep_prefix_spin.setValue(rule.keep_prefix)
        self.keep_suffix_spin.setValue(rule.keep_suffix)
        self.code_width_spin.setValue(rule.code_width)
        self.mandatory_check.setChecked(rule.mandatory)
        self.enabled_check.setChecked(rule.enabled)
        self.case_sensitive_check.setChecked(rule.case_sensitive)
        for scope, checkbox in self.scope_checks.items():
            checkbox.setChecked(scope in rule.applies_to)
        self.positive_examples_edit.setText(" | ".join(rule.positive_examples))
        self.negative_examples_edit.setText(" | ".join(rule.negative_examples))
        self.category_edit.setText(rule.category)

    def _selected_scopes(self) -> tuple[str, ...]:
        if self.scope_checks["all"].isChecked():
            return ("all",)
        selected = tuple(
            scope
            for scope in _SCOPE_LABELS
            if scope != "all" and self.scope_checks[scope].isChecked()
        )
        if not selected:
            raise RuleLibraryError("请至少选择一个适用范围")
        return selected

    def _scope_toggled(self, scope: str, checked: bool) -> None:
        if scope == "all" and checked:
            for other_scope, checkbox in self.scope_checks.items():
                if other_scope != "all":
                    checkbox.setChecked(False)
        elif scope != "all" and checked:
            self.scope_checks["all"].setChecked(False)

    def _sync_match_fields(self) -> None:
        mode = MatchMode(str(self.match_mode_combo.currentData()))
        manual = mode is MatchMode.MANUAL
        self.form.setRowVisible(self.pattern_edit, not manual)
        if manual:
            self.pattern_edit.clear()

    def _sync_action_fields(self) -> None:
        action = ActionKind(str(self.action_combo.currentData()))
        needs_replacement = action in {
            ActionKind.FIXED_REPLACEMENT,
            ActionKind.SEQUENCE_CODE,
        }
        self.form.setRowVisible(self.replacement_edit, needs_replacement)
        if self.kind is RuleKind.STANDARD:
            self.form.setRowVisible(
                self.mask_row,
                action is ActionKind.MASK_MIDDLE,
            )
            self.form.setRowVisible(
                self.code_width_spin,
                action is ActionKind.SEQUENCE_CODE,
            )


class RuleLibraryDialog(QDialog):
    """System-rule viewer and two-tab editor backed by an encrypted RuleStore."""

    library_changed = Signal(object)

    def __init__(self, store: RuleStore, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.store = store
        self._library: RuleLibrary | None = None
        self._tables: dict[RuleKind, QTableWidget] = {}
        self._edit_buttons: dict[RuleKind, QPushButton] = {}
        self._toggle_buttons: dict[RuleKind, QPushButton] = {}
        self._delete_buttons: dict[RuleKind, QPushButton] = {}

        self.setWindowTitle("规则库")
        self.setModal(True)
        self.resize(980, 640)
        layout = QVBoxLayout(self)

        title = QLabel("规则库")
        title.setStyleSheet("font-size: 20px; font-weight: 700;")
        layout.addWidget(title)
        copy = QLabel(
            "规则只保存在当前 Windows 用户的本机加密存储中。"
            "固定替换用于明确词语，判断标准用于可复用的命中与处理规则。"
        )
        copy.setWordWrap(True)
        layout.addWidget(copy)
        builtins = QLabel(
            "系统识别标准完整列在第一个页签中，只读且始终生效；包括姓名、联系方式、"
            "证件、银行卡、密码密钥、车牌和网络信息等。时间保持原文。"
        )
        builtins.setObjectName("mutedText")
        builtins.setWordWrap(True)
        layout.addWidget(builtins)

        feedback_row = QHBoxLayout()
        self.feedback_label = QLabel()
        self.feedback_label.setObjectName("ruleLibraryFeedback")
        self.feedback_label.setWordWrap(True)
        feedback_row.addWidget(self.feedback_label, 1)
        self.restore_button = QPushButton("恢复默认预置")
        self.restore_button.setToolTip("只恢复缺失的系统预置，不覆盖我的规则")
        self.restore_button.clicked.connect(self._restore_clicked)
        feedback_row.addWidget(self.restore_button)
        self.retry_button = QPushButton("重新读取")
        self.retry_button.clicked.connect(self._reload_clicked)
        self.retry_button.hide()
        feedback_row.addWidget(self.retry_button)
        layout.addLayout(feedback_row)

        self.tabs = QTabWidget()
        self.tabs.setObjectName("ruleLibraryTabs")
        self.tabs.addTab(
            self._build_builtin_tab(),
            f"系统内置识别 {len(BUILTIN_RULES)}",
        )
        self.tabs.addTab(self._build_tab(RuleKind.FIXED), "固定替换")
        self.tabs.addTab(self._build_tab(RuleKind.STANDARD), "判断标准")
        layout.addWidget(self.tabs, 1)

        close_buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        close_buttons.button(QDialogButtonBox.StandardButton.Close).setText("关闭")
        close_buttons.rejected.connect(self.reject)
        layout.addWidget(close_buttons)

        try:
            self.reload()
        except RuleLibraryError as exc:
            self._set_load_error(exc)

    @property
    def library(self) -> RuleLibrary:
        if self._library is None:
            raise RuleStoreError("规则库尚未成功读取，不能进行修改")
        return self._library.snapshot()

    def reload(self) -> RuleLibrary:
        library = self.store.load()
        self._library = library
        self.tabs.setEnabled(True)
        self.retry_button.hide()
        self._set_feedback(
            f"已读取 {len(library.rules)} 条规则，规则已保存。",
            error=False,
        )
        self._refresh_tables()
        return library.snapshot()

    def add_rule(self, rule: RuleDefinition) -> RuleLibrary:
        _validate_dialog_rule(rule)
        return self._commit(lambda library: library.add(rule))

    def update_rule(
        self,
        rule: RuleDefinition,
        *,
        expected_rule: RuleDefinition | None = None,
    ) -> RuleLibrary:
        _validate_dialog_rule(rule)
        expected = expected_rule or self._snapshot_rule(rule.id)

        def operation(library: RuleLibrary) -> RuleLibrary:
            self._require_unchanged(library, expected)
            return library.update(rule)

        return self._commit(operation)

    def set_rule_enabled(
        self,
        rule_id: str,
        enabled: bool,
        *,
        expected_rule: RuleDefinition | None = None,
    ) -> RuleLibrary:
        expected = expected_rule or self._snapshot_rule(rule_id)

        def operation(library: RuleLibrary) -> RuleLibrary:
            current = self._require_unchanged(library, expected)
            if current.enabled is enabled:
                return library
            return library.set_enabled(rule_id, enabled)

        return self._commit(operation)

    def delete_rule(
        self,
        rule_id: str,
        *,
        expected_rule: RuleDefinition | None = None,
    ) -> RuleLibrary:
        expected = expected_rule or self._snapshot_rule(rule_id)

        def operation(library: RuleLibrary) -> RuleLibrary:
            self._require_unchanged(library, expected)
            return library.delete(rule_id)

        return self._commit(operation)

    def import_file(
        self,
        path: Path,
        *,
        conflict_policy: ConflictPolicy = "error",
    ) -> RuleImportResult:
        current = self.store.load()
        result = import_rules(
            path,
            current,
            conflict_policy=conflict_policy,
        )
        for rule in result.library.rules:
            _validate_dialog_rule(rule)
        if result.library != current:
            self.store.save_if_revision(
                result.library,
                expected_revision=current.revision,
            )
            self._library = result.library
            self._refresh_tables()
            self.library_changed.emit(result.library.snapshot())
        else:
            self._library = current
            self._refresh_tables()
        return result

    def export_template(self, path: Path) -> Path:
        return export_blank_rule_template(path)

    def _commit(
        self,
        operation: Callable[[RuleLibrary], RuleLibrary],
    ) -> RuleLibrary:
        current = self.store.load()
        updated = operation(current)
        changed = updated != current
        if changed:
            self.store.save_if_revision(
                updated,
                expected_revision=current.revision,
            )
        self._library = updated
        self._refresh_tables()
        if changed:
            self.library_changed.emit(updated.snapshot())
        return updated.snapshot()

    def _build_builtin_tab(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        summary = QLabel(
            "这些是程序自带的识别与处理标准，不能编辑、停用或删除。"
            "结构校验充分的项目可批量采用；地点、单位等语义候选仍保留人工判断。"
        )
        summary.setWordWrap(True)
        layout.addWidget(summary)
        table = QTableWidget(len(BUILTIN_RULES), 5)
        table.setObjectName("builtinRulesTable")
        table.setHorizontalHeaderLabels(
            ("数据类型", "识别标准", "默认处理", "示例", "复核方式")
        )
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        for row, rule in enumerate(BUILTIN_RULES):
            for column, value in enumerate(
                (rule.name, rule.recognition, rule.treatment, rule.example, rule.review)
            ):
                table.setItem(row, column, QTableWidgetItem(value))
        layout.addWidget(table, 1)
        return page

    def _snapshot_rule(self, rule_id: str) -> RuleDefinition:
        return _find_rule(self.library, rule_id)

    @staticmethod
    def _require_unchanged(
        library: RuleLibrary,
        expected: RuleDefinition,
    ) -> RuleDefinition:
        try:
            current = _find_rule(library, expected.id)
        except KeyError as exc:
            raise RuleConflictError("规则已被其他窗口删除，请重新读取后再操作") from exc
        if current != expected:
            raise RuleConflictError("规则已被其他窗口修改，未覆盖最新内容，请重新读取")
        return current

    def _build_tab(self, kind: RuleKind) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        summary = QLabel(
            "把明确词语替换为固定代号；适合单位简称、项目名和约定称谓。"
            if kind is RuleKind.FIXED
            else "按你给出的关键词、格式和例子，统一替换、遮住或删除。"
        )
        summary.setWordWrap(True)
        layout.addWidget(summary)

        headers = (
            ("原词", "替换为", "规则名称", "来源", "范围", "必须替换", "状态")
            if kind is RuleKind.FIXED
            else (
                "规则名称",
                "来源",
                "判断方式",
                "关键词或格式",
                "处理方式",
                "正例",
                "反例",
                "范围",
                "必须替换",
                "状态",
            )
        )
        table = QTableWidget(0, len(headers))
        table.setObjectName(
            "fixedRulesTable" if kind is RuleKind.FIXED else "standardRulesTable"
        )
        table.setHorizontalHeaderLabels(headers)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        table.itemSelectionChanged.connect(
            lambda selected_kind=kind: self._update_action_buttons(selected_kind)
        )
        table.cellDoubleClicked.connect(
            lambda _row, _column, selected_kind=kind: self._edit_clicked(selected_kind)
        )
        self._tables[kind] = table
        layout.addWidget(table, 1)

        actions = QHBoxLayout()
        add_button = QPushButton("新增")
        add_button.clicked.connect(lambda _checked=False, selected_kind=kind: self._add_clicked(selected_kind))
        edit_button = QPushButton("编辑")
        edit_button.clicked.connect(
            lambda _checked=False, selected_kind=kind: self._edit_clicked(selected_kind)
        )
        toggle_button = QPushButton("停用")
        toggle_button.clicked.connect(
            lambda _checked=False, selected_kind=kind: self._toggle_clicked(selected_kind)
        )
        delete_button = QPushButton("删除")
        delete_button.clicked.connect(
            lambda _checked=False, selected_kind=kind: self._delete_clicked(selected_kind)
        )
        self._edit_buttons[kind] = edit_button
        self._toggle_buttons[kind] = toggle_button
        self._delete_buttons[kind] = delete_button
        actions.addWidget(add_button)
        actions.addWidget(edit_button)
        actions.addWidget(toggle_button)
        actions.addWidget(delete_button)
        if kind is RuleKind.FIXED:
            actions.addSpacing(16)
            self.import_button = QPushButton("导入词库")
            self.import_button.setToolTip("支持 CSV 和 XLSX 文件")
            self.import_button.clicked.connect(self._import_clicked)
            self.template_button = QPushButton("导出空白 XLSX 模板")
            self.template_button.clicked.connect(self._template_clicked)
            actions.addWidget(self.import_button)
            actions.addWidget(self.template_button)
        actions.addStretch(1)
        layout.addLayout(actions)
        self._update_action_buttons(kind)
        return page

    def _refresh_tables(self) -> None:
        if self._library is None:
            return
        for kind, table in self._tables.items():
            rules = [rule for rule in self._library.rules if rule.kind is kind]
            table.setRowCount(len(rules))
            for row, rule in enumerate(rules):
                scope_text = "、".join(
                    _SCOPE_LABELS[scope] for scope in rule.applies_to
                )
                values: tuple[str, ...]
                if kind is RuleKind.FIXED:
                    values = (
                        rule.patterns[0],
                        rule.replacement,
                        rule.name,
                        "行业预置" if is_preset_rule(rule) else "我的规则",
                        scope_text,
                        "是" if rule.mandatory else "否",
                        "启用" if rule.enabled else "停用",
                    )
                else:
                    values = (
                        rule.name,
                        "行业预置" if is_preset_rule(rule) else "我的规则",
                        _MATCH_LABELS[rule.match_mode],
                        " | ".join(rule.patterns) if rule.patterns else "仅人工选择",
                        _ACTION_LABELS[rule.action],
                        " | ".join(rule.positive_examples),
                        " | ".join(rule.negative_examples),
                        scope_text,
                        "是" if rule.mandatory else "否",
                        "启用" if rule.enabled else "停用",
                    )
                for column, value in enumerate(values):
                    item = QTableWidgetItem(value)
                    if column == 0:
                        item.setData(Qt.ItemDataRole.UserRole, rule.id)
                    table.setItem(row, column, item)
            table.clearSelection()
            self._update_action_buttons(kind)

    def _selected_rule(self, kind: RuleKind) -> RuleDefinition | None:
        table = self._tables[kind]
        row = table.currentRow()
        if row < 0:
            return None
        item = table.item(row, 0)
        if item is None:
            return None
        rule_id = str(item.data(Qt.ItemDataRole.UserRole))
        try:
            return _find_rule(self.library, rule_id)
        except KeyError:
            return None

    def _update_action_buttons(self, kind: RuleKind) -> None:
        selected = self._selected_rule(kind) if self._library is not None else None
        enabled = selected is not None
        editable = enabled and selected is not None
        self._edit_buttons[kind].setEnabled(editable)
        self._toggle_buttons[kind].setEnabled(enabled)
        self._delete_buttons[kind].setEnabled(editable)
        self._toggle_buttons[kind].setText(
            "停用" if selected is not None and selected.enabled else "启用"
        )

    def _add_clicked(self, kind: RuleKind) -> None:
        editor = RuleEditorDialog(kind, parent=self)
        if editor.exec() != QDialog.DialogCode.Accepted:
            return
        result_rule = editor.result_rule
        if result_rule is None:
            return
        self._run_ui_action(
            "新增规则失败",
            lambda: self.add_rule(result_rule),
            "规则已新增并保存到本机。",
        )

    def _edit_clicked(self, kind: RuleKind) -> None:
        selected = self._selected_rule(kind)
        if selected is None:
            self._set_feedback("请先选择一条要编辑的规则。", error=True)
            return
        editor = RuleEditorDialog(kind, selected, self)
        if editor.exec() != QDialog.DialogCode.Accepted:
            return
        result_rule = editor.result_rule
        if result_rule is None:
            return
        self._run_ui_action(
            "编辑规则失败",
            lambda: self.update_rule(result_rule, expected_rule=selected),
            "规则修改已保存。",
        )

    def _toggle_clicked(self, kind: RuleKind) -> None:
        selected = self._selected_rule(kind)
        if selected is None:
            self._set_feedback("请先选择一条要启用或停用的规则。", error=True)
            return
        enabled = not selected.enabled
        self._run_ui_action(
            "修改规则状态失败",
            lambda: self.set_rule_enabled(
                selected.id,
                enabled,
                expected_rule=selected,
            ),
            "规则已启用。" if enabled else "规则已停用。",
        )

    def _delete_clicked(self, kind: RuleKind) -> None:
        selected = self._selected_rule(kind)
        if selected is None:
            self._set_feedback("请先选择一条要删除的规则。", error=True)
            return
        answer = QMessageBox.question(
            self,
            "确认删除",
            f"确定删除规则“{selected.name}”吗？此操作只影响本机规则库。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        self._run_ui_action(
            "删除规则失败",
            lambda: self.delete_rule(selected.id, expected_rule=selected),
            "规则已删除。",
        )

    def _restore_clicked(self) -> None:
        answer = QMessageBox.question(
            self,
            "恢复默认预置",
            "恢复已删除的系统预置吗？现有自定义规则不会被覆盖。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        self._run_ui_action(
            "恢复默认预置失败",
            lambda: self._commit(restore_default_rules),
            "缺失的默认预置已恢复。",
        )

    def _import_clicked(self) -> None:
        selected, _filter = QFileDialog.getOpenFileName(
            self,
            "导入规则",
            "",
            "规则文件 (*.csv *.xlsx)",
        )
        if not selected:
            return

        selected_path = Path(selected)
        try:
            result: object | None = self.import_file(selected_path)
        except RuleImportConflictError as exc:
            choice = QMessageBox(self)
            choice.setIcon(QMessageBox.Icon.Warning)
            choice.setWindowTitle("发现冲突规则")
            choice.setText("导入内容与本机已有规则的处理结果不同。")
            choice.setInformativeText(
                f"{exc}\n\n请选择本次导入如何处理这些冲突项。"
            )
            keep_button = choice.addButton(
                "保留现有规则",
                QMessageBox.ButtonRole.AcceptRole,
            )
            use_button = choice.addButton(
                "采用导入规则",
                QMessageBox.ButtonRole.DestructiveRole,
            )
            cancel_button = choice.addButton(
                "取消导入",
                QMessageBox.ButtonRole.RejectRole,
            )
            choice.setDefaultButton(keep_button)
            choice.exec()
            clicked = choice.clickedButton()
            if clicked is cancel_button or clicked is None:
                self._set_feedback("已取消导入，现有规则没有变化。", error=False)
                return
            policy: ConflictPolicy = (
                "use_imported" if clicked is use_button else "keep_existing"
            )
            result = self._run_ui_action(
                "导入规则失败",
                lambda: self.import_file(
                    selected_path,
                    conflict_policy=policy,
                ),
            )
        except (RuleLibraryError, KeyError, OSError, ValueError) as exc:
            message = str(exc).strip() or "规则库操作未完成"
            self._set_feedback(f"导入规则失败：{message}", error=True)
            QMessageBox.warning(self, "导入规则失败", message)
            return

        if isinstance(result, RuleImportResult):
            self._set_feedback(
                f"已新增 {result.added_count} 条；跳过完全重复 "
                f"{result.skipped_duplicate_count} 条；保留现有 "
                f"{result.kept_existing_count} 条；采用导入 "
                f"{result.replaced_existing_count} 条。",
                error=False,
            )

    def _template_clicked(self) -> None:
        selected, _filter = QFileDialog.getSaveFileName(
            self,
            "导出空白规则模板",
            "规则导入模板.xlsx",
            "Excel 工作簿 (*.xlsx)",
        )
        if not selected:
            return
        self._run_ui_action(
            "导出模板失败",
            lambda: self.export_template(Path(selected)),
            "空白规则模板已导出。",
        )

    def _reload_clicked(self) -> None:
        self._run_ui_action("读取规则库失败", self.reload, "规则库已重新读取。")

    def _run_ui_action(
        self,
        title: str,
        action: Callable[[], object],
        success_text: str = "",
    ) -> object | None:
        try:
            result = action()
        except (RuleLibraryError, KeyError, OSError, ValueError) as exc:
            message = str(exc).strip() or "规则库操作未完成"
            self._set_feedback(f"{title}：{message}", error=True)
            QMessageBox.warning(self, title, message)
            return None
        if success_text:
            self._set_feedback(success_text, error=False)
        return result

    def _set_load_error(self, error: RuleLibraryError) -> None:
        self._library = None
        self.tabs.setEnabled(False)
        self.retry_button.show()
        self._set_feedback(
            f"无法读取规则库：{error}。未创建空白规则库，也未覆盖原文件。",
            error=True,
        )

    def _set_feedback(self, message: str, *, error: bool) -> None:
        self.feedback_label.setText(message)
        self.feedback_label.setStyleSheet(
            "color: #B42318;" if error else "color: #176B62;"
        )


def _find_rule(library: RuleLibrary, rule_id: str) -> RuleDefinition:
    rule = next((item for item in library.rules if item.id == rule_id), None)
    if rule is None:
        raise KeyError("规则不存在")
    return rule


def _validate_dialog_rule(rule: RuleDefinition) -> None:
    if rule.kind is RuleKind.FIXED and (
        rule.match_mode is not MatchMode.CONTAINS
        or rule.action is not ActionKind.FIXED_REPLACEMENT
    ):
        raise RuleValidationError("固定替换必须使用“包含”和“固定代号”")
    if rule.kind is RuleKind.STANDARD and not rule.positive_examples:
        raise RuleValidationError("判断标准至少需要填写一个正确示例（正例）")


def _split_user_values(value: str) -> tuple[str, ...]:
    return tuple(
        item.strip()
        for item in re.split(r"[\r\n|｜]+", value)
        if item.strip()
    )


def _select_combo_value(combo: QComboBox, value: str) -> None:
    index = combo.findData(value)
    if index < 0:
        raise ValueError("规则包含当前界面不支持的选项")
    combo.setCurrentIndex(index)


def _new_rule_id() -> str:
    from uuid import uuid4

    return uuid4().hex


__all__ = [
    "RuleEditorDialog",
    "RuleLibraryDialog",
    "RuleTemplateError",
    "export_blank_rule_template",
]
