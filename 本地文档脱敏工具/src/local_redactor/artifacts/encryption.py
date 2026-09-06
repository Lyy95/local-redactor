from __future__ import annotations

import io
import os
import zipfile
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

import msoffcrypto  # type: ignore[import-untyped]
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]

from local_redactor.models import MappingEntry

MAPPING_HEADERS = (
    "映射编号",
    "类别",
    "原始值",
    "脱敏值",
    "处理方式",
    "出现次数",
    "位置",
    "规则来源",
    "还原说明",
)
MAPPING_WARNING = "仅限本地保管，严禁上传"


class MappingEncryptionError(RuntimeError):
    """Raised when a mapping workbook cannot be encrypted and verified."""


def create_encrypted_mapping(
    mappings: Sequence[MappingEntry],
    target_path: Path,
    password: str,
) -> None:
    """Create, Agile-encrypt and independently verify a mapping workbook."""

    target = Path(target_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    plaintext = io.BytesIO()
    try:
        _write_plaintext_mapping(mappings, plaintext)
        plaintext.seek(0)
        _encrypt_office_file(plaintext, target, password)
        verify_encrypted_mapping(target, password, expected_rows=len(mappings))
    except MappingEncryptionError:
        target.unlink(missing_ok=True)
        raise
    except Exception as exc:
        target.unlink(missing_ok=True)
        raise MappingEncryptionError("脱敏映射表加密未完成") from exc
    finally:
        plaintext.close()


def create_local_mapping(
    mappings: Sequence[MappingEntry],
    target_path: Path,
) -> None:
    """Create and independently verify a local plaintext XLSX mapping workbook."""

    target = Path(target_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_name(f".{target.name}.partial")
    partial.unlink(missing_ok=True)
    try:
        with partial.open("xb") as stream:
            _write_plaintext_mapping(mappings, stream)
        os.replace(partial, target)
        if not zipfile.is_zipfile(target):
            raise MappingEncryptionError("本地映射表不是有效的 XLSX 文件")
        workbook = load_workbook(target, read_only=True, data_only=False)
        worksheet = workbook["脱敏映射表"]
        valid = (
            worksheet["A1"].value == MAPPING_WARNING
            and worksheet.max_row == len(mappings) + 2
        )
        workbook.close()
        if not valid:
            raise MappingEncryptionError("本地映射表内容校验未通过")
    except MappingEncryptionError:
        target.unlink(missing_ok=True)
        raise
    except Exception as exc:
        target.unlink(missing_ok=True)
        raise MappingEncryptionError("本地映射表生成未完成") from exc
    finally:
        partial.unlink(missing_ok=True)


def verify_encrypted_mapping(
    encrypted_path: Path,
    password: str,
    *,
    expected_rows: int | None = None,
) -> None:
    """Verify encryption, rejection without a key, and correct-key decryption."""

    path = Path(encrypted_path)
    if not path.is_file() or path.stat().st_size == 0:
        raise MappingEncryptionError("加密映射表不存在或为空")
    if zipfile.is_zipfile(path):
        raise MappingEncryptionError("映射表仍是可直接解包的明文 OOXML")

    try:
        with path.open("rb") as stream:
            office = msoffcrypto.OfficeFile(stream)
            if not office.is_encrypted():
                raise MappingEncryptionError("映射表没有启用 Office 文件加密")
            if getattr(office, "type", None) != "agile":
                raise MappingEncryptionError("映射表没有使用 Office Agile Encryption")
    except MappingEncryptionError:
        raise
    except Exception as exc:
        raise MappingEncryptionError("映射文件无法识别为 Office 加密文档") from exc

    decrypted_without_key = io.BytesIO()
    try:
        with path.open("rb") as stream:
            no_key_office = msoffcrypto.OfficeFile(stream)
            no_key_office.decrypt(decrypted_without_key)
    except Exception:
        pass
    else:
        raise MappingEncryptionError("映射表在未提供密码时仍可解密")

    wrong_password = "__local_redactor_invalid_password__"
    if wrong_password == password:
        wrong_password += "_2"
    try:
        with path.open("rb") as stream:
            wrong_key_office = msoffcrypto.OfficeFile(stream)
            wrong_key_office.load_key(password=wrong_password, verify_password=True)
            wrong_key_office.decrypt(io.BytesIO())
    except Exception:
        pass
    else:
        raise MappingEncryptionError("映射表使用错误密码时仍可解密")

    decrypted = io.BytesIO()
    try:
        with path.open("rb") as stream:
            office = msoffcrypto.OfficeFile(stream)
            office.load_key(password=password, verify_password=True)
            office.decrypt(decrypted)
        decrypted.seek(0)
        if not zipfile.is_zipfile(decrypted):
            raise MappingEncryptionError("正确密码解密后不是有效的 XLSX 文件")
        decrypted.seek(0)
        workbook = load_workbook(decrypted, read_only=True, data_only=False)
        worksheet = workbook["脱敏映射表"]
        if worksheet["A1"].value != MAPPING_WARNING:
            workbook.close()
            raise MappingEncryptionError("映射表缺少本地保管警示")
        if expected_rows is not None and worksheet.max_row != expected_rows + 2:
            workbook.close()
            raise MappingEncryptionError("映射表行数与已确认映射不一致")
        workbook.close()
    except MappingEncryptionError:
        raise
    except Exception as exc:
        raise MappingEncryptionError("正确密码无法验证映射表") from exc
    finally:
        decrypted.close()
        decrypted_without_key.close()


def _write_plaintext_mapping(
    mappings: Sequence[MappingEntry],
    stream: Any,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "脱敏映射表"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    warning = worksheet.cell(row=1, column=1)
    _set_literal_text(warning, MAPPING_WARNING)
    warning.font = Font(bold=True, color="FFFFFF", size=14)
    warning.fill = PatternFill("solid", fgColor="B42318")
    warning.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28

    for column, header in enumerate(MAPPING_HEADERS, start=1):
        cell = worksheet.cell(row=2, column=column)
        _set_literal_text(cell, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B62")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row, mapping in enumerate(mappings, start=3):
        values: tuple[str | int, ...] = (
            mapping.mapping_id,
            mapping.category.value,
            mapping.original,
            mapping.replacement,
            mapping.method.value,
            mapping.occurrence_count,
            "\n".join(mapping.locations),
            mapping.rule_source,
            mapping.restore_note,
        )
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=column)
            if isinstance(value, str):
                _set_literal_text(cell, value)
            else:
                cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:I{max(2, len(mappings) + 2)}"
    widths = (18, 16, 28, 28, 18, 12, 42, 28, 36)
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            worksheet.cell(row=2, column=column).column_letter
        ].width = width

    properties = workbook.properties
    properties.creator = ""
    properties.lastModifiedBy = ""
    properties.title = "本地脱敏映射表"
    properties.subject = ""
    properties.description = ""
    properties.keywords = ""
    properties.category = ""
    properties.contentStatus = ""
    properties.identifier = ""
    properties.language = ""
    properties.version = ""
    properties.created = datetime(2000, 1, 1)
    properties.modified = datetime(2000, 1, 1)
    properties.lastPrinted = None
    properties.revision = "1"
    workbook.save(stream)
    workbook.close()


def _set_literal_text(cell: Any, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def _encrypt_office_file(
    plaintext: io.BytesIO,
    target: Path,
    password: str,
) -> None:
    partial = target.with_name(f".{target.name}.partial")
    partial.unlink(missing_ok=True)
    try:
        plaintext.seek(0)
        with partial.open("xb") as destination:
            office = msoffcrypto.OfficeFile(plaintext)
            office.encrypt(password, destination)
        os.replace(partial, target)
    except Exception as exc:
        partial.unlink(missing_ok=True)
        raise MappingEncryptionError("Office Agile 映射表加密失败") from exc
