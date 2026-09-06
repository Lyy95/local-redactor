from __future__ import annotations

from io import BytesIO
from pathlib import Path

import cv2
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = PROJECT_ROOT / "fixtures"
REGULAR_FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
BOLD_FONT = Path(r"C:\Windows\Fonts\msyhbd.ttc")


def fictional_id(prefix17: str) -> str:
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    check_codes = "10X98765432"
    total = sum(int(digit) * weight for digit, weight in zip(prefix17, weights, strict=True))
    return prefix17 + check_codes[total % 11]


def demo_image() -> bytes:
    canvas = Image.new("RGB", (1000, 520), "white")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype(str(REGULAR_FONT), 32)
    small = ImageFont.truetype(str(REGULAR_FONT), 24)
    bold = ImageFont.truetype(str(BOLD_FONT), 38)
    draw.text((42, 30), "虚构项目现场确认单", font=bold, fill="#1f2937")
    draw.text((42, 105), "联系人：林青河  手机号：13800000000", font=regular, fill="#1f2937")
    draw.text(
        (42, 160),
        "系统：晨星协同系统  内网IP：10.23.8.17",
        font=regular,
        fill="#1f2937",
    )
    draw.text((42, 216), "全部内容仅用于脱敏工具测试", font=small, fill="#5f6b76")
    draw.ellipse((640, 55, 880, 295), outline="#c81e1e", width=14)
    draw.ellipse((670, 85, 850, 265), outline="#c81e1e", width=5)
    draw.text((706, 145), "测试印章", font=regular, fill="#c81e1e")
    draw.line((80, 365, 180, 330, 260, 390, 350, 322), fill="#111827", width=7)
    draw.text((82, 405), "虚构签名候选", font=small, fill="#374151")

    encoder = cv2.QRCodeEncoder_create()
    qr = encoder.encode("https://demo.example.invalid/case/CASE-2026-0001?user=user01")
    qr = cv2.resize(qr, (190, 190), interpolation=cv2.INTER_NEAREST)
    qr_rgb = Image.fromarray(np.repeat(qr[:, :, None], 3, axis=2))
    canvas.paste(qr_rgb, (760, 315))
    output = BytesIO()
    canvas.save(output, format="PNG", compress_level=9)
    return output.getvalue()


def add_external_hyperlink(paragraph, text: str, target: str) -> None:  # type: ignore[no-untyped-def]
    relationship_id = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend((run_properties, text_element))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_docx(path: Path, image_bytes: bytes) -> None:
    document = Document()
    document.core_properties.author = "虚构作者甲"
    document.core_properties.last_modified_by = "虚构审核人乙"
    document.core_properties.title = "虚构晨星项目方案"
    heading = document.add_heading("晨星协同系统试点方案（虚构测试资料）", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    paragraph.add_run(
        "2026年7月18日，星澜数据技术有限公司项目负责人林青河在云港市云河区，"
        "与数据治理处沟通晨星协同系统试点，预算137.6万元。"
    )
    document.add_paragraph(
        f"姓名：林青河，身份证号：{fictional_id('44010019900101001')}，"
        "手机号：13800000000，邮箱：linqinghe@starlake.example.invalid。"
    )
    document.add_paragraph(
        "案件编号：CASE-2026-0048，设备编号：DEV-YG-2391，"
        "用户名：linqinghe，内网IP：10.23.8.17，公网IP：203.0.113.25，"
        "域名：api.starlake.example.invalid，车牌：粤B12345。"
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("部门", "项目", "金额", "时间")
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        cell.paragraphs[0].runs[0].bold = True
    values = ("数据治理处", "晨星协同项目", "137.6万元", "2026年7月18日 14:36")
    row = table.add_row()
    for index, value in enumerate(values):
        row.cells[index].text = value

    hidden_paragraph = document.add_paragraph("可见说明：以下运行标记为隐藏文字。")
    hidden_run = hidden_paragraph.add_run("隐藏代号：BLUE-SKY-2026")
    hidden_properties = hidden_run._r.get_or_add_rPr()
    hidden_properties.append(OxmlElement("w:vanish"))

    link_paragraph = document.add_paragraph("虚构内部链接：")
    add_external_hyperlink(
        link_paragraph,
        "demo.example.invalid",
        "https://demo.example.invalid/internal/project",
    )
    image_stream = BytesIO(image_bytes)
    document.add_picture(image_stream, width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    section = document.sections[0]
    section.header.paragraphs[0].text = "内部资料 - 虚构测试水印代号 BLUE-SKY"
    section.footer.paragraphs[0].text = "仅用于本地文档脱敏工具验收"
    if hasattr(document, "add_comment"):
        comment_target = document.paragraphs[1].runs[0]
        document.add_comment(
            runs=[comment_target],
            text="虚构批注：提交前核对项目单位和负责人信息。",
            author="虚构批注人",
            initials="TEST",
        )
    for style in document.styles:
        if style.name in {"Normal", "Body Text"}:
            style.font.name = "Microsoft YaHei"
            style.font.size = Pt(10.5)
    document.save(path)


def make_xlsx(path: Path, image_bytes: bytes) -> None:
    workbook = Workbook()
    workbook.properties.creator = "虚构作者甲"
    workbook.properties.lastModifiedBy = "虚构审核人乙"
    workbook.properties.title = "虚构晨星项目台账"
    sheet = workbook.active
    sheet.title = "晨星项目台账"
    headers = (
        "姓名",
        "单位",
        "部门",
        "手机号",
        "时间",
        "金额",
        "案件编号",
        "内网IP",
    )
    sheet.append(headers)
    sheet.append(
        (
            "林青河",
            "星澜数据技术有限公司",
            "数据治理处",
            "13800000000",
            "2026年7月18日 14:36",
            1_376_000,
            "CASE-2026-0048",
            "10.23.8.17",
        )
    )
    sheet.append(
        (
            "林青河",
            "星澜数据技术有限公司",
            "数据治理处",
            "13800000000",
            "2026年7月19日 09:20",
            860_000,
            "CASE-2026-0051",
            "10.23.8.17",
        )
    )
    sheet["I1"] = "设备编号"
    sheet["I2"] = "DEV-YG-2391"
    sheet["J1"] = "用户名"
    sheet["J2"] = "linqinghe"
    sheet["K1"] = "系统"
    sheet["K2"] = "晨星协同系统"
    sheet["L1"] = "合计"
    sheet["L2"] = "=SUM(F2:F3)"
    sheet["A2"].comment = Comment(
        "虚构批注：负责人信息需要复核。",
        "虚构批注人",
    )
    sheet["K2"].hyperlink = "https://demo.example.invalid/internal/system"
    sheet["K2"].style = "Hyperlink"
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["J"].hidden = True
    sheet.oddHeader.center.text = "内部台账 - BLUE-SKY"
    sheet.oddFooter.center.text = "仅用于脱敏工具验收"
    sheet.freeze_panes = "A2"
    sheet.merge_cells("A5:D5")
    sheet["A5"] = "图片包含虚构文字、印章、签名和二维码候选"

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B62")
        cell.alignment = Alignment(horizontal="center")
    for column in "ABCDEFGHIJKL":
        sheet.column_dimensions[column].width = 18
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["E"].width = 24
    image_stream = BytesIO(image_bytes)
    excel_image = ExcelImage(image_stream)
    excel_image.width = 620
    excel_image.height = 322
    sheet.add_image(excel_image, "A7")

    hidden = workbook.create_sheet("隐藏配置")
    hidden.sheet_state = "veryHidden"
    hidden.append(("配置项", "配置值"))
    hidden.append(("内部域名", "api.starlake.example.invalid"))
    hidden.append(("平台账号", "ACCT-DEMO-0008"))
    workbook.save(path)
    workbook.close()


def main() -> None:
    if not REGULAR_FONT.is_file() or not BOLD_FONT.is_file():
        raise RuntimeError("缺少生成虚构图片所需的中文字体")
    FIXTURES.mkdir(parents=True, exist_ok=True)
    image = demo_image()
    make_docx(FIXTURES / "虚构项目方案.docx", image)
    make_xlsx(FIXTURES / "虚构项目台账.xlsx", image)
    print(FIXTURES)


if __name__ == "__main__":
    main()
