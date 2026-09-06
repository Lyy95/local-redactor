from __future__ import annotations

import hashlib
import io
from datetime import datetime
from pathlib import Path

import msoffcrypto
import pytest
from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from msoffcrypto.exceptions import DecryptionError, InvalidKeyError
from openpyxl import Workbook, load_workbook

from local_redactor.artifacts import (
    ArtifactExportError,
    ArtifactValidationError,
    ArtifactWriter,
)
from local_redactor.models import (
    Category,
    DocumentKind,
    DocumentModel,
    Finding,
    FindingStatus,
    MappingEntry,
    Modality,
    ObjectDisposition,
    PackageItem,
    SourceLocation,
    TextBlock,
    TransformMethod,
)
from local_redactor.office import DocxAdapter, XlsxAdapter

PASSWORD = "Example#Pass123"
ORIGINAL = "虚构人员林舟"
REPLACEMENT = "人员甲"
FIXED_TIME = datetime(2026, 7, 29, 15, 30, 0)


class FixtureAdapter:
    def __init__(
        self,
        kind: DocumentKind,
        *,
        leak_original: bool = False,
        fail: bool = False,
        external_relationship: bool = False,
    ) -> None:
        self.kind = kind
        self.leak_original = leak_original
        self.fail = fail
        self.external_relationship = external_relationship

    def export(
        self,
        document: DocumentModel,
        findings: list[Finding],
        target_path: Path,
    ) -> None:
        target_path.parent.mkdir(parents=True, exist_ok=True)
        if self.fail:
            target_path.write_bytes(b"incomplete")
            raise RuntimeError("fixture failure")
        text = ORIGINAL if self.leak_original else findings[0].replacement
        if self.kind is DocumentKind.DOCX:
            document_file = WordDocument()
            paragraph = document_file.add_paragraph(text)
            if self.external_relationship:
                _add_external_hyperlink(paragraph)
            document_file.save(target_path)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "分析清单"
            worksheet["A1"] = text
            workbook.save(target_path)
            workbook.close()

    def scan(self, path: Path) -> DocumentModel:
        blocks: list[TextBlock] = []
        if self.kind is DocumentKind.DOCX:
            document_file = WordDocument(path)
            texts = [paragraph.text for paragraph in document_file.paragraphs]
        else:
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            texts = [
                str(cell.value)
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            workbook.close()
        for index, text in enumerate(texts, start=1):
            blocks.append(
                TextBlock(
                    text=text,
                    location=SourceLocation(
                        part="fixture",
                        display=f"虚构位置 {index}",
                    ),
                )
            )
        return DocumentModel(
            kind=self.kind,
            source_path=path,
            source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            display_name=path.name,
            blocks=blocks,
        )


def _add_external_hyperlink(paragraph: object) -> None:
    relationship_id = paragraph.part.relate_to(  # type: ignore[attr-defined]
        "https://example.invalid/should-not-survive",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "外部链接"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # type: ignore[attr-defined]


def _fixture_input(
    tmp_path: Path,
    kind: DocumentKind,
) -> tuple[DocumentModel, list[Finding], list[MappingEntry]]:
    suffix = ".docx" if kind is DocumentKind.DOCX else ".xlsx"
    source = tmp_path / f"虚构原文件{suffix}"
    source.write_bytes(b"fictional source placeholder")
    location = SourceLocation(part="fixture", display="虚构位置 1")
    document = DocumentModel(
        kind=kind,
        source_path=source,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        display_name=source.name,
    )
    finding = Finding(
        category=Category.NAME,
        modality=Modality.TEXT,
        original=ORIGINAL,
        replacement=REPLACEMENT,
        locations=[location],
        detector="fixture",
        confidence=1.0,
        suggested_method=TransformMethod.ALIAS,
        status=FindingStatus.TRANSFORM,
    )
    mapping = MappingEntry(
        mapping_id="PERSON-001",
        category=Category.NAME,
        original=ORIGINAL,
        replacement=REPLACEMENT,
        method=TransformMethod.ALIAS,
        occurrence_count=1,
        locations=["虚构位置 1"],
        restore_note="按映射编号人工核对",
        rule_source="虚构固定规则",
    )
    return document, [finding], [mapping]


def _real_fixture_input(
    tmp_path: Path,
    kind: DocumentKind,
) -> tuple[DocumentModel, list[Finding], list[MappingEntry]]:
    if kind is DocumentKind.DOCX:
        source = tmp_path / "真实适配器虚构输入.docx"
        document_file = WordDocument()
        document_file.add_paragraph(f"{ORIGINAL}负责虚构项目。")
        document_file.save(source)
        adapter = DocxAdapter()
    else:
        source = tmp_path / "真实适配器虚构输入.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "虚构清单"
        worksheet["A1"] = ORIGINAL
        workbook.save(source)
        workbook.close()
        adapter = XlsxAdapter()
    model = adapter.scan(source)
    for item in model.inventory.parts:
        if item.requires_decision:
            item.disposition = ObjectDisposition.REMOVE
    locations: list[SourceLocation] = []
    for block in model.blocks:
        start = block.text.find(ORIGINAL)
        if start < 0:
            continue
        locations.append(
            SourceLocation(
                part=block.location.part,
                display=block.location.display,
                block_id=block.id,
                sheet=block.location.sheet,
                cell=block.location.cell,
                start=start,
                end=start + len(ORIGINAL),
            )
        )
    assert locations
    finding = Finding(
        category=Category.NAME,
        modality=Modality.TEXT if kind is DocumentKind.DOCX else Modality.CELL,
        original=ORIGINAL,
        replacement=REPLACEMENT,
        locations=locations,
        detector="fixture",
        confidence=1.0,
        suggested_method=TransformMethod.ALIAS,
        status=FindingStatus.TRANSFORM,
    )
    mapping = MappingEntry(
        mapping_id="PERSON-001",
        category=Category.NAME,
        original=ORIGINAL,
        replacement=REPLACEMENT,
        method=TransformMethod.ALIAS,
        occurrence_count=len(locations),
        locations=[location.display for location in locations],
    )
    return model, [finding], [mapping]


def _decrypt_mapping(path: Path, password: str) -> io.BytesIO:
    output = io.BytesIO()
    with path.open("rb") as stream:
        office = msoffcrypto.OfficeFile(stream)
        office.load_key(password=password, verify_password=True)
        office.decrypt(output)
    output.seek(0)
    return output


@pytest.mark.parametrize("kind", [DocumentKind.DOCX, DocumentKind.XLSX])
def test_artifact_writer_creates_isolated_verified_delivery(
    tmp_path: Path,
    kind: DocumentKind,
) -> None:
    document, findings, mappings = _fixture_input(tmp_path, kind)
    document.warnings.append("IMAGE_OCR_FAILED: 已按整图安全策略处理")
    document.inventory.parts.append(
        PackageItem(
            part="fixture/hidden-object",
            kind="embedded_object",
            requires_decision=True,
            disposition=ObjectDisposition.REMOVE,
        )
    )
    writer = ArtifactWriter(FixtureAdapter(kind), clock=lambda: FIXED_TIME)

    artifacts = writer.export_task(
        document,
        findings,
        mappings,
        tmp_path / "结果",
        PASSWORD,
    )

    suffix = ".docx" if kind is DocumentKind.DOCX else ".xlsx"
    assert artifacts.result_root.name == "脱敏结果_20260729_153000"
    assert artifacts.ai_copy == artifacts.result_root / "AI交付" / f"AI分析副本{suffix}"
    assert artifacts.encrypted_mapping == (artifacts.result_root / "本地保管" / "脱敏映射表.xlsx")
    assert artifacts.report == artifacts.result_root / "本地保管" / "脱敏检查报告.html"
    assert {path.name for path in artifacts.ai_copy.parent.iterdir()} == {f"AI分析副本{suffix}"}
    assert {path.name for path in artifacts.encrypted_mapping.parent.iterdir()} == {
        "脱敏映射表.xlsx",
        "脱敏检查报告.html",
    }
    assert not artifacts.encrypted_mapping.read_bytes().startswith(b"PK")
    assert ORIGINAL.encode("utf-8") not in artifacts.encrypted_mapping.read_bytes()
    assert PASSWORD.encode("utf-8") not in artifacts.encrypted_mapping.read_bytes()

    with artifacts.encrypted_mapping.open("rb") as stream:
        encrypted = msoffcrypto.OfficeFile(stream)
        assert encrypted.is_encrypted()
        assert encrypted.type == "agile"
    with (
        pytest.raises((DecryptionError, TypeError)),
        artifacts.encrypted_mapping.open("rb") as stream,
    ):
        no_key = msoffcrypto.OfficeFile(stream)
        no_key.decrypt(io.BytesIO())
    with (
        pytest.raises(InvalidKeyError),
        artifacts.encrypted_mapping.open("rb") as stream,
    ):
        wrong_key = msoffcrypto.OfficeFile(stream)
        wrong_key.load_key(password="Wrong#Password123", verify_password=True)

    decrypted = _decrypt_mapping(artifacts.encrypted_mapping, PASSWORD)
    workbook = load_workbook(decrypted, read_only=True, data_only=False)
    worksheet = workbook["脱敏映射表"]
    assert worksheet["A1"].value == "仅限本地保管，严禁上传"
    assert worksheet["C3"].value == ORIGINAL
    assert worksheet["D3"].value == REPLACEMENT
    assert worksheet["H3"].value == "虚构固定规则"
    assert worksheet["I3"].value == "按映射编号人工核对"
    workbook.close()
    decrypted.close()

    report = artifacts.report.read_text(encoding="utf-8")
    assert "待处理项</td><td>0" in report
    assert "技术复检状态：通过" in report
    assert "不等于文件已脱密" in report
    assert "可安全上传" in report
    assert "隐藏内容与对象清理" in report
    assert "embedded_object / remove" in report
    assert "识别与结构警告代码" in report
    assert "IMAGE_OCR_FAILED" in report
    assert "低置信度候选" in report
    assert ORIGINAL not in report
    assert REPLACEMENT not in report
    assert PASSWORD not in report
    assert "虚构固定规则" not in report
    assert set(artifacts.hashes) == {"ai_copy", "encrypted_mapping", "report"}
    assert all(len(value) == 64 for value in artifacts.hashes.values())
    assert not list((tmp_path / "结果").glob(".*.staging-*"))


def test_artifact_writer_allows_password_free_local_mapping(tmp_path: Path) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    writer = ArtifactWriter(FixtureAdapter(DocumentKind.DOCX), clock=lambda: FIXED_TIME)

    artifacts = writer.export_task(
        document,
        findings,
        mappings,
        tmp_path / "免密结果",
        "",
    )

    assert artifacts.encrypted_mapping.read_bytes().startswith(b"PK")
    workbook = load_workbook(artifacts.encrypted_mapping, read_only=True, data_only=False)
    worksheet = workbook["脱敏映射表"]
    assert worksheet["A1"].value == "仅限本地保管，严禁上传"
    assert worksheet["C3"].value == ORIGINAL
    workbook.close()


@pytest.mark.parametrize(
    ("kind", "adapter"),
    [
        (DocumentKind.DOCX, DocxAdapter()),
        (DocumentKind.XLSX, XlsxAdapter()),
    ],
)
def test_artifact_writer_integrates_with_real_office_adapter(
    tmp_path: Path,
    kind: DocumentKind,
    adapter: DocxAdapter | XlsxAdapter,
) -> None:
    document, findings, mappings = _real_fixture_input(tmp_path, kind)
    writer = ArtifactWriter(adapter, clock=lambda: FIXED_TIME)

    artifacts = writer.export_task(
        document,
        findings,
        mappings,
        tmp_path / "真实适配器结果",
        PASSWORD,
    )

    assert artifacts.ai_copy.is_file()
    assert ORIGINAL.encode("utf-8") not in artifacts.ai_copy.read_bytes()
    assert artifacts.encrypted_mapping.is_file()
    assert artifacts.report.is_file()


def test_artifact_writer_cleans_staging_when_adapter_fails(tmp_path: Path) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    result_root = tmp_path / "结果"
    writer = ArtifactWriter(
        FixtureAdapter(DocumentKind.DOCX, fail=True),
        clock=lambda: FIXED_TIME,
    )

    with pytest.raises(ArtifactExportError):
        writer.export_task(document, findings, mappings, result_root, PASSWORD)

    assert not list(result_root.glob("脱敏结果_*"))
    assert not list(result_root.glob(".*.staging-*"))


def test_artifact_writer_blocks_mandatory_rule_kept_as_original(tmp_path: Path) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    findings[0].status = FindingStatus.KEEP_FALSE_POSITIVE
    findings[0].replacement = findings[0].original
    findings[0].metadata["rule_mandatory"] = True
    writer = ArtifactWriter(FixtureAdapter(DocumentKind.DOCX), clock=lambda: FIXED_TIME)

    with pytest.raises(ArtifactExportError, match="强制规则"):
        writer.export_task(
            document,
            findings,
            mappings,
            tmp_path / "强制规则结果",
            PASSWORD,
        )


@pytest.mark.parametrize(
    ("original", "replacement"),
    [
        ("SecretProject", "secretproject"),
        ("AB", "ABX"),
    ],
)
def test_artifact_writer_blocks_normalized_original_inside_replacement(
    tmp_path: Path,
    original: str,
    replacement: str,
) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    findings[0].original = original
    findings[0].replacement = replacement
    findings[0].metadata["rule_mandatory"] = True
    writer = ArtifactWriter(FixtureAdapter(DocumentKind.DOCX), clock=lambda: FIXED_TIME)

    with pytest.raises(ArtifactExportError, match="完整原文"):
        writer.export_task(
            document,
            findings,
            mappings,
            tmp_path / f"残留原词-{original}",
            PASSWORD,
        )


def test_artifact_writer_blocks_original_value_leak_and_cleans_output(
    tmp_path: Path,
) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    result_root = tmp_path / "结果"
    writer = ArtifactWriter(
        FixtureAdapter(DocumentKind.DOCX, leak_original=True),
        clock=lambda: FIXED_TIME,
    )

    with pytest.raises(ArtifactValidationError):
        writer.export_task(document, findings, mappings, result_root, PASSWORD)

    assert not list(result_root.glob("脱敏结果_*"))
    assert not list(result_root.glob(".*.staging-*"))


def test_artifact_writer_blocks_external_relationship(tmp_path: Path) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.DOCX)
    result_root = tmp_path / "结果"
    writer = ArtifactWriter(
        FixtureAdapter(DocumentKind.DOCX, external_relationship=True),
        clock=lambda: FIXED_TIME,
    )

    with pytest.raises(ArtifactValidationError):
        writer.export_task(document, findings, mappings, result_root, PASSWORD)

    assert not list(result_root.glob("脱敏结果_*"))
    assert not list(result_root.glob(".*.staging-*"))


def test_artifact_writer_rejects_pending_review_before_creating_output(
    tmp_path: Path,
) -> None:
    document, findings, mappings = _fixture_input(tmp_path, DocumentKind.XLSX)
    findings[0].status = FindingStatus.PENDING
    result_root = tmp_path / "结果"
    writer = ArtifactWriter(FixtureAdapter(DocumentKind.XLSX), clock=lambda: FIXED_TIME)

    with pytest.raises(ArtifactExportError):
        writer.export_task(document, findings, mappings, result_root, PASSWORD)

    assert not result_root.exists()
