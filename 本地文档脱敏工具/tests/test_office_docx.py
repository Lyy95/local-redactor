from __future__ import annotations

import hashlib
import io
import struct
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image

from local_redactor.models import (
    Category,
    Finding,
    FindingStatus,
    ImageDisposition,
    ImageObject,
    ImageRegion,
    Modality,
    ObjectDisposition,
    SourceLocation,
    TransformMethod,
)
from local_redactor.office import DocxAdapter, ExportSafetyError, OoxmlSecurityError
from local_redactor.office.common import prepare_image_for_export, sanitize_image

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"


def _png_bytes() -> bytes:
    image = Image.new("RGB", (90, 48), color=(35, 92, 140))
    output = io.BytesIO()
    image.save(output, format="PNG", compress_level=6)
    return output.getvalue()


def test_docx_rebuild_preserves_mixed_run_format_after_shorter_replacement(
    tmp_path: Path,
) -> None:
    source = tmp_path / "mixed-format.docx"
    document = Document()
    paragraph = document.add_paragraph()
    prefix = paragraph.add_run("前缀")
    prefix.font.name = "Arial"
    prefix.font.size = Pt(10)
    sensitive = paragraph.add_run("海南省")
    sensitive.bold = True
    sensitive.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    suffix = paragraph.add_run("后缀")
    suffix.italic = True
    suffix.font.size = Pt(14)
    paragraph.paragraph_format.space_before = Pt(12)
    document.save(source)

    adapter = DocxAdapter()
    model = adapter.scan(source)
    _resolve_inventory(model)
    finding = _finding(model, "海南省", "HN", Category.LOCATION)
    target = tmp_path / "rebuilt.docx"
    adapter.export(model, [finding], target)

    rebuilt = Document(target)
    result = next(item for item in rebuilt.paragraphs if item.text == "前缀HN后缀")
    assert [run.text for run in result.runs] == ["前缀", "HN", "后缀"]
    assert result.runs[1].bold is True
    assert result.runs[1].font.color.rgb == RGBColor(0xC0, 0x00, 0x00)
    assert result.runs[2].italic is True
    assert result.runs[2].font.size == Pt(14)
    assert result.paragraph_format.space_before == Pt(12)


def _add_hyperlink(paragraph: object, text: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(  # type: ignore[attr-defined]
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # type: ignore[attr-defined]


def _build_docx_fixture(path: Path) -> bytes:
    document = Document()
    title = document.add_paragraph("蓝穹演练项目")
    title.style = "Title"
    paragraph = document.add_paragraph("虚构人员林舟负责蓝穹演练项目。")
    _add_hyperlink(
        paragraph,
        "公开说明",
        "https://example.invalid/fictional-guide",
    )
    hidden_run = document.add_paragraph().add_run("隐藏标识-霜桥")
    hidden_run.font.hidden = True
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "单位"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "虚构机构青屿公司"
    table.cell(1, 1).text = "林舟"
    header = document.sections[0].header
    header.paragraphs[0].text = "内部页眉代号-蓝穹"
    footer = document.sections[0].footer
    footer.paragraphs[0].text = "内部页脚代号-霜桥"
    image_bytes = _png_bytes()
    picture = document.add_picture(io.BytesIO(image_bytes))
    picture._inline.docPr.set("descr", "虚构蓝色示意图")
    document.save(path)
    return image_bytes


def _inject_revision_and_textbox(source: Path, target: Path) -> None:
    with (
        zipfile.ZipFile(source, "r") as incoming,
        zipfile.ZipFile(
            target,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            data = incoming.read(info)
            if info.filename == "word/document.xml":
                root = etree.fromstring(data)
                body = root.find(f"{{{W_NS}}}body")
                assert body is not None
                first_paragraph = body.find(f"{{{W_NS}}}p")
                assert first_paragraph is not None
                deleted = etree.SubElement(first_paragraph, f"{{{W_NS}}}del")
                deleted_run = etree.SubElement(deleted, f"{{{W_NS}}}r")
                deleted_text = etree.SubElement(deleted_run, f"{{{W_NS}}}delText")
                deleted_text.text = "修订删除标识-云舟"

                textbox_paragraph = etree.Element(f"{{{W_NS}}}p")
                run = etree.SubElement(textbox_paragraph, f"{{{W_NS}}}r")
                pict = etree.SubElement(run, f"{{{W_NS}}}pict")
                shape = etree.SubElement(pict, f"{{{V_NS}}}shape")
                textbox = etree.SubElement(shape, f"{{{V_NS}}}textbox")
                textbox_content = etree.SubElement(textbox, f"{{{W_NS}}}txbxContent")
                inner_paragraph = etree.SubElement(
                    textbox_content,
                    f"{{{W_NS}}}p",
                )
                inner_run = etree.SubElement(inner_paragraph, f"{{{W_NS}}}r")
                inner_text = etree.SubElement(inner_run, f"{{{W_NS}}}t")
                inner_text.text = "文本框标识-星槎"
                section_properties = body.find(f"{{{W_NS}}}sectPr")
                if section_properties is None:
                    body.append(textbox_paragraph)
                else:
                    section_properties.addprevious(textbox_paragraph)
                data = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            outgoing.writestr(info.filename, data)


def _oversized_png_header(content: bytes) -> bytes:
    assert content[:8] == b"\x89PNG\r\n\x1a\n"
    assert content[12:16] == b"IHDR"
    header = bytearray(content[16:29])
    header[:8] = struct.pack(">II", 9_000, 9_000)
    checksum = struct.pack(">I", zlib.crc32(b"IHDR" + header) & 0xFFFFFFFF)
    return content[:16] + bytes(header) + checksum + content[33:]


def _finding(
    document: object,
    original: str,
    replacement: str,
    category: Category,
) -> Finding:
    locations: list[SourceLocation] = []
    for block in document.blocks:  # type: ignore[attr-defined]
        offset = 0
        while True:
            start = block.text.find(original, offset)
            if start < 0:
                break
            locations.append(
                SourceLocation(
                    part=block.location.part,
                    display=block.location.display,
                    block_id=block.id,
                    start=start,
                    end=start + len(original),
                )
            )
            offset = start + len(original)
    assert locations
    return Finding(
        category=category,
        modality=Modality.TEXT,
        original=original,
        replacement=replacement,
        locations=locations,
        detector="fictional-test",
        confidence=1.0,
        suggested_method=TransformMethod.ALIAS,
        status=FindingStatus.TRANSFORM,
    )


def _resolve_inventory(document: object) -> None:
    for item in document.inventory.parts:  # type: ignore[attr-defined]
        if not item.requires_decision:
            continue
        item.disposition = (
            ObjectDisposition.VISIBLE_NOTE if item.kind == "header" else ObjectDisposition.REMOVE
        )


def _all_docx_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.extend(paragraph.text for paragraph in cell.paragraphs)
    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(values)


def test_docx_scan_covers_body_table_header_hidden_revision_textbox_link_and_image(
    tmp_path: Path,
) -> None:
    base = tmp_path / "基础.docx"
    source = tmp_path / "多部件虚构文档.docx"
    _build_docx_fixture(base)
    _inject_revision_and_textbox(base, source)

    model = DocxAdapter().scan(source)
    all_text = "\n".join(block.text for block in model.blocks)
    block_kinds = {block.block_kind for block in model.blocks}

    assert "虚构人员林舟" in all_text
    assert "虚构机构青屿公司" in all_text
    assert "内部页眉代号-蓝穹" in all_text
    assert "隐藏标识-霜桥" in all_text
    assert "修订删除标识-云舟" in all_text
    assert "文本框标识-星槎" in all_text
    assert "https://example.invalid/fictional-guide" in all_text
    assert "hidden_text" in block_kinds
    assert "revision_deleted" in block_kinds
    assert "textbox" in block_kinds
    assert "external_relationship" in block_kinds
    assert len(model.images) == 1
    assert model.images[0].alt_text == "Picture 1 虚构蓝色示意图"
    assert model.inventory.unresolved_items


def test_docx_export_requires_every_image_and_object_decision(tmp_path: Path) -> None:
    source = tmp_path / "待复核.docx"
    _build_docx_fixture(source)
    model = DocxAdapter().scan(source)
    target = tmp_path / "不应生成.docx"

    with pytest.raises(ExportSafetyError):
        DocxAdapter().export(model, [], target)
    assert not target.exists()


def test_image_pixel_redaction_rewrites_pixels_and_metadata() -> None:
    original = _png_bytes()
    image = ImageObject(
        source_part="word/media/fictional.png",
        location=SourceLocation(
            part="word/document.xml",
            display="虚构图片",
            image_id="fictional-image",
        ),
        content=original,
        disposition=ImageDisposition.PIXEL_REDACT,
        regions=[
            ImageRegion(
                bbox=(4, 5, 20, 18),
                category=Category.PHOTO,
                replacement_label="[照片区域]",
            )
        ],
    )

    result = sanitize_image(image)

    assert result is not None
    content, media_type = result
    assert media_type == "image/png"
    assert content != original
    with Image.open(io.BytesIO(content)) as processed:
        processed.load()
        pixels = [processed.getpixel((x, y)) for y in range(5, 18) for x in range(4, 20)]
        assert (0, 0, 0) in pixels
        assert any(min(pixel) >= 200 for pixel in pixels)
        assert not processed.getexif()


def test_docx_scan_blocks_oversized_image_before_pixel_decode(tmp_path: Path) -> None:
    normal = tmp_path / "正常图片.docx"
    oversized = tmp_path / "超大图片.docx"
    _build_docx_fixture(normal)
    with (
        zipfile.ZipFile(normal, "r") as incoming,
        zipfile.ZipFile(
            oversized,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            data = incoming.read(info)
            if "/media/" in info.filename:
                data = _oversized_png_header(data)
            outgoing.writestr(info.filename, data)

    with pytest.raises(OoxmlSecurityError) as error:
        DocxAdapter().scan(oversized)
    assert error.value.code == "image-pixel-quota"


def test_export_image_guard_checks_dimensions_before_load() -> None:
    image = ImageObject(
        source_part="word/media/oversized.png",
        location=SourceLocation(
            part="word/document.xml",
            display="虚构超大图片",
            image_id="oversized-image",
        ),
        content=_oversized_png_header(_png_bytes()),
        disposition=ImageDisposition.KEEP_REENCODED,
    )

    with pytest.raises(ExportSafetyError, match="像素数量|安全解码"):
        sanitize_image(image)


def test_reviewed_image_finding_cannot_be_exported_as_keep_only() -> None:
    image = ImageObject(
        source_part="word/media/fictional.png",
        location=SourceLocation(
            part="word/document.xml",
            display="虚构图片",
            image_id="fictional-image",
        ),
        content=_png_bytes(),
        disposition=ImageDisposition.KEEP_REENCODED,
    )
    finding = Finding(
        category=Category.NAME,
        modality=Modality.IMAGE,
        original="林舟",
        replacement="人员甲",
        locations=[
            SourceLocation(
                part="word/document.xml",
                display="虚构图片文字",
                image_id=image.id,
                bbox=(4, 5, 40, 22),
            )
        ],
        detector="fictional-test",
        confidence=1.0,
        suggested_method=TransformMethod.PIXEL_REDACT,
        status=FindingStatus.TRANSFORM,
    )

    with pytest.raises(ExportSafetyError, match="不能仅重新编码"):
        prepare_image_for_export(image, [finding])


def test_whole_image_region_keeps_all_reviewed_aliases_in_visible_legend() -> None:
    image = ImageObject(
        source_part="word/media/fictional.png",
        location=SourceLocation(
            part="word/document.xml",
            display="虚构图片",
            image_id="fictional-image",
        ),
        content=_png_bytes(),
        disposition=ImageDisposition.PIXEL_REDACT,
        regions=[
            ImageRegion(
                bbox=(0, 0, 90, 48),
                category=Category.OTHER,
                replacement_label="[整图像素已处理]",
            )
        ],
    )
    findings = [
        Finding(
            category=Category.NAME,
            modality=Modality.IMAGE,
            original="林舟",
            replacement="人员甲",
            locations=[
                SourceLocation(
                    part="word/document.xml",
                    display="虚构图片文字",
                    image_id=image.id,
                    bbox=(4, 5, 40, 22),
                )
            ],
            detector="fictional-test",
            confidence=1.0,
            suggested_method=TransformMethod.PIXEL_REDACT,
            status=FindingStatus.TRANSFORM,
        ),
        Finding(
            category=Category.ORGANIZATION,
            modality=Modality.IMAGE,
            original="青屿公司",
            replacement="机构甲",
            locations=[
                SourceLocation(
                    part="word/document.xml",
                    display="虚构图片文字",
                    image_id=image.id,
                    bbox=(42, 5, 82, 22),
                )
            ],
            detector="fictional-test",
            confidence=1.0,
            suggested_method=TransformMethod.PIXEL_REDACT,
            status=FindingStatus.TRANSFORM,
        ),
    ]

    prepared = prepare_image_for_export(image, findings)

    assert prepared.regions[0].replacement_label == "人员甲；机构甲"


def test_docx_export_rebuilds_clean_copy_and_applies_consistent_findings(
    tmp_path: Path,
) -> None:
    source = tmp_path / "原始虚构文档.docx"
    source_image = _build_docx_fixture(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    model = DocxAdapter().scan(source)
    _resolve_inventory(model)
    image = model.images[0]
    image.disposition = ImageDisposition.PIXEL_REDACT
    region_bbox = (4, 5, 86, 40)
    image.regions.append(
        ImageRegion(
            bbox=region_bbox,
            category=Category.NAME,
            replacement_label="[不一致旧标签]",
        )
    )
    person_finding = _finding(model, "林舟", "人员甲", Category.NAME)
    person_finding.locations.append(
        SourceLocation(
            part=image.location.part,
            display=image.location.display,
            image_id=image.id,
            bbox=(10, 10, 40, 25),
        )
    )
    findings = [
        person_finding,
        _finding(model, "虚构机构青屿公司", "甲公司（技术服务商）", Category.ORGANIZATION),
        _finding(
            model,
            "内部页眉代号-蓝穹",
            "项目页眉",
            Category.HEADER_FOOTER,
        ),
    ]
    target = tmp_path / "AI分析副本.docx"

    DocxAdapter().export(model, findings, target)

    assert target.exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    text = _all_docx_text(target)
    assert "人员甲" in text
    assert text.count("人员甲") == 2
    assert "不一致旧标签" not in text
    assert "甲公司（技术服务商）" in text
    assert "项目页眉" in text
    assert "林舟" not in text
    assert "虚构机构青屿公司" not in text
    assert "隐藏标识-霜桥" not in text
    assert "内部页脚代号-霜桥" not in text

    with zipfile.ZipFile(target) as archive:
        relation_bytes = b"\n".join(
            archive.read(name) for name in archive.namelist() if name.endswith(".rels")
        )
        output_images = [archive.read(name) for name in archive.namelist() if "/media/" in name]
    assert b"example.invalid" not in relation_bytes
    assert output_images
    assert all(image != source_image for image in output_images)
