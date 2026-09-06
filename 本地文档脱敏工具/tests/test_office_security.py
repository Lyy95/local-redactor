from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from local_redactor.models import DocumentKind
from local_redactor.office import OoxmlLimits, OoxmlSecurityError, preflight_ooxml


def _minimal_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("虚构测试内容")
    document.save(path)


def _minimal_xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "虚构测试内容"
    workbook.save(path)
    workbook.close()


def _rewrite_zip(
    source: Path,
    target: Path,
    mutator: object,
) -> None:
    with (
        zipfile.ZipFile(source, "r") as incoming,
        zipfile.ZipFile(
            target,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            data = incoming.read(info)
            replacement = mutator(info.filename, data)  # type: ignore[operator]
            outgoing.writestr(info.filename, replacement)


def test_preflight_accepts_real_docx_and_xlsx_and_lists_every_part(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "虚构文档.docx"
    xlsx_path = tmp_path / "虚构表格.xlsx"
    _minimal_docx(docx_path)
    _minimal_xlsx(xlsx_path)

    docx_result = preflight_ooxml(docx_path, DocumentKind.DOCX)
    xlsx_result = preflight_ooxml(xlsx_path, DocumentKind.XLSX)

    with zipfile.ZipFile(docx_path) as archive:
        docx_names = {item.filename for item in archive.infolist() if not item.is_dir()}
    with zipfile.ZipFile(xlsx_path) as archive:
        xlsx_names = {item.filename for item in archive.infolist() if not item.is_dir()}

    assert docx_result.part_names == docx_names
    assert xlsx_result.part_names == xlsx_names
    assert "word/document.xml" in docx_result.part_names
    assert "xl/workbook.xml" in xlsx_result.part_names


@pytest.mark.parametrize("suffix", [".doc", ".docm", ".xls", ".xlsm", ".pdf"])
def test_preflight_rejects_unsupported_extensions(
    tmp_path: Path,
    suffix: str,
) -> None:
    path = tmp_path / f"伪装文件{suffix}"
    path.write_bytes(b"not-an-office-file")
    with pytest.raises(OoxmlSecurityError) as error:
        preflight_ooxml(path)
    assert error.value.code == "unsupported-extension"


def test_preflight_rejects_extension_spoofing(tmp_path: Path) -> None:
    path = tmp_path / "伪装文件.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0" + "旧格式伪装".encode())
    with pytest.raises(OoxmlSecurityError) as error:
        preflight_ooxml(path)
    assert error.value.code == "not-ooxml-zip"


def test_preflight_rejects_zip_path_traversal(tmp_path: Path) -> None:
    path = tmp_path / "危险路径.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../逃逸.xml", "<root/>")
    with pytest.raises(OoxmlSecurityError) as error:
        preflight_ooxml(path)
    assert error.value.code == "zip-path-traversal"


def test_preflight_rejects_xml_doctype_and_entity(tmp_path: Path) -> None:
    source = tmp_path / "正常.docx"
    dangerous = tmp_path / "外部实体.docx"
    _minimal_docx(source)

    def inject_doctype(name: str, data: bytes) -> bytes:
        if name == "word/document.xml":
            return (
                b'<?xml version="1.0" encoding="UTF-8"?>'
                b'<!DOCTYPE w:document [<!ENTITY leak SYSTEM "file:///fictional">]>'
                + data[data.find(b"<w:document") :]
            )
        return data

    _rewrite_zip(source, dangerous, inject_doctype)
    with pytest.raises(OoxmlSecurityError) as error:
        preflight_ooxml(dangerous)
    assert error.value.code == "xml-external-entity"


def test_preflight_enforces_compression_ratio_before_parsing(tmp_path: Path) -> None:
    source = tmp_path / "正常.docx"
    compressed = tmp_path / "压缩配额.docx"
    _minimal_docx(source)

    with (
        zipfile.ZipFile(source, "r") as incoming,
        zipfile.ZipFile(
            compressed,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            outgoing.writestr(info.filename, incoming.read(info))
        outgoing.writestr("word/media/repeat.bin", b"A" * (2 * 1024 * 1024))

    with pytest.raises(OoxmlSecurityError) as error:
        preflight_ooxml(
            compressed,
            limits=OoxmlLimits(max_compression_ratio=10.0),
        )
    assert error.value.code == "compression-ratio"


def test_preflight_lists_external_relationship_without_resolving_it(
    tmp_path: Path,
) -> None:
    source = tmp_path / "正常.xlsx"
    linked = tmp_path / "外部链接.xlsx"
    _minimal_xlsx(source)

    def add_external_link(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/_rels/sheet1.xml.rels":
            return data
        return data.replace(
            b"</Relationships>",
            (
                b'<Relationship Id="rId999" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/'
                b'2006/relationships/hyperlink" '
                b'Target="https://example.invalid/fictional" '
                b'TargetMode="External"/></Relationships>'
            ),
        )

    with zipfile.ZipFile(source, "r") as incoming:
        names = {info.filename for info in incoming.infolist()}
        with zipfile.ZipFile(linked, "w", zipfile.ZIP_DEFLATED) as outgoing:
            for info in incoming.infolist():
                outgoing.writestr(info.filename, incoming.read(info))
            if "xl/worksheets/_rels/sheet1.xml.rels" not in names:
                outgoing.writestr(
                    "xl/worksheets/_rels/sheet1.xml.rels",
                    (
                        '<?xml version="1.0" encoding="UTF-8"?>'
                        '<Relationships xmlns="http://schemas.openxmlformats.org/'
                        'package/2006/relationships">'
                        '<Relationship Id="rId999" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/'
                        '2006/relationships/hyperlink" '
                        'Target="https://example.invalid/fictional" '
                        'TargetMode="External"/>'
                        "</Relationships>"
                    ),
                )

    result = preflight_ooxml(linked, DocumentKind.XLSX)
    external = [record for record in result.relationships if record.external]
    assert len(external) == 1
    assert external[0].target == "https://example.invalid/fictional"
    assert any(item.external and item.requires_decision for item in result.inventory.parts)


def test_preflight_keeps_unknown_xml_as_an_unresolved_inventory_item(
    tmp_path: Path,
) -> None:
    source = tmp_path / "正常.docx"
    unknown = tmp_path / "未知部件.docx"
    _minimal_docx(source)
    with (
        zipfile.ZipFile(source, "r") as incoming,
        zipfile.ZipFile(
            unknown,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            outgoing.writestr(info.filename, incoming.read(info))
        outgoing.writestr(
            "word/unknown-feature.xml",
            '<?xml version="1.0" encoding="UTF-8"?><fictional/>',
        )

    result = preflight_ooxml(unknown, DocumentKind.DOCX)
    item = next(item for item in result.inventory.parts if item.part == "word/unknown-feature.xml")
    assert item.kind == "unknown_xml"
    assert item.requires_decision
    assert item in result.inventory.unresolved_items
