from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image

from local_redactor.models import (
    Category,
    Finding,
    FindingStatus,
    ImageDisposition,
    ImageRegion,
    Modality,
    ObjectDisposition,
    SourceLocation,
    TransformMethod,
)
from local_redactor.office import ExportSafetyError, XlsxAdapter

S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _png_bytes() -> bytes:
    image = Image.new("RGB", (72, 44), color=(48, 112, 77))
    output = io.BytesIO()
    image.save(output, format="PNG", compress_level=6)
    return output.getvalue()


def _build_xlsx_fixture(path: Path, *, patch_formula_cache: bool = True) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "工作清单"
    worksheet["A1"] = "虚构人员林舟"
    worksheet["A1"].font = Font(bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="356A9A")
    worksheet["A1"].comment = Comment("批注意见-星槎", "虚构作者")
    worksheet["A2"] = "林舟"
    worksheet["A2"].hyperlink = "https://example.invalid/fictional-sheet"
    worksheet["B2"] = 100
    worksheet["B3"] = 200
    worksheet["C1"] = "隐藏列标识-苇汀"
    worksheet.column_dimensions["C"].hidden = True
    worksheet["A3"] = "隐藏行标识-霜桥"
    worksheet.row_dimensions[3].hidden = True
    worksheet["D2"] = "=SUM(B2:B3)"
    worksheet.merge_cells("A5:B5")
    worksheet["A5"] = "合并标题"
    worksheet.column_dimensions["A"].width = 24
    worksheet.row_dimensions[1].height = 26
    worksheet.freeze_panes = "A2"
    worksheet.oddHeader.center.text = "内部页眉-青屿"
    worksheet.oddFooter.right.text = "内部页脚-云舟"
    image_bytes = _png_bytes()
    worksheet.add_image(OpenpyxlImage(io.BytesIO(image_bytes)), "E2")

    hidden = workbook.create_sheet("隐页-霜桥")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "虚构隐页内容"
    workbook.defined_names.add(
        DefinedName(
            "虚构范围",
            attr_text="'工作清单'!$A$1:$A$2",
        )
    )
    workbook.save(path)
    workbook.close()

    if patch_formula_cache:
        patched = path.with_name(f"{path.stem}-cached.xlsx")
        with (
            zipfile.ZipFile(path, "r") as incoming,
            zipfile.ZipFile(
                patched,
                "w",
                zipfile.ZIP_DEFLATED,
            ) as outgoing,
        ):
            for info in incoming.infolist():
                data = incoming.read(info)
                if info.filename == "xl/worksheets/sheet1.xml":
                    root = etree.fromstring(data)
                    namespace = {"s": S_NS}
                    formula_cell = root.xpath(
                        ".//s:c[@r='D2']",
                        namespaces=namespace,
                    )[0]
                    value = formula_cell.find(f"{{{S_NS}}}v")
                    assert value is not None
                    value.text = "300"
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                outgoing.writestr(info.filename, data)
        path.unlink()
        patched.replace(path)
    return image_bytes


def _finding(
    document: object,
    original: str,
    replacement: str,
    category: Category,
) -> Finding:
    locations: list[SourceLocation] = []
    for block in document.blocks:  # type: ignore[attr-defined]
        offset = 0
        while True:
            start = block.text.find(original, offset)
            if start < 0:
                break
            locations.append(
                SourceLocation(
                    part=block.location.part,
                    display=block.location.display,
                    block_id=block.id,
                    sheet=block.location.sheet,
                    cell=block.location.cell,
                    start=start,
                    end=start + len(original),
                )
            )
            offset = start + len(original)
    assert locations
    return Finding(
        category=category,
        modality=Modality.CELL,
        original=original,
        replacement=replacement,
        locations=locations,
        detector="fictional-test",
        confidence=1.0,
        suggested_method=TransformMethod.ALIAS,
        status=FindingStatus.TRANSFORM,
    )


def _resolve_inventory(document: object) -> None:
    for item in document.inventory.parts:  # type: ignore[attr-defined]
        if not item.requires_decision:
            continue
        if item.kind in {"hidden_sheet", "comment", "header_footer"}:
            item.disposition = ObjectDisposition.VISIBLE_NOTE
        else:
            item.disposition = ObjectDisposition.REMOVE


def test_xlsx_scan_covers_all_sheets_cells_formula_cache_hidden_content_and_objects(
    tmp_path: Path,
) -> None:
    source = tmp_path / "多部件虚构表格.xlsx"
    _build_xlsx_fixture(source)

    model = XlsxAdapter().scan(source)
    all_text = "\n".join(block.text for block in model.blocks)
    block_kinds = {block.block_kind for block in model.blocks}
    formula_cells = [
        cell for sheet in model.structure["sheets"] for cell in sheet["cells"] if cell["is_formula"]
    ]

    assert "虚构人员林舟" in all_text
    assert "隐页-霜桥" in all_text
    assert "虚构隐页内容" in all_text
    assert "批注意见-星槎" in all_text
    assert "内部页眉-青屿" in all_text
    assert "https://example.invalid/fictional-sheet" in all_text
    assert "虚构范围" in all_text
    assert "formula" in block_kinds
    assert "comment" in block_kinds
    assert "header_footer" in block_kinds
    assert "hyperlink_target" in block_kinds
    assert "defined_name" in block_kinds
    assert formula_cells[0]["cached_value"] == 300
    assert len(model.images) == 1
    assert any(item.kind == "hidden_sheet" for item in model.inventory.parts)
    assert any(item.kind == "hidden_row" for item in model.inventory.parts)
    assert any(item.kind == "hidden_column" for item in model.inventory.parts)
    assert model.inventory.unresolved_items


def test_xlsx_export_blocks_uncached_formula_until_reviewed(tmp_path: Path) -> None:
    source = tmp_path / "无缓存公式.xlsx"
    _build_xlsx_fixture(source, patch_formula_cache=False)
    model = XlsxAdapter().scan(source)
    for image in model.images:
        image.disposition = ImageDisposition.REMOVE
    target = tmp_path / "不应生成.xlsx"

    with pytest.raises(ExportSafetyError):
        XlsxAdapter().export(model, [], target)
    assert not target.exists()
    assert any(item.kind == "formula" and item.requires_decision for item in model.inventory.parts)


def test_xlsx_export_rebuilds_static_visible_workbook_without_active_links(
    tmp_path: Path,
) -> None:
    source = tmp_path / "原始虚构表格.xlsx"
    source_image = _build_xlsx_fixture(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    model = XlsxAdapter().scan(source)
    _resolve_inventory(model)
    image = model.images[0]
    image.disposition = ImageDisposition.PIXEL_REDACT
    region_bbox = (2, 2, 68, 40)
    image.regions.append(
        ImageRegion(
            bbox=region_bbox,
            category=Category.NAME,
            replacement_label="[不一致旧标签]",
        )
    )
    person_finding = _finding(model, "林舟", "人员甲", Category.NAME)
    person_finding.locations.append(
        SourceLocation(
            part=image.location.part,
            display=image.location.display,
            sheet=image.location.sheet,
            cell=image.location.cell,
            image_id=image.id,
            bbox=region_bbox,
        )
    )
    findings = [
        person_finding,
        _finding(model, "工作清单", "分析清单", Category.PROJECT),
        _finding(model, "内部页眉-青屿", "页眉说明", Category.HEADER_FOOTER),
    ]
    target = tmp_path / "AI分析副本.xlsx"

    XlsxAdapter().export(model, findings, target)

    assert target.exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    output = load_workbook(target, data_only=False, keep_links=False)
    assert output.sheetnames[0] == "脱敏说明"
    assert output["脱敏说明"]["A1"].value == "AI 分析副本说明"
    assert "不代表真实事实" in output["脱敏说明"]["A3"].value
    assert "分析清单" in output.sheetnames
    assert all(sheet.sheet_state == "visible" for sheet in output.worksheets)
    worksheet = output["分析清单"]
    assert worksheet["A1"].value == "虚构人员人员甲"
    assert worksheet["A2"].value == "人员甲"
    assert worksheet["D2"].value == 300
    assert worksheet["D2"].data_type != "f"
    assert worksheet["C1"].value is None
    assert worksheet["A3"].value is None
    assert worksheet["A1"].comment is None
    assert worksheet["A2"].hyperlink is None
    assert worksheet.row_dimensions[3].hidden is False
    assert worksheet.column_dimensions["C"].hidden is False
    assert "A5:B5" in {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
    assert worksheet["A1"].font.bold is True
    assert "可见附注" in output.sheetnames
    note_values = [
        cell.value
        for row in output["可见附注"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert any("[1] 人员甲" in str(value) for value in note_values)
    assert all("不一致旧标签" not in str(value) for value in note_values)
    assert all("工作清单" not in str(value) for value in note_values)
    output.close()

    with zipfile.ZipFile(target) as archive:
        relation_bytes = b"\n".join(
            archive.read(name) for name in archive.namelist() if name.endswith(".rels")
        )
        worksheet_xml = b"\n".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
        output_images = [archive.read(name) for name in archive.namelist() if "/media/" in name]
    assert b"example.invalid" not in relation_bytes
    assert b"<f" not in worksheet_xml
    assert output_images
    assert all(image != source_image for image in output_images)
