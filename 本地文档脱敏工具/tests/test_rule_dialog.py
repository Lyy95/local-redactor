from __future__ import annotations

import csv
import os
from dataclasses import replace
from pathlib import Path
from typing import cast

import pytest
from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from PySide6.QtWidgets import (  # noqa: E402
    QApplication,
    QCheckBox,
    QFileDialog,
    QLabel,
    QMessageBox,
    QTableWidget,
)

from local_redactor.rule_library import (  # noqa: E402
    ActionKind,
    MatchMode,
    RuleConflictError,
    RuleDefinition,
    RuleKind,
    RuleLibrary,
    RuleStore,
    RuleValidationError,
)
from local_redactor.ui.rule_dialog import (  # noqa: E402
    RuleEditorDialog,
    RuleLibraryDialog,
    RuleTemplateError,
)


class _TestProtector:
    def protect(self, plaintext: bytes) -> bytes:
        return b"dialog-v1:" + bytes(value ^ 0xA5 for value in plaintext)

    def unprotect(self, ciphertext: bytes) -> bytes:
        if not ciphertext.startswith(b"dialog-v1:"):
            raise ValueError("damaged")
        return bytes(value ^ 0xA5 for value in ciphertext.removeprefix(b"dialog-v1:"))


class _CountingStore(RuleStore):
    def __init__(self, path: Path) -> None:
        super().__init__(path, _TestProtector())
        self.load_count = 0
        self.save_count = 0

    def load(self):  # type: ignore[no-untyped-def]
        self.load_count += 1
        return super().load()

    def save_if_revision(
        self,
        library: RuleLibrary,
        *,
        expected_revision: int,
    ) -> None:
        self.save_count += 1
        super().save_if_revision(
            library,
            expected_revision=expected_revision,
        )


class _InterleavingStore(RuleStore):
    def __init__(self, path: Path, external_rule: RuleDefinition) -> None:
        super().__init__(path, _TestProtector())
        self.external_rule = external_rule
        self.inject_once = True

    def save_if_revision(
        self,
        library: RuleLibrary,
        *,
        expected_revision: int,
    ) -> None:
        if self.inject_once:
            self.inject_once = False
            competing = RuleStore(self.path, _TestProtector())
            latest = competing.load().add(self.external_rule)
            competing.save(latest)
        super().save_if_revision(
            library,
            expected_revision=expected_revision,
        )


@pytest.fixture(scope="session")
def dialog_qapp() -> QApplication:
    existing = QApplication.instance()
    return cast(QApplication, existing) if existing is not None else QApplication([])


def _dialog(store: RuleStore, qapp: QApplication) -> RuleLibraryDialog:
    del qapp
    dialog = RuleLibraryDialog(store)
    dialog.show()
    QApplication.processEvents()
    return dialog


def test_dialog_crud_loads_fresh_and_persists_immutable_library(
    tmp_path: Path,
    dialog_qapp: QApplication,
) -> None:
    store = _CountingStore(tmp_path / "rules.dat")
    dialog = _dialog(store, dialog_qapp)
    rule = RuleDefinition.fixed(
        "虚构单位",
        "机构甲",
        name="单位固定代号",
        match_mode=MatchMode.CONTAINS,
    )

    dialog.add_rule(rule)
    fixed_table = dialog._tables[RuleKind.FIXED]  # noqa: SLF001
    assert fixed_table.columnCount() == 7
    assert fixed_table.horizontalHeaderItem(0).text() == "原词"
    assert fixed_table.horizontalHeaderItem(1).text() == "替换为"
    assert fixed_table.item(0, 0).text() == "虚构单位"
    assert fixed_table.item(0, 1).text() == "机构甲"
    assert fixed_table.item(0, 3).text() == "我的规则"
    edited = replace(rule, name="单位统一代号", replacement="机构乙")
    dialog.update_rule(edited)
    dialog.set_rule_enabled(rule.id, False)

    persisted = store.load()
    assert persisted.rules == (replace(edited, enabled=False),)
    assert persisted.revision == 3
    assert store.load_count >= 5  # constructor + every change + assertion
    assert store.save_count == 3
    assert "虚构单位".encode() not in store.path.read_bytes()

    dialog.delete_rule(rule.id)
    assert store.load().rules == ()
    assert store.save_count == 4
    dialog.close()


def test_conflict_and_stale_edit_do_not_overwrite_latest_store(
    tmp_path: Path,
    dialog_qapp: QApplication,
) -> None:
    store = RuleStore(tmp_path / "rules.dat", _TestProtector())
    first_dialog = _dialog(store, dialog_qapp)
    original = RuleDefinition.fixed(
        "同一词语",
        "代号甲",
        name="原规则",
        match_mode=MatchMode.CONTAINS,
    )
    first_dialog.add_rule(original)
    second_dialog = _dialog(store, dialog_qapp)

    conflicting = RuleDefinition.fixed(
        "同一词语",
        "代号乙",
        name="冲突规则",
        match_mode=MatchMode.CONTAINS,
    )
    before_conflict = store.path.read_bytes()
    with pytest.raises(RuleConflictError, match="不同处理结果"):
        first_dialog.add_rule(conflicting)
    assert store.path.read_bytes() == before_conflict

    newer = replace(original, replacement="代号丙", name="其他窗口修改")
    second_dialog.update_rule(newer)
    latest_bytes = store.path.read_bytes()
    with pytest.raises(RuleConflictError, match="未覆盖最新内容"):
        first_dialog.update_rule(replace(original, name="旧窗口修改"))
    assert store.path.read_bytes() == latest_bytes
    assert store.load().rules == (newer,)
    first_dialog.close()
    second_dialog.close()


def test_interleaved_writer_is_detected_before_conditional_save(
    tmp_path: Path,
    dialog_qapp: QApplication,
) -> None:
    external = RuleDefinition.fixed(
        "外部窗口词语",
        "外部代号",
        match_mode=MatchMode.CONTAINS,
    )
    store = _InterleavingStore(tmp_path / "rules.dat", external)
    dialog = _dialog(store, dialog_qapp)
    local = RuleDefinition.fixed(
        "当前窗口词语",
        "当前代号",
        match_mode=MatchMode.CONTAINS,
    )

    with pytest.raises(RuleConflictError, match="未覆盖最新内容"):
        dialog.add_rule(local)

    assert store.load().rules == (external,)
    assert store.load().revision == 1
    dialog.close()


def test_ui_conflict_shows_explicit_chinese_error_without_saving(
    tmp_path: Path,
    dialog_qapp: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = RuleStore(tmp_path / "rules.dat", _TestProtector())
    dialog = _dialog(store, dialog_qapp)
    dialog.add_rule(
        RuleDefinition.fixed("冲突词", "代号甲", match_mode=MatchMode.CONTAINS)
    )
    before = store.path.read_bytes()
    messages: list[tuple[str, str]] = []

    def record_warning(
        _parent: object,
        title: str,
        message: str,
        *_args: object,
        **_kwargs: object,
    ) -> QMessageBox.StandardButton:
        messages.append((title, message))
        return QMessageBox.StandardButton.Ok

    monkeypatch.setattr(QMessageBox, "warning", record_warning)
    result = dialog._run_ui_action(  # noqa: SLF001
        "新增规则失败",
        lambda: dialog.add_rule(
            RuleDefinition.fixed("冲突词", "代号乙", match_mode=MatchMode.CONTAINS)
        ),
    )

    assert result is None
    assert messages
    assert messages[0][0] == "新增规则失败"
    assert "同一命中条件存在不同处理结果" in messages[0][1]
    assert "未覆盖" not in messages[0][1]
    assert "新增规则失败" in dialog.feedback_label.text()
    assert store.path.read_bytes() == before
    dialog.close()


def test_blank_template_is_safe_compatible_and_never_overwritten(
    tmp_path: Path,
    dialog_qapp: QApplication,
) -> None:
    store = RuleStore(tmp_path / "rules.dat", _TestProtector())
    dialog = _dialog(store, dialog_qapp)
    target = tmp_path / "规则导入模板.xlsx"

    assert dialog.export_template(target) == target
    workbook = load_workbook(target)
    worksheet = workbook["规则导入模板"]
    expected_headers = (
        "规则类型",
        "规则名称",
        "匹配方式",
        "关键词",
        "处理方式",
        "替换内容",
        "规则编号",
        "必须处理",
        "启用",
        "优先级",
        "正例",
        "反例",
        "适用范围",
        "区分大小写",
        "保留开头",
        "保留结尾",
        "代号位数",
        "内容类别",
    )
    assert tuple(cell.value for cell in worksheet[1]) == expected_headers
    assert worksheet.max_row == 1
    assert "填写说明" in workbook.sheetnames
    instructions = "\n".join(
        str(cell.value or "")
        for row in workbook["填写说明"].iter_rows()
        for cell in row
    )
    assert "匹配方式必须填写“包含”" in instructions
    assert "正例必填" in instructions
    worksheet.append(
        [
            "固定替换",
            "虚构单位代号",
            "包含",
            "虚构单位",
            "固定代号",
            "机构甲",
        ]
    )
    workbook.save(target)
    workbook.close()

    result = dialog.import_file(target)
    assert result.added_count == 1
    assert store.load().rules[0].replacement == "机构甲"

    existing = target.read_bytes()
    with pytest.raises(RuleTemplateError, match="已存在"):
        dialog.export_template(target)
    assert target.read_bytes() == existing
    dialog.close()


def test_import_conflict_is_atomic_and_editor_uses_user_language(
    tmp_path: Path,
    dialog_qapp: QApplication,
) -> None:
    store = RuleStore(tmp_path / "rules.dat", _TestProtector())
    dialog = _dialog(store, dialog_qapp)
    dialog.add_rule(
        RuleDefinition.fixed("虚构项目", "项目甲", match_mode=MatchMode.CONTAINS)
    )
    before = store.path.read_bytes()
    source = tmp_path / "冲突.csv"
    with source.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            ("规则类型", "规则名称", "匹配方式", "关键词", "处理方式", "替换内容")
        )
        writer.writerow(
            ("固定替换", "冲突", "包含", "虚构项目", "固定代号", "项目乙")
        )

    with pytest.raises(RuleConflictError):
        dialog.import_file(source)
    assert store.path.read_bytes() == before

    kept = dialog.import_file(source, conflict_policy="keep_existing")
    assert kept.kept_existing_count == 1
    assert store.path.read_bytes() == before

    replaced = dialog.import_file(source, conflict_policy="use_imported")
    assert replaced.replaced_existing_count == 1
    assert store.load().rules[0].replacement == "项目乙"

    editor = RuleEditorDialog(RuleKind.STANDARD)
    assert dialog.tabs.count() == 3
    assert [dialog.tabs.tabText(index) for index in range(3)] == [
        "系统内置识别 25",
        "固定替换",
        "判断标准",
    ]
    builtin_table = dialog.findChild(QTableWidget, "builtinRulesTable")
    assert builtin_table is not None
    builtin_names = {
        builtin_table.item(row, 0).text() for row in range(builtin_table.rowCount())
    }
    assert {"中文姓名", "手机号码", "身份证号", "银行卡号", "电子邮箱", "固定电话", "密码与口令", "车牌号"} <= builtin_names
    assert [
        editor.match_mode_combo.itemText(index)
        for index in range(editor.match_mode_combo.count())
    ][:3] == ["包含", "等于", "符合格式"]
    assert [
        editor.action_combo.itemText(index)
        for index in range(editor.action_combo.count())
    ] == ["固定代号", "中间星号", "顺序代号", "删除"]
    checkbox_labels = {checkbox.text() for checkbox in editor.findChildren(QCheckBox)}
    assert {
        "必须替换，不能保留原文",
        "全部内容",
        "正文",
        "图片文字",
    } <= checkbox_labels
    form_labels = {label.text() for label in editor.findChildren(QLabel)}
    assert {"正确示例（原内容）", "反例", "适用范围"} <= form_labels

    fixed_editor = RuleEditorDialog(RuleKind.FIXED)
    assert fixed_editor.match_mode_combo.count() == 1
    assert fixed_editor.match_mode_combo.currentText() == "包含"
    assert fixed_editor.match_mode_combo.isHidden()
    assert fixed_editor.action_combo.count() == 1
    assert fixed_editor.action_combo.currentText() == "固定代号"
    assert fixed_editor.action_combo.isHidden()
    fixed_editor.name_edit.setText("完整原词替换")
    fixed_editor.pattern_edit.setText("虚构原词")
    fixed_editor.replacement_edit.setText("代号甲")
    fixed_rule = fixed_editor.rule_definition()
    assert fixed_rule.match_mode is MatchMode.CONTAINS
    assert fixed_rule.action is ActionKind.FIXED_REPLACEMENT
    fixed_editor.close()

    with pytest.raises(RuleValidationError, match="包含.*固定代号"):
        dialog.add_rule(
            RuleDefinition.fixed(
                "不合规固定项",
                "X",
                match_mode=MatchMode.EXACT,
            )
        )

    editor.name_edit.setText("账号格式")
    editor.pattern_edit.setText(r"ACCT-\d{4}")
    editor.match_mode_combo.setCurrentIndex(
        editor.match_mode_combo.findData(MatchMode.REGEX.value)
    )
    editor.action_combo.setCurrentIndex(
        editor.action_combo.findData(ActionKind.MASK_MIDDLE.value)
    )
    editor.negative_examples_edit.setText("ACCOUNT-1234")
    editor.scope_checks["all"].setChecked(False)
    editor.scope_checks["text"].setChecked(True)
    editor.scope_checks["ocr"].setChecked(True)
    with pytest.raises(RuleValidationError, match="正例"):
        editor.rule_definition()
    editor.positive_examples_edit.setText("ACCT-1234")
    built = editor.rule_definition()
    assert built.match_mode is MatchMode.REGEX
    assert built.action is ActionKind.MASK_MIDDLE
    assert built.mandatory
    assert built.positive_examples == ("ACCT-1234",)
    assert built.negative_examples == ("ACCOUNT-1234",)
    assert built.applies_to == ("text", "ocr")
    editor.close()
    dialog.close()


def test_import_conflict_dialog_can_adopt_imported_rule(
    tmp_path: Path,
    dialog_qapp: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = RuleStore(tmp_path / "rules.dat", _TestProtector())
    dialog = _dialog(store, dialog_qapp)
    dialog.add_rule(RuleDefinition.fixed("虚构系统", "系统甲"))
    source = tmp_path / "冲突.csv"
    with source.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            ("规则类型", "规则名称", "匹配方式", "关键词", "处理方式", "替换内容")
        )
        writer.writerow(
            ("固定替换", "导入系统代号", "包含", "虚构系统", "固定代号", "系统乙")
        )

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        lambda *_args, **_kwargs: (str(source), "规则文件 (*.csv *.xlsx)"),
    )

    def choose_imported(message_box: QMessageBox) -> int:
        button = next(
            item for item in message_box.buttons() if item.text() == "采用导入规则"
        )
        button.click()
        return 0

    monkeypatch.setattr(QMessageBox, "exec", choose_imported)

    dialog._import_clicked()  # noqa: SLF001

    assert store.load().rules[0].replacement == "系统乙"
    assert "采用导入 1 条" in dialog.feedback_label.text()
    dialog.close()
