from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--screenshot", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root / "src"))
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ.setdefault("QTWEBENGINE_CHROMIUM_FLAGS", "--disable-gpu --no-sandbox")

    from docx import Document
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication

    with tempfile.TemporaryDirectory(prefix="local-redactor-v2-step3-") as temporary:
        isolated = Path(temporary)
        os.environ["LOCALAPPDATA"] = str(isolated / "local-app-data")

        from local_redactor_v2.host import DesktopWindow

        source = isolated / "Gate2验收测试文档.docx"
        document = Document()
        document.add_heading("项目联系信息（虚构测试数据）", level=1)
        document.add_paragraph("项目代号：青云项目")
        document.add_paragraph("联系人：张三")
        document.add_paragraph("联系电话：13800138000")
        document.add_paragraph("邮箱：zhangsan@example.com")
        document.add_paragraph("办公地址：北京市朝阳区测试路88号")
        document.save(source)

        app = QApplication.instance() or QApplication([])
        window = DesktopWindow(root / "ui" / "dist" / "client")
        window.bridge._file_selector = lambda: source
        page = window.view.page()
        results: dict[str, Any] = {"ok": False, "viewports": []}
        state = {"polls": 0, "viewport": 0, "finished": False}
        sizes = [(1536, 1024), (1440, 900), (1280, 720)]

        def run(script: str, callback: Any | None = None) -> None:
            page.runJavaScript(script, callback or (lambda _value: None))

        def finish(code: int) -> None:
            if state["finished"]:
                return
            state["finished"] = True
            results["ok"] = code == 0
            if args.screenshot:
                args.screenshot.parent.mkdir(parents=True, exist_ok=True)
                results["screenshotSaved"] = window.grab().save(str(args.screenshot.resolve()))
                results["screenshot"] = str(args.screenshot.resolve())
            print(json.dumps(results, ensure_ascii=False), flush=True)
            window.close()
            QTimer.singleShot(0, lambda: app.exit(code))

        def begin() -> None:
            run(
                "[...document.querySelectorAll('button')].find(b => "
                "b.textContent.includes('选择单个文件'))?.click()"
            )
            QTimer.singleShot(500, accept_and_start)

        def accept_and_start() -> None:
            run(
                "document.querySelector('input[type=checkbox]')?.click();"
                "[...document.querySelectorAll('button')].find(b => "
                "b.textContent.includes('开始检查'))?.click()"
            )
            QTimer.singleShot(800, poll_step3)

        def poll_step3() -> None:
            state["polls"] += 1

            def checked(value: Any) -> None:
                if value:
                    state["viewport"] = 0
                    measure_next()
                elif state["polls"] >= 60:
                    results["error"] = "step3-timeout"
                    finish(1)
                else:
                    QTimer.singleShot(800, poll_step3)

            run("document.body.innerText.includes('第 3 步 · 确认处理')", checked)

        def measure_next() -> None:
            index = state["viewport"]
            if index >= len(sizes):
                checks = results["viewports"]
                passed = bool(checks) and all(
                    item["step3"]
                    and item["inlineAnchors"] > 0
                    and item["suggestion"]
                    and item["rail"]
                    and not item["separateRecognitionHeading"]
                    and not item["horizontalOverflow"]
                    for item in checks
                )
                finish(0 if passed else 1)
                return
            width, height = sizes[index]
            window.resize(width, height)
            QTimer.singleShot(300, lambda: measure(width, height))

        def measure(width: int, height: int) -> None:
            script = """
                (() => {
                  const root = document.documentElement;
                  return JSON.stringify({
                    step3: document.body.innerText.includes('第 3 步 · 确认处理'),
                    inlineAnchors: document.querySelectorAll('.tf-document [data-finding-id]').length,
                    separateRecognitionHeading: [...document.querySelectorAll('.tf-document h2')].some(x => x.textContent === '本机识别结果'),
                    suggestion: Boolean(document.querySelector('.tf-suggestion-card')),
                    rail: Boolean(document.querySelector('.tf-finding-rail')),
                    horizontalOverflow: root.scrollWidth > root.clientWidth + 1
                  });
                })()
            """

            def measured(value: Any) -> None:
                metrics = json.loads(value) if isinstance(value, str) else {}
                results["viewports"].append({"width": width, "height": height, **metrics})
                state["viewport"] += 1
                measure_next()

            run(script, measured)

        page.loadFinished.connect(
            lambda loaded: QTimer.singleShot(450, begin)
            if loaded
            else (results.update(error="ui-load-failed"), finish(1))
        )
        window.resize(*sizes[0])
        window.show()
        QTimer.singleShot(70_000, lambda: (results.update(error="harness-timeout"), finish(1)))
        return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
