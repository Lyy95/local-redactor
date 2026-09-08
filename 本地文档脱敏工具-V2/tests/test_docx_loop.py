from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from local_redactor.history import HistoryEntry, HistoryStore
from local_redactor.models import (
    Category,
    Finding,
    Modality,
    ProcessingMode,
    TransformMethod,
)
from local_redactor.rule_library import FernetFileProtector, RuleLibrary, RuleStore
from local_redactor.service import LocalDesktopController
from local_redactor_v2.bridge import DesktopBridge


class _Detector:
    def detect(self, document):
        token = "13800138000"
        findings = []
        for block in document.blocks:
            start = block.text.find(token)
            if start < 0:
                continue
            findings.append(
                Finding(
                    category=Category.PHONE,
                    modality=Modality.TEXT,
                    original=token,
                    locations=[replace(block.location, start=start, end=start + len(token))],
                    detector="test:phone",
                    confidence=1.0,
                    suggested_method=TransformMethod.SIMULATE,
                    context="集成测试手机号",
                )
            )
        return findings


class _RuleStore:
    def load(self):
        return RuleLibrary()


class _HistoryStore:
    def __init__(self):
        self.entries: list[HistoryEntry] = []

    def load(self):
        return tuple(self.entries)

    def append(self, entry):
        self.entries = [entry, *(item for item in self.entries if item.id != entry.id)]

    def delete(self, entry_id):
        self.entries = [item for item in self.entries if item.id != entry_id]

    def clear(self):
        self.entries = []


class _TestProtector:
    def protect(self, plaintext: bytes) -> bytes:
        return b"test:" + plaintext

    def unprotect(self, ciphertext: bytes) -> bytes:
        if not ciphertext.startswith(b"test:"):
            raise ValueError("invalid test ciphertext")
        return ciphertext[5:]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _payload(raw: str) -> dict:
    value = json.loads(raw)
    assert value["ok"] is True, value
    return value.get("data", value)


def test_default_local_recognizer_is_available_for_docx(tmp_path: Path) -> None:
    source = tmp_path / "识别能力检查.docx"
    document = Document()
    document.add_paragraph("联系电话：13800138000")
    document.save(source)

    controller = LocalDesktopController(
        rule_store=_RuleStore(),
        history_store=_HistoryStore(),
    )
    bundle = controller.scan(source, ProcessingMode.BALANCED)

    assert bundle.blocking_issues == []
    assert any(
        finding.category is Category.PHONE and finding.original == "13800138000"
        for finding in bundle.findings
    )


def test_bridge_completes_real_docx_review_and_export(tmp_path: Path) -> None:
    source = tmp_path / "待脱敏文档.docx"
    document = Document()
    document.add_heading("项目联系表", level=1)
    document.add_paragraph("项目代号：青云项目")
    document.add_paragraph("联系电话：13800138000")
    document.save(source)
    original_hash = _sha256(source)

    history = _HistoryStore()

    def controller_factory():
        return LocalDesktopController(
            detector=_Detector(),
            image_analyzer_factory=lambda: SimpleNamespace(
                analyze=lambda _document: SimpleNamespace(findings=[], warnings=[])
            ),
            rule_store=_RuleStore(),
            history_store=history,
        )

    bridge = DesktopBridge(
        tmp_path,
        controller_factory=controller_factory,
        file_selector=lambda: source,
        folder_selector=lambda: tmp_path / "输出",
    )

    chosen = _payload(bridge.choose_file())
    created = _payload(bridge.create_task(json.dumps({"source": chosen["source"]})))
    task_id = created["taskId"]
    scanned = _payload(bridge.start_scan(task_id))
    assert scanned["state"] == "review_required"
    assert scanned["pendingCount"] == 1
    assert scanned["findings"][0]["original"] == "13800138000"
    assert scanned["preview"]["blocks"]
    preview_block_ids = {item["id"] for item in scanned["preview"]["blocks"]}
    phone_locations = scanned["findings"][0]["locations"]
    assert phone_locations
    assert phone_locations[0]["blockId"] in preview_block_ids
    assert phone_locations[0]["start"] is not None
    assert phone_locations[0]["end"] is not None

    with_manual = _payload(
        bridge.add_manual_finding(
            task_id,
            json.dumps({"original": "青云项目", "suggestion": "P-001"}),
        )
    )
    assert with_manual["pendingCount"] == 2
    manual = next(item for item in with_manual["findings"] if item["original"] == "青云项目")
    after_manual = _payload(bridge.resolve_finding(task_id, manual["id"], "edit", "P-001"))
    assert after_manual["pendingCount"] == 1

    finding_id = scanned["findings"][0]["id"]
    reviewed = _payload(bridge.resolve_finding(task_id, finding_id, "adopt", ""))
    assert reviewed["state"] == "ready_to_generate"
    assert reviewed["pendingCount"] == 0

    exported = _payload(bridge.export_task(task_id, ""))
    assert exported["cancelled"] is False
    artifacts = exported["artifacts"]
    ai_copy = Path(artifacts["ai"]["path"])
    mapping = Path(artifacts["mapping"]["path"])
    report = Path(artifacts["report"]["path"])
    assert ai_copy.is_file()
    assert mapping.is_file()
    assert report.is_file()
    assert mapping.parent.name == "本地保管"
    assert report.parent.name == "本地保管"
    assert ai_copy.parent.name == "AI交付"
    assert not any(path.name == mapping.name for path in ai_copy.parent.iterdir())

    delivered = "\n".join(paragraph.text for paragraph in Document(ai_copy).paragraphs)
    assert "13800138000" not in delivered
    assert "青云项目" not in delivered
    assert "P-001" in delivered
    assert _sha256(source) == original_hash

    history_rows = _payload(bridge.list_history())["entries"]
    assert len(history_rows) == 1
    assert history_rows[0]["status"] == "completed"
    serialized_history = json.dumps(history_rows, ensure_ascii=False)
    assert str(source) not in serialized_history
    assert "13800138000" not in serialized_history

    reopened_bridge = DesktopBridge(tmp_path, controller_factory=controller_factory)
    reopened_rows = _payload(reopened_bridge.list_history())["entries"]
    assert reopened_rows == history_rows


def test_encrypted_history_store_survives_a_new_process_store_instance(tmp_path: Path) -> None:
    path = tmp_path / "history.dat"
    first = HistoryStore(path=path, protector=_TestProtector())
    first.append(
        HistoryEntry(
            id="completed-1",
            created_at="2026-08-31T10:00:00+08:00",
            source_name="虚构项目方案.docx",
            source_sha256="0" * 64,
            document_kind="docx",
            status="completed",
            finding_count=2,
            transformed_count=2,
            removed_count=0,
            result_path="C:/local-only/result",
        )
    )

    reopened = HistoryStore(path=path, protector=_TestProtector())
    entries = reopened.load()
    assert len(entries) == 1
    assert entries[0].id == "completed-1"
    assert entries[0].source_name == "虚构项目方案.docx"


def test_bridge_applies_rules_incrementally_to_active_task(tmp_path: Path) -> None:
    """Saving a fixed rule mid-review refreshes findings without a full re-scan."""

    source = tmp_path / "增量规则复核.docx"
    document = Document()
    document.add_paragraph("项目代号：青云项目")
    document.add_paragraph("联系电话：13800138000")
    document.save(source)

    store = RuleStore(
        tmp_path / "rules.dat",
        protector=FernetFileProtector(tmp_path / "key"),
        include_presets=False,
    )
    history = _HistoryStore()

    def controller_factory():
        return LocalDesktopController(
            detector=_Detector(),
            image_analyzer_factory=lambda: SimpleNamespace(
                analyze=lambda _document: SimpleNamespace(findings=[], warnings=[])
            ),
            rule_store=store,
            history_store=history,
        )

    bridge = DesktopBridge(
        tmp_path,
        controller_factory=controller_factory,
        file_selector=lambda: source,
        folder_selector=lambda: tmp_path / "输出",
    )

    # No active task -> soft error path for the UI.
    missing = json.loads(bridge.apply_rules_incrementally("missing-task"))
    assert missing["ok"] is False
    assert missing["error"]["code"] == "TASK_NOT_FOUND"

    chosen = _payload(bridge.choose_file())
    created = _payload(bridge.create_task(json.dumps({"source": chosen["source"]})))
    task_id = created["taskId"]

    skipped = _payload(bridge.apply_rules_incrementally(task_id))
    assert skipped["skipped"] is True
    assert skipped["applied"] is False
    assert skipped["addedCount"] == 0

    scanned = _payload(bridge.start_scan(task_id))
    assert scanned["state"] == "review_required"
    before_ids = {item["id"] for item in scanned["findings"]}
    assert not any(item["original"] == "青云项目" for item in scanned["findings"])

    saved = _payload(
        bridge.save_rule(
            json.dumps(
                {
                    "id": "rule-qingyun",
                    "type": "fixed",
                    "name": "青云项目固定替换",
                    "matchMode": "包含",
                    "pattern": "青云项目",
                    "action": "换成固定代号",
                    "replacement": "项目A",
                    "scope": "正文、表格",
                    "mandatory": True,
                    "enabled": True,
                },
                ensure_ascii=False,
            )
        )
    )
    assert any(rule["id"] == "rule-qingyun" for rule in saved["rules"])

    applied = _payload(bridge.apply_rules_incrementally(task_id))
    assert applied["applied"] is True
    assert applied["skipped"] is False
    assert applied["addedCount"] >= 1
    assert any(item["original"] == "青云项目" for item in applied["findings"])
    qingyun = next(item for item in applied["findings"] if item["original"] == "青云项目")
    assert "项目A" in qingyun["suggestion"]
    assert qingyun["rule"]["id"] == "rule-qingyun"
    # Prior detector findings remain addressable after remapping.
    assert before_ids.intersection({item["id"] for item in applied["findings"]}) or any(
        item["original"] == "13800138000" for item in applied["findings"]
    )

    disabled = _payload(bridge.set_rule_enabled("rule-qingyun", False))
    assert disabled["rules"][0]["enabled"] is False
    after_disable = _payload(bridge.apply_rules_incrementally(task_id))
    assert after_disable["applied"] is True
    assert not any(item["original"] == "青云项目" for item in after_disable["findings"])
