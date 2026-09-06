from __future__ import annotations

import contextlib
import csv
import ctypes
import json
import os
import re
import tempfile
import threading
import time
import zipfile
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from enum import StrEnum
from pathlib import Path
from typing import Any, Literal, Protocol
from uuid import uuid4

from openpyxl import load_workbook  # type: ignore[import-untyped]

from local_redactor.core.common import contains_normalized_original

SCHEMA_VERSION = 1
MAX_RULES = 10_000
MAX_RULE_NAME_LENGTH = 120
MAX_PATTERN_LENGTH = 256
MAX_REPLACEMENT_LENGTH = 512
MAX_EXAMPLE_LENGTH = 2_048
MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
MAX_XLSX_ENTRIES = 512
MAX_XLSX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
DEFAULT_SCOPES = ("all",)
PRESET_RULE_ID_PREFIX = "preset-v1-"
CURRENT_PRESET_VERSION = 2
RULE_STORE_LOCK_TIMEOUT_SECONDS = 5.0
ALLOWED_SCOPES = frozenset(
    {
        "all",
        "text",
        "cell",
        "ocr",
        "metadata",
        "hidden",
    }
)

IMPORT_HEADERS = (
    "规则类型",
    "规则名称",
    "匹配方式",
    "关键词",
    "处理方式",
    "替换内容",
)
OPTIONAL_IMPORT_HEADERS = (
    "规则编号",
    "必须处理",
    "启用",
    "优先级",
    "正例",
    "反例",
    "适用范围",
    "区分大小写",
    "保留开头",
    "保留结尾",
    "代号位数",
    "内容类别",
)

_THREAD_LOCKS: dict[str, threading.Lock] = {}
_THREAD_LOCKS_GUARD = threading.Lock()


class RuleLibraryError(RuntimeError):
    """Base error for persistent rule-library operations."""


class RuleValidationError(RuleLibraryError):
    """A rule or imported row is invalid."""


class RuleConflictError(RuleLibraryError):
    """Two equally applicable rules request incompatible outcomes."""


class RuleProtectionError(RuleLibraryError):
    """Current-user encryption or decryption failed."""


class RuleStoreError(RuleLibraryError):
    """The encrypted rule store could not be safely loaded or saved."""


class RuleImportError(RuleLibraryError):
    """A CSV/XLSX import is malformed or unsafe."""


class RuleKind(StrEnum):
    FIXED = "fixed"
    STANDARD = "standard"


class MatchMode(StrEnum):
    EXACT = "exact"
    CONTAINS = "contains"
    REGEX = "regex"
    ANY_KEYWORD = "any_keyword"
    ALL_KEYWORDS = "all_keywords"
    MANUAL = "manual"


class ActionKind(StrEnum):
    FIXED_REPLACEMENT = "fixed_replacement"
    MASK_MIDDLE = "mask_middle"
    SEQUENCE_CODE = "sequence_code"
    DELETE = "delete"


@dataclass(frozen=True, slots=True)
class RuleDefinition:
    """One user-defined rule.

    `EXACT` matches only when the whole reviewed value equals its one pattern.
    `CONTAINS` and `REGEX` return each matching span. `ANY_KEYWORD` returns each
    matching keyword span. `ALL_KEYWORDS` does the same only when every keyword
    occurs in the same reviewed value. `MANUAL` runs only when its rule id is
    explicitly supplied by the caller and targets the whole selected value.

    Rules never decide whether information is classified. They only enforce
    explicit user terms or an explicit manual classification.
    """

    name: str
    kind: RuleKind
    match_mode: MatchMode
    patterns: tuple[str, ...]
    action: ActionKind
    replacement: str = ""
    mandatory: bool = True
    enabled: bool = True
    priority: int = 0
    positive_examples: tuple[str, ...] = ()
    negative_examples: tuple[str, ...] = ()
    applies_to: tuple[str, ...] = DEFAULT_SCOPES
    case_sensitive: bool = True
    keep_prefix: int = 1
    keep_suffix: int = 1
    code_width: int = 2
    category: str = ""
    id: str = field(default_factory=lambda: uuid4().hex)

    def __post_init__(self) -> None:
        object.__setattr__(self, "name", self.name.strip())
        object.__setattr__(self, "replacement", self.replacement.strip())
        object.__setattr__(self, "category", self.category.strip())
        object.__setattr__(self, "id", self.id.strip())
        object.__setattr__(
            self,
            "patterns",
            tuple(pattern.strip() for pattern in self.patterns if pattern.strip()),
        )
        object.__setattr__(
            self,
            "positive_examples",
            tuple(example for example in self.positive_examples if example),
        )
        object.__setattr__(
            self,
            "negative_examples",
            tuple(example for example in self.negative_examples if example),
        )
        object.__setattr__(
            self,
            "applies_to",
            tuple(
                dict.fromkeys(
                    scope.strip().casefold() for scope in self.applies_to if scope.strip()
                )
            ),
        )
        self._validate()

    @classmethod
    def fixed(
        cls,
        source: str,
        replacement_text: str,
        *,
        name: str | None = None,
        match_mode: MatchMode = MatchMode.CONTAINS,
        mandatory: bool = True,
        enabled: bool = True,
        priority: int = 0,
        applies_to: Sequence[str] = DEFAULT_SCOPES,
        case_sensitive: bool = True,
        category: str = "",
        rule_id: str | None = None,
    ) -> RuleDefinition:
        """Build the common one-term fixed-replacement rule."""

        return cls(
            id=rule_id or uuid4().hex,
            name=name or f"{source} → {replacement_text}",
            kind=RuleKind.FIXED,
            match_mode=match_mode,
            patterns=(source,),
            action=ActionKind.FIXED_REPLACEMENT,
            replacement=replacement_text,
            mandatory=mandatory,
            enabled=enabled,
            priority=priority,
            applies_to=tuple(applies_to),
            case_sensitive=case_sensitive,
            category=category,
        )

    @classmethod
    def standard(
        cls,
        name: str,
        *,
        match_mode: MatchMode,
        patterns: Sequence[str] = (),
        action: ActionKind,
        replacement: str = "",
        mandatory: bool = True,
        enabled: bool = True,
        priority: int = 0,
        positive_examples: Sequence[str] = (),
        negative_examples: Sequence[str] = (),
        applies_to: Sequence[str] = DEFAULT_SCOPES,
        case_sensitive: bool = True,
        keep_prefix: int = 1,
        keep_suffix: int = 1,
        code_width: int = 2,
        category: str = "",
        rule_id: str | None = None,
    ) -> RuleDefinition:
        """Build a reusable user criterion without inferring secrecy level."""

        return cls(
            id=rule_id or uuid4().hex,
            name=name,
            kind=RuleKind.STANDARD,
            match_mode=match_mode,
            patterns=tuple(patterns),
            action=action,
            replacement=replacement,
            mandatory=mandatory,
            enabled=enabled,
            priority=priority,
            positive_examples=tuple(positive_examples),
            negative_examples=tuple(negative_examples),
            applies_to=tuple(applies_to),
            case_sensitive=case_sensitive,
            keep_prefix=keep_prefix,
            keep_suffix=keep_suffix,
            code_width=code_width,
            category=category,
        )

    def effect_signature(self) -> tuple[object, ...]:
        return (
            self.action,
            self.replacement,
            self.keep_prefix,
            self.keep_suffix,
            self.code_width,
            self.mandatory,
            self.category,
        )

    def source_signature(self) -> tuple[object, ...]:
        patterns = (
            self.patterns
            if self.case_sensitive
            else tuple(pattern.casefold() for pattern in self.patterns)
        )
        return (
            self.match_mode,
            patterns,
            self.case_sensitive,
            tuple(sorted(self.applies_to)),
        )

    def duplicate_signature(self) -> tuple[object, ...]:
        return (
            self.source_signature(),
            self.effect_signature(),
            self.enabled,
            self.priority,
            self.kind,
        )

    def to_dict(self) -> dict[str, object]:
        return {
            "id": self.id,
            "name": self.name,
            "kind": self.kind.value,
            "match_mode": self.match_mode.value,
            "patterns": list(self.patterns),
            "action": self.action.value,
            "replacement": self.replacement,
            "mandatory": self.mandatory,
            "enabled": self.enabled,
            "priority": self.priority,
            "positive_examples": list(self.positive_examples),
            "negative_examples": list(self.negative_examples),
            "applies_to": list(self.applies_to),
            "case_sensitive": self.case_sensitive,
            "keep_prefix": self.keep_prefix,
            "keep_suffix": self.keep_suffix,
            "code_width": self.code_width,
            "category": self.category,
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, object]) -> RuleDefinition:
        try:
            return cls(
                id=_strict_string(value["id"], "id"),
                name=_strict_string(value["name"], "name"),
                kind=RuleKind(_strict_string(value["kind"], "kind")),
                match_mode=MatchMode(_strict_string(value["match_mode"], "match_mode")),
                patterns=_string_tuple(value["patterns"], "patterns"),
                action=ActionKind(_strict_string(value["action"], "action")),
                replacement=_strict_string(
                    value.get("replacement", ""),
                    "replacement",
                ),
                mandatory=_strict_bool(value.get("mandatory", True), "mandatory"),
                enabled=_strict_bool(value.get("enabled", True), "enabled"),
                priority=_strict_int(value.get("priority", 0), "priority"),
                positive_examples=_string_tuple(
                    value.get("positive_examples", ()),
                    "positive_examples",
                ),
                negative_examples=_string_tuple(
                    value.get("negative_examples", ()),
                    "negative_examples",
                ),
                applies_to=_string_tuple(value.get("applies_to", DEFAULT_SCOPES), "applies_to"),
                case_sensitive=_strict_bool(
                    value.get("case_sensitive", True),
                    "case_sensitive",
                ),
                keep_prefix=_strict_int(value.get("keep_prefix", 1), "keep_prefix"),
                keep_suffix=_strict_int(value.get("keep_suffix", 1), "keep_suffix"),
                code_width=_strict_int(value.get("code_width", 2), "code_width"),
                category=_strict_string(value.get("category", ""), "category"),
            )
        except (KeyError, TypeError, ValueError) as exc:
            raise RuleValidationError("规则数据缺少必要字段或枚举值无效") from exc

    def _validate(self) -> None:
        if not self.id or len(self.id) > 128:
            raise RuleValidationError("规则编号不能为空且不能超过 128 个字符")
        if not self.name or len(self.name) > MAX_RULE_NAME_LENGTH:
            raise RuleValidationError(f"规则名称不能为空且不能超过 {MAX_RULE_NAME_LENGTH} 个字符")
        if not -10_000 <= self.priority <= 10_000:
            raise RuleValidationError("规则优先级必须在 -10000 到 10000 之间")
        if not self.applies_to or any(scope not in ALLOWED_SCOPES for scope in self.applies_to):
            raise RuleValidationError("规则适用范围无效")
        if "all" in self.applies_to and len(self.applies_to) != 1:
            raise RuleValidationError("适用范围为 all 时不能再混用其他范围")
        if self.match_mode is MatchMode.MANUAL:
            if self.patterns:
                raise RuleValidationError("仅人工规则不能配置自动命中关键词")
        elif not self.patterns:
            raise RuleValidationError("非人工规则至少需要一个关键词或表达式")
        if self.match_mode in {MatchMode.EXACT, MatchMode.CONTAINS, MatchMode.REGEX} and (
            len(self.patterns) != 1
        ):
            raise RuleValidationError("精确、包含和受控正则规则只能配置一个匹配内容")
        for pattern in self.patterns:
            if len(pattern) > MAX_PATTERN_LENGTH:
                raise RuleValidationError(f"关键词或表达式不能超过 {MAX_PATTERN_LENGTH} 个字符")
        if self.match_mode is MatchMode.REGEX:
            _compile_controlled_regex(self.patterns[0], self.case_sensitive)
        if len(self.replacement) > MAX_REPLACEMENT_LENGTH:
            raise RuleValidationError(f"替换内容不能超过 {MAX_REPLACEMENT_LENGTH} 个字符")
        if (
            self.action
            in {
                ActionKind.FIXED_REPLACEMENT,
                ActionKind.SEQUENCE_CODE,
            }
            and not self.replacement
        ):
            raise RuleValidationError("固定替换和顺序代号必须填写替换内容或代号前缀")
        if (
            self.action
            in {
                ActionKind.MASK_MIDDLE,
                ActionKind.DELETE,
            }
            and self.replacement
        ):
            raise RuleValidationError("中间星号和删除规则不能填写固定替换内容")
        if not 0 <= self.keep_prefix <= 64 or not 0 <= self.keep_suffix <= 64:
            raise RuleValidationError("保留开头和结尾字符数必须在 0 到 64 之间")
        if not 1 <= self.code_width <= 8:
            raise RuleValidationError("顺序代号位数必须在 1 到 8 之间")
        if any(len(example) > MAX_EXAMPLE_LENGTH for example in self.positive_examples):
            raise RuleValidationError(f"正例不能超过 {MAX_EXAMPLE_LENGTH} 个字符")
        if any(len(example) > MAX_EXAMPLE_LENGTH for example in self.negative_examples):
            raise RuleValidationError(f"反例不能超过 {MAX_EXAMPLE_LENGTH} 个字符")
        if self.match_mode is not MatchMode.MANUAL:
            for example in self.positive_examples:
                if not _raw_spans(self, example):
                    raise RuleValidationError(f"正例未命中规则：{self.name}")
            for example in self.negative_examples:
                if _raw_spans(self, example):
                    raise RuleValidationError(f"反例命中了规则：{self.name}")
        if self.kind is RuleKind.FIXED:
            if self.match_mode not in {MatchMode.EXACT, MatchMode.CONTAINS}:
                raise RuleValidationError("固定替换只允许按完整原词进行文字匹配")
            if self.action is not ActionKind.FIXED_REPLACEMENT:
                raise RuleValidationError("固定替换只能换成一个固定代号")
            if len(self.patterns) != 1:
                raise RuleValidationError("固定替换只能填写一个原词")
            if self.mandatory and contains_normalized_original(
                self.patterns[0],
                self.replacement,
            ):
                raise RuleValidationError("必须替换的代号不能保留完整原词")
        elif not self.positive_examples:
            raise RuleValidationError("判断标准至少需要填写一个正确示例（正例）")


ConflictPolicy = Literal["error", "keep_existing", "use_imported"]


class RuleImportConflictError(RuleConflictError):
    """An import needs an explicit local choice before it can be committed."""

    def __init__(
        self,
        conflicts: Sequence[tuple[RuleDefinition, RuleDefinition]],
    ) -> None:
        self.conflicts = tuple(conflicts)
        names = "、".join(
            f"“{existing.name}”与“{imported.name}”" for existing, imported in self.conflicts[:3]
        )
        remaining = len(self.conflicts) - 3
        suffix = f"等 {len(self.conflicts)} 组" if remaining > 0 else ""
        super().__init__(f"发现处理结果不同的规则：{names}{suffix}")


@dataclass(frozen=True, slots=True)
class RuleLibrary:
    rules: tuple[RuleDefinition, ...] = ()
    revision: int = 0
    schema_version: int = SCHEMA_VERSION
    preset_version: int = 0

    def __post_init__(self) -> None:
        if self.schema_version != SCHEMA_VERSION:
            raise RuleValidationError("规则库版本不受支持")
        if self.revision < 0:
            raise RuleValidationError("规则库修订号无效")
        if self.preset_version < 0:
            raise RuleValidationError("预置规则版本无效")
        if len(self.rules) > MAX_RULES:
            raise RuleValidationError(f"规则总数不能超过 {MAX_RULES}")
        identifiers = [rule.id for rule in self.rules]
        if len(identifiers) != len(set(identifiers)):
            raise RuleConflictError("规则编号重复")
        _validate_source_conflicts(self.rules)

    @property
    def enabled_rules(self) -> tuple[RuleDefinition, ...]:
        return tuple(rule for rule in self.rules if rule.enabled)

    def snapshot(self) -> RuleLibrary:
        """Return an immutable task snapshot detached from later library edits."""

        return RuleLibrary(
            rules=tuple(self.rules),
            revision=self.revision,
            schema_version=self.schema_version,
            preset_version=self.preset_version,
        )

    def add(self, rule: RuleDefinition) -> RuleLibrary:
        if any(existing.id == rule.id for existing in self.rules):
            raise RuleConflictError("规则编号已存在")
        return RuleLibrary(
            rules=(*self.rules, rule),
            revision=self.revision + 1,
            preset_version=self.preset_version,
        )

    def update(self, rule: RuleDefinition) -> RuleLibrary:
        if not any(existing.id == rule.id for existing in self.rules):
            raise KeyError("规则不存在")
        return RuleLibrary(
            rules=tuple(rule if existing.id == rule.id else existing for existing in self.rules),
            revision=self.revision + 1,
            preset_version=self.preset_version,
        )

    def delete(self, rule_id: str) -> RuleLibrary:
        remaining = tuple(rule for rule in self.rules if rule.id != rule_id)
        if len(remaining) == len(self.rules):
            raise KeyError("规则不存在")
        return RuleLibrary(
            rules=remaining,
            revision=self.revision + 1,
            preset_version=self.preset_version,
        )

    def set_enabled(self, rule_id: str, enabled: bool) -> RuleLibrary:
        existing = next((rule for rule in self.rules if rule.id == rule_id), None)
        if existing is None:
            raise KeyError("规则不存在")
        return self.update(replace(existing, enabled=enabled))

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_version": self.schema_version,
            "revision": self.revision,
            "preset_version": self.preset_version,
            "rules": [rule.to_dict() for rule in self.rules],
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, object]) -> RuleLibrary:
        try:
            schema_version = _strict_int(value["schema_version"], "schema_version")
            revision = _strict_int(value.get("revision", 0), "revision")
            preset_version = _strict_int(value.get("preset_version", 0), "preset_version")
            raw_rules = value["rules"]
        except (KeyError, TypeError, ValueError) as exc:
            raise RuleValidationError("规则库结构无效") from exc
        if not isinstance(raw_rules, list):
            raise RuleValidationError("规则列表必须是数组")
        rules = tuple(RuleDefinition.from_dict(_mapping(item, "rule")) for item in raw_rules)
        return cls(
            rules=rules,
            revision=revision,
            schema_version=schema_version,
            preset_version=preset_version,
        )


@dataclass(frozen=True, slots=True)
class RuleMatch:
    rule_id: str
    rule_name: str
    start: int
    end: int
    original: str
    replacement: str
    mandatory: bool
    kind: RuleKind
    action: ActionKind
    category: str
    priority: int


@dataclass(frozen=True, slots=True)
class RuleApplication:
    original: str
    replacement: str
    matches: tuple[RuleMatch, ...]


@dataclass(frozen=True, slots=True)
class _Candidate:
    rule: RuleDefinition
    start: int
    end: int

    @property
    def length(self) -> int:
        return self.end - self.start


class RuleSession:
    """Apply one immutable rule snapshot with task-local sequential aliases."""

    def __init__(self, library: RuleLibrary) -> None:
        self._rules = tuple(library.enabled_rules)
        self._sequence_counts: dict[str, int] = {}
        self._sequence_values: dict[tuple[str, str], str] = {}

    def find(
        self,
        text: str,
        *,
        applies_to: str = "text",
        manual_rule_ids: Iterable[str] = (),
    ) -> tuple[RuleMatch, ...]:
        if applies_to not in ALLOWED_SCOPES - {"all"}:
            raise RuleValidationError("待匹配内容的适用范围无效")
        manual_ids = frozenset(manual_rule_ids)
        available_manual_ids = {
            rule.id
            for rule in self._rules
            if rule.match_mode is MatchMode.MANUAL and _scope_applies(rule, applies_to)
        }
        unavailable_manual_ids = manual_ids - available_manual_ids
        if unavailable_manual_ids:
            raise RuleValidationError("指定的人工判断标准不存在、未启用或不适用于当前内容")
        candidates: list[_Candidate] = []
        for rule in self._rules:
            if not _scope_applies(rule, applies_to):
                continue
            if rule.match_mode is MatchMode.MANUAL:
                if rule.id in manual_ids and text:
                    candidates.append(_Candidate(rule, 0, len(text)))
                continue
            for start, end in _raw_spans(rule, text):
                candidates.append(_Candidate(rule, start, end))
        selected = _select_candidates(candidates)
        matches: list[RuleMatch] = []
        for candidate in selected:
            original = text[candidate.start : candidate.end]
            replacement_text = self._replacement_for(candidate.rule, original)
            if (
                candidate.rule.mandatory
                and candidate.rule.action is not ActionKind.DELETE
                and contains_normalized_original(original, replacement_text)
            ):
                raise RuleValidationError(f"规则“{candidate.rule.name}”的结果仍保留完整原词")
            matches.append(
                RuleMatch(
                    rule_id=candidate.rule.id,
                    rule_name=candidate.rule.name,
                    start=candidate.start,
                    end=candidate.end,
                    original=original,
                    replacement=replacement_text,
                    mandatory=candidate.rule.mandatory,
                    kind=candidate.rule.kind,
                    action=candidate.rule.action,
                    category=candidate.rule.category,
                    priority=candidate.rule.priority,
                )
            )
        return tuple(matches)

    def apply(
        self,
        text: str,
        *,
        applies_to: str = "text",
        manual_rule_ids: Iterable[str] = (),
    ) -> RuleApplication:
        """Apply findings once against the original text.

        Replacements are spliced back from right to left and are deliberately
        never scanned again during this call.
        """

        matches = self.find(
            text,
            applies_to=applies_to,
            manual_rule_ids=manual_rule_ids,
        )
        output = text
        for match in reversed(matches):
            output = f"{output[: match.start]}{match.replacement}{output[match.end :]}"
        return RuleApplication(
            original=text,
            replacement=output,
            matches=matches,
        )

    def _replacement_for(self, rule: RuleDefinition, original: str) -> str:
        if rule.action is ActionKind.FIXED_REPLACEMENT:
            return rule.replacement
        if rule.action is ActionKind.DELETE:
            return ""
        if rule.action is ActionKind.MASK_MIDDLE:
            return _mask_middle(
                original,
                keep_prefix=rule.keep_prefix,
                keep_suffix=rule.keep_suffix,
            )
        key_text = original if rule.case_sensitive else original.casefold()
        key = (rule.id, key_text)
        existing = self._sequence_values.get(key)
        if existing is not None:
            return existing
        index = self._sequence_counts.get(rule.id, 0) + 1
        self._sequence_counts[rule.id] = index
        replacement_text = f"{rule.replacement}{index:0{rule.code_width}d}"
        self._sequence_values[key] = replacement_text
        return replacement_text


class Protector(Protocol):
    """Encrypt/decrypt bytes without persisting plaintext."""

    def protect(self, plaintext: bytes) -> bytes: ...

    def unprotect(self, ciphertext: bytes) -> bytes: ...


class WindowsDpapiProtector:
    """Windows DPAPI protection bound to the current Windows user."""

    _CRYPTPROTECT_UI_FORBIDDEN = 0x01
    _ENTROPY = b"LocalRedactor.RuleLibrary.v1"

    class _DataBlob(ctypes.Structure):
        _fields_ = [
            ("cbData", ctypes.c_ulong),
            ("pbData", ctypes.POINTER(ctypes.c_ubyte)),
        ]

    def protect(self, plaintext: bytes) -> bytes:
        return self._crypt(plaintext, decrypt=False)

    def unprotect(self, ciphertext: bytes) -> bytes:
        return self._crypt(ciphertext, decrypt=True)

    def _crypt(self, value: bytes, *, decrypt: bool) -> bytes:
        if os.name != "nt":
            raise RuleProtectionError("当前系统不支持 Windows 用户级加密")
        win_dll = getattr(ctypes, "WinDLL", None)
        if win_dll is None:
            raise RuleProtectionError("Windows 加密组件不可用")
        crypt32: Any = win_dll("crypt32.dll", use_last_error=True)
        kernel32: Any = win_dll("kernel32.dll", use_last_error=True)

        source_buffer, source_blob = self._make_blob(value)
        entropy_buffer, entropy_blob = self._make_blob(self._ENTROPY)
        output_blob = self._DataBlob()
        arguments: tuple[Any, ...]

        if decrypt:
            function = crypt32.CryptUnprotectData
            function.argtypes = [
                ctypes.POINTER(self._DataBlob),
                ctypes.c_void_p,
                ctypes.POINTER(self._DataBlob),
                ctypes.c_void_p,
                ctypes.c_void_p,
                ctypes.c_ulong,
                ctypes.POINTER(self._DataBlob),
            ]
            arguments = (
                ctypes.byref(source_blob),
                None,
                ctypes.byref(entropy_blob),
                None,
                None,
                self._CRYPTPROTECT_UI_FORBIDDEN,
                ctypes.byref(output_blob),
            )
        else:
            function = crypt32.CryptProtectData
            function.argtypes = [
                ctypes.POINTER(self._DataBlob),
                ctypes.c_wchar_p,
                ctypes.POINTER(self._DataBlob),
                ctypes.c_void_p,
                ctypes.c_void_p,
                ctypes.c_ulong,
                ctypes.POINTER(self._DataBlob),
            ]
            arguments = (
                ctypes.byref(source_blob),
                "本地文档脱敏工具规则库",
                ctypes.byref(entropy_blob),
                None,
                None,
                self._CRYPTPROTECT_UI_FORBIDDEN,
                ctypes.byref(output_blob),
            )
        function.restype = ctypes.c_int
        if not function(*arguments):
            code = ctypes.get_last_error()
            raise RuleProtectionError(f"Windows 用户级加密操作失败（代码 {code}）")
        try:
            return ctypes.string_at(output_blob.pbData, output_blob.cbData)
        finally:
            # Keep both native input buffers alive until the DPAPI call has
            # completed; DATA_BLOB contains only raw pointers.
            _ = (source_buffer, entropy_buffer)
            kernel32.LocalFree.argtypes = [ctypes.c_void_p]
            kernel32.LocalFree.restype = ctypes.c_void_p
            kernel32.LocalFree(output_blob.pbData)

    @classmethod
    def _make_blob(cls, value: bytes) -> tuple[Any, _DataBlob]:
        buffer = ctypes.create_string_buffer(value, max(1, len(value)))
        pointer = ctypes.cast(buffer, ctypes.POINTER(ctypes.c_ubyte))
        return buffer, cls._DataBlob(len(value), pointer)


def default_rule_store_path() -> Path:
    """Return the private per-user Windows location for `rules.dat`."""

    local_app_data = os.environ.get("LOCALAPPDATA", "").strip()
    if not local_app_data:
        raise RuleStoreError("无法确定当前 Windows 用户的本地应用数据目录")
    local_root = Path(local_app_data)
    if not local_root.is_absolute():
        raise RuleStoreError("当前 Windows 用户的本地应用数据目录无效")
    return local_root / "LocalRedactor" / "本地文档脱敏工具" / "rules.dat"


def is_preset_rule(rule: RuleDefinition) -> bool:
    return rule.id.startswith(PRESET_RULE_ID_PREFIX)


def default_industry_rules() -> tuple[RuleDefinition, ...]:
    """Return deterministic read-only presets shipped with the application."""

    public_security = (
        ("公安", "GA", "organization"),
        ("网络警察", "WJ", "department"),
        ("网警", "WJ", "department"),
        ("技术侦察", "JZ", "department"),
        ("技侦", "JZ", "department"),
        ("刑事侦查", "XZ", "department"),
        ("刑侦", "XZ", "department"),
        ("治安", "ZA", "department"),
        ("交通警察", "JJ", "department"),
        ("交警", "JJ", "department"),
        ("特警", "TJ", "department"),
        ("巡警", "XJ", "department"),
        ("禁毒", "JD", "department"),
        ("经济犯罪侦查", "JJFZZC", "department"),
        ("经侦", "JJFZZC", "department"),
        ("情报", "QB", "department"),
        ("法制部门", "FZBM", "department"),
        ("反诈中心", "FZZX", "department"),
        ("指挥中心", "ZHZX", "department"),
        ("警务督察", "JWDC", "department"),
        ("出入境", "CRJ", "department"),
        ("户政", "HZ", "department"),
    )
    places = (
        ("海南省", "HN"),
        ("海南", "HN"),
        ("福建省", "FJ"),
        ("福建", "FJ"),
        ("广东省", "GD"),
        ("广东", "GD"),
        ("北京市", "BJ"),
        ("北京", "BJ"),
        ("天津市", "TJ"),
        ("天津", "TJ"),
        ("河北省", "HB"),
        ("河北", "HB"),
        ("山西省", "SX"),
        ("山西", "SX"),
        ("内蒙古自治区", "NMG"),
        ("内蒙古", "NMG"),
        ("辽宁省", "LN"),
        ("辽宁", "LN"),
        ("吉林省", "JL"),
        ("吉林", "JL"),
        ("黑龙江省", "HLJ"),
        ("黑龙江", "HLJ"),
        ("上海市", "SH"),
        ("上海", "SH"),
        ("江苏省", "JS"),
        ("江苏", "JS"),
        ("浙江省", "ZJ"),
        ("浙江", "ZJ"),
        ("安徽省", "AH"),
        ("安徽", "AH"),
        ("江西省", "JX"),
        ("江西", "JX"),
        ("山东省", "SD"),
        ("山东", "SD"),
        ("河南省", "HEN"),
        ("河南", "HEN"),
        ("湖北省", "HUB"),
        ("湖北", "HUB"),
        ("湖南省", "HUN"),
        ("湖南", "HUN"),
        ("广西壮族自治区", "GX"),
        ("广西", "GX"),
        ("重庆市", "CQ"),
        ("重庆", "CQ"),
        ("四川省", "SC"),
        ("四川", "SC"),
        ("贵州省", "GZ"),
        ("贵州", "GZ"),
        ("云南省", "YN"),
        ("云南", "YN"),
        ("西藏自治区", "XZ"),
        ("西藏", "XZ"),
        ("陕西省", "SAX"),
        ("陕西", "SAX"),
        ("甘肃省", "GS"),
        ("甘肃", "GS"),
        ("青海省", "QH"),
        ("青海", "QH"),
        ("宁夏回族自治区", "NX"),
        ("宁夏", "NX"),
        ("新疆维吾尔自治区", "XJ"),
        ("新疆", "XJ"),
        ("香港特别行政区", "HK"),
        ("香港", "HK"),
        ("澳门特别行政区", "MO"),
        ("澳门", "MO"),
        ("台湾省", "TW"),
        ("台湾", "TW"),
    )
    definitions: list[RuleDefinition] = []
    for index, (source, replacement, category) in enumerate(public_security):
        definitions.append(
            RuleDefinition.fixed(
                source,
                replacement,
                name=f"公安行业预置：{source}",
                priority=100 + len(source),
                category=category,
                rule_id=f"{PRESET_RULE_ID_PREFIX}police-{index:03d}",
            )
        )
    for index, (source, replacement) in enumerate(places):
        definitions.append(
            RuleDefinition.fixed(
                source,
                replacement,
                name=f"地名预置：{source}",
                priority=50 + len(source),
                category="location",
                rule_id=f"{PRESET_RULE_ID_PREFIX}place-{index:03d}",
            )
        )
    return tuple(definitions)


def default_standard_rules() -> tuple[RuleDefinition, ...]:
    """Return conservative, visible criteria that users may edit or remove."""

    specs = (
        (
            "手机号码",
            r"1[3-9]\d{9}",
            ActionKind.MASK_MIDDLE,
            "",
            "13812345678",
            "12345678901",
            3,
            4,
            "phone",
            False,
        ),
        (
            "身份证号",
            r"\d{17}[\dXx]",
            ActionKind.MASK_MIDDLE,
            "",
            "420101199003051234",
            "20260807123000001",
            6,
            4,
            "id_card",
            False,
        ),
        (
            "16 位银行卡号",
            r"\d{16}",
            ActionKind.MASK_MIDDLE,
            "",
            "6222021234567890",
            "202608071230000",
            6,
            4,
            "bank_card",
            False,
        ),
        (
            "标准车牌号",
            r"[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼][A-Z][A-Z0-9]{5}",
            ActionKind.SEQUENCE_CODE,
            "车辆",
            "闽A12345",
            "CASEA12345",
            0,
            0,
            "vehicle_plate",
            False,
        ),
    )
    definitions: list[RuleDefinition] = []
    for index, (
        name,
        pattern,
        action,
        replacement,
        positive,
        negative,
        keep_prefix,
        keep_suffix,
        category,
        enabled,
    ) in enumerate(specs):
        definitions.append(
            RuleDefinition.standard(
                name,
                match_mode=MatchMode.REGEX,
                patterns=(pattern,),
                action=action,
                replacement=replacement if action is ActionKind.SEQUENCE_CODE else "",
                mandatory=category == "secret",
                enabled=enabled,
                priority=200,
                positive_examples=(positive,),
                negative_examples=(negative,),
                keep_prefix=keep_prefix,
                keep_suffix=keep_suffix,
                category=category,
                case_sensitive=False,
                rule_id=f"{PRESET_RULE_ID_PREFIX}standard-{index:03d}",
            )
        )
    manual_specs = (
        ("中文姓名", "人员", "张三", "张三丰", "name"),
        ("电子邮箱", "email", "abcde@163.com", "abcde.example.com", "email"),
        ("固定电话", "TEL", "010-88888888", "13812345678", "landline"),
        (
            "密码、口令与 PIN（内置识别后自动删除）",
            "SECRET",
            "密码：Abc#1234",
            "密码管理规范",
            "secret",
        ),
        (
            "Token/API Key/AccessKey/SecretKey（内置识别后自动删除）",
            "SECRET",
            "token: eyJhbGciOi",
            "token 管理办法",
            "secret",
        ),
        ("IPv4 / IPv6", "IP", "192.168.1.10", "192.168.1", "private_ip"),
        ("域名", "domain", "service.example.com", "report_docx", "domain"),
        ("案件编号", "CASE", "案号：闽0101执法20260001", "案件编号管理", "case_id"),
        ("设备编号", "DEVICE", "SN：ABC-2026-001", "SN 字段说明", "device_id"),
        ("详细地址", "地址", "福州市鼓楼区鼓屏路 1 号", "福建省", "address"),
        ("账号与用户名", "user", "用户名：zhangsan", "用户管理", "username"),
        ("单位与机构", "ORG", "福建省某信息中心", "有关单位", "organization"),
        ("部门与警种", "DEPT", "网络警察支队", "业务部门", "department"),
        ("项目名称", "PROJECT", "智慧警务建设项目", "建设项目", "project"),
        ("系统名称", "SYSTEM", "智慧警务平台", "业务系统", "system"),
        ("精确地点", "PLACE", "福建省福州市", "各地市", "location"),
    )
    offset = len(definitions)
    for index, (name, replacement, positive, negative, category) in enumerate(manual_specs):
        definitions.append(
            RuleDefinition.standard(
                name,
                match_mode=MatchMode.MANUAL,
                action=ActionKind.SEQUENCE_CODE,
                replacement=replacement,
                mandatory=False,
                enabled=False,
                positive_examples=(positive,),
                negative_examples=(negative,),
                category=category,
                rule_id=f"{PRESET_RULE_ID_PREFIX}standard-{offset + index:03d}",
            )
        )
    return tuple(definitions)


def _merge_default_industry_rules(library: RuleLibrary) -> RuleLibrary:
    if library.preset_version >= CURRENT_PRESET_VERSION:
        return library
    existing_sources = {rule.source_signature() for rule in library.rules}
    missing = [
        rule
        for rule in (*default_industry_rules(), *default_standard_rules())
        if rule.source_signature() not in existing_sources
    ]
    return RuleLibrary(
        (*library.rules, *missing),
        revision=library.revision,
        preset_version=CURRENT_PRESET_VERSION,
    )


def restore_default_rules(library: RuleLibrary) -> RuleLibrary:
    """Restore missing shipped presets only after an explicit user action."""

    existing_ids = {rule.id for rule in library.rules}
    missing = [
        rule
        for rule in (*default_industry_rules(), *default_standard_rules())
        if rule.id not in existing_ids
    ]
    if not missing:
        return library
    return RuleLibrary(
        (*library.rules, *missing),
        revision=library.revision + 1,
        preset_version=CURRENT_PRESET_VERSION,
    )


def _thread_lock_for(path: Path) -> threading.Lock:
    key = str(path.resolve(strict=False)).casefold()
    with _THREAD_LOCKS_GUARD:
        return _THREAD_LOCKS.setdefault(key, threading.Lock())


@contextmanager
def _rule_store_lock(path: Path) -> Iterator[None]:
    """Serialize one rules.dat transaction across threads and app processes."""

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise RuleStoreError("无法创建本地规则库目录") from exc
    lock_path = path.with_name(f".{path.name}.lock")
    deadline = time.monotonic() + RULE_STORE_LOCK_TIMEOUT_SECONDS
    thread_lock = _thread_lock_for(path)
    if not thread_lock.acquire(timeout=RULE_STORE_LOCK_TIMEOUT_SECONDS):
        raise RuleStoreError("规则库正在被另一窗口修改，请稍后重试")
    try:
        try:
            stream = lock_path.open("a+b")
        except OSError as exc:
            raise RuleStoreError("无法打开规则库写入锁") from exc
        with stream:
            stream.seek(0, os.SEEK_END)
            if stream.tell() == 0:
                stream.write(b"\0")
                stream.flush()
            stream.seek(0)
            if os.name == "nt":
                locker: Any = __import__("msvcrt")
                while True:
                    try:
                        locker.locking(stream.fileno(), locker.LK_NBLCK, 1)
                        break
                    except OSError as exc:
                        if time.monotonic() >= deadline:
                            raise RuleStoreError("规则库正在被另一窗口修改，请稍后重试") from exc
                        time.sleep(0.05)
                try:
                    yield
                finally:
                    stream.seek(0)
                    with contextlib.suppress(OSError):
                        locker.locking(stream.fileno(), locker.LK_UNLCK, 1)
            else:
                locker = __import__("fcntl")
                while True:
                    try:
                        locker.flock(stream.fileno(), locker.LOCK_EX | locker.LOCK_NB)
                        break
                    except OSError as exc:
                        if time.monotonic() >= deadline:
                            raise RuleStoreError("规则库正在被另一窗口修改，请稍后重试") from exc
                        time.sleep(0.05)
                try:
                    yield
                finally:
                    with contextlib.suppress(OSError):
                        locker.flock(stream.fileno(), locker.LOCK_UN)
    finally:
        thread_lock.release()


class RuleStore:
    """Encrypted persistence with atomic writes and revision-checked updates."""

    def __init__(
        self,
        path: Path | None = None,
        protector: Protector | None = None,
        *,
        include_presets: bool = False,
    ) -> None:
        self.path = default_rule_store_path() if path is None else Path(path)
        self.protector = protector or WindowsDpapiProtector()
        self.include_presets = include_presets

    def load(self) -> RuleLibrary:
        with _rule_store_lock(self.path):
            return self._load_unlocked()

    def save(self, library: RuleLibrary) -> None:
        with _rule_store_lock(self.path):
            self._save_unlocked(library)

    def save_if_revision(
        self,
        library: RuleLibrary,
        *,
        expected_revision: int,
    ) -> None:
        """Save only if another window has not changed the loaded revision."""

        with _rule_store_lock(self.path):
            current = self._load_unlocked()
            if current.revision != expected_revision:
                raise RuleConflictError("规则库已在其他窗口更新，本次未覆盖最新内容；请刷新后重试")
            self._save_unlocked(library)

    def update(
        self,
        change: Callable[[RuleLibrary], RuleLibrary],
    ) -> RuleLibrary:
        """Run one complete read-modify-write transaction under the path lock."""

        with _rule_store_lock(self.path):
            current = self._load_unlocked()
            updated = change(current)
            self._save_unlocked(updated)
            return updated

    def _load_unlocked(self) -> RuleLibrary:
        if not self.path.exists():
            library = RuleLibrary()
            return _merge_default_industry_rules(library) if self.include_presets else library
        if not self.path.is_file():
            raise RuleStoreError("规则库路径不是文件")
        try:
            encrypted = self.path.read_bytes()
        except OSError as exc:
            raise RuleStoreError("无法读取本地规则库") from exc
        if not encrypted:
            raise RuleStoreError("本地规则库为空或已损坏")
        plaintext = bytearray()
        try:
            plaintext.extend(self.protector.unprotect(encrypted))
            decoded = json.loads(plaintext.decode("utf-8"))
            library = RuleLibrary.from_dict(_mapping(decoded, "library"))
        except RuleLibraryError:
            raise
        except (UnicodeDecodeError, json.JSONDecodeError, OSError, ValueError) as exc:
            raise RuleStoreError("本地规则库无法解密或内容已损坏") from exc
        finally:
            for index in range(len(plaintext)):
                plaintext[index] = 0
        return _merge_default_industry_rules(library) if self.include_presets else library

    def _save_unlocked(self, library: RuleLibrary) -> None:
        plaintext = bytearray(
            json.dumps(
                library.to_dict(),
                ensure_ascii=False,
                separators=(",", ":"),
                sort_keys=True,
            ).encode("utf-8")
        )
        try:
            encrypted = self.protector.protect(bytes(plaintext))
        except RuleLibraryError:
            raise
        except Exception as exc:
            raise RuleStoreError("本地规则库加密失败") from exc
        finally:
            for index in range(len(plaintext)):
                plaintext[index] = 0
        if not encrypted:
            raise RuleStoreError("本地规则库加密结果为空")

        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise RuleStoreError("无法创建本地规则库目录") from exc

        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="wb",
                dir=self.path.parent,
                prefix=f".{self.path.name}.",
                suffix=".tmp",
                delete=False,
            ) as stream:
                temporary_path = Path(stream.name)
                stream.write(encrypted)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary_path, self.path)
        except OSError as exc:
            if temporary_path is not None:
                with contextlib.suppress(OSError):
                    temporary_path.unlink(missing_ok=True)
            raise RuleStoreError("无法原子保存本地规则库") from exc


@dataclass(frozen=True, slots=True)
class RuleImportResult:
    library: RuleLibrary
    added_count: int
    skipped_duplicate_count: int
    kept_existing_count: int = 0
    replaced_existing_count: int = 0


def import_rules(
    path: Path,
    library: RuleLibrary,
    *,
    conflict_policy: ConflictPolicy = "error",
) -> RuleImportResult:
    """Read a local CSV/XLSX rule sheet without evaluating formulas.

    The function returns a new immutable library. It never mutates or saves the
    existing library, so any conflict or malformed row aborts the whole import.
    """

    source = Path(path)
    try:
        size = source.stat().st_size
    except OSError as exc:
        raise RuleImportError("无法读取规则导入文件") from exc
    if not source.is_file() or size <= 0:
        raise RuleImportError("规则导入文件为空或不存在")
    if size > MAX_IMPORT_BYTES:
        raise RuleImportError("规则导入文件过大")

    suffix = source.suffix.casefold()
    if suffix == ".csv":
        rows = _read_csv_rows(source)
    elif suffix == ".xlsx":
        rows = _read_xlsx_rows(source)
    else:
        raise RuleImportError("规则导入只支持 .csv 和 .xlsx")
    imported = tuple(_rule_from_import_row(row, index) for index, row in enumerate(rows, 2))
    if not imported:
        raise RuleImportError("规则导入文件没有数据行")
    if len(imported) + len(library.rules) > MAX_RULES:
        raise RuleImportError(f"导入后规则总数不能超过 {MAX_RULES}")
    return merge_imported_rules(
        library,
        imported,
        conflict_policy=conflict_policy,
    )


def merge_imported_rules(
    library: RuleLibrary,
    imported: Sequence[RuleDefinition],
    *,
    conflict_policy: ConflictPolicy = "error",
) -> RuleImportResult:
    if conflict_policy not in {"error", "keep_existing", "use_imported"}:
        raise ValueError("导入冲突处理方式无效")
    rules = list(library.rules)
    added = 0
    skipped = 0
    kept_existing = 0
    replaced_existing = 0
    imported_effects: dict[tuple[object, ...], tuple[object, ...]] = {}

    for rule in imported:
        source = rule.source_signature()
        imported_effect = imported_effects.get(source)
        if imported_effect is not None and imported_effect != rule.effect_signature():
            raise RuleImportError("导入文件内部存在同一命中条件、不同处理结果的规则")
        imported_effects[source] = rule.effect_signature()

        duplicate_signatures = {item.duplicate_signature() for item in rules}
        if rule.duplicate_signature() in duplicate_signatures:
            skipped += 1
            continue

        same_source = [
            (index, item) for index, item in enumerate(rules) if item.source_signature() == source
        ]
        conflicting = [
            (index, item)
            for index, item in same_source
            if item.effect_signature() != rule.effect_signature()
        ]
        if conflicting:
            if conflict_policy == "error":
                raise RuleImportConflictError(tuple((item, rule) for _index, item in conflicting))
            if conflict_policy == "keep_existing":
                kept_existing += 1
                continue

            first_index, first_existing = same_source[0]
            candidate = replace(rule, id=first_existing.id)
            source_indexes = {index for index, _item in same_source}
            rules = [
                candidate if index == first_index else item
                for index, item in enumerate(rules)
                if index == first_index or index not in source_indexes
            ]
            replaced_existing += 1
            continue

        seen_ids = {item.id for item in rules}
        candidate = rule
        if candidate.id in seen_ids:
            candidate = replace(candidate, id=uuid4().hex)
        rules.append(candidate)
        added += 1
    changed = bool(added or replaced_existing)
    return RuleImportResult(
        library=RuleLibrary(
            rules=tuple(rules),
            revision=library.revision + (1 if changed else 0),
            preset_version=library.preset_version,
        ),
        added_count=added,
        skipped_duplicate_count=skipped,
        kept_existing_count=kept_existing,
        replaced_existing_count=replaced_existing,
    )


def _raw_spans(rule: RuleDefinition, text: str) -> list[tuple[int, int]]:
    if not text or rule.match_mode is MatchMode.MANUAL:
        return []
    if rule.match_mode is MatchMode.EXACT:
        equal = (
            text == rule.patterns[0]
            if rule.case_sensitive
            else text.casefold() == rule.patterns[0].casefold()
        )
        return [(0, len(text))] if equal else []
    if rule.match_mode is MatchMode.CONTAINS:
        return _literal_spans(
            text,
            rule.patterns[0],
            case_sensitive=rule.case_sensitive,
        )
    if rule.match_mode is MatchMode.REGEX:
        compiled = _compile_controlled_regex(rule.patterns[0], rule.case_sensitive)
        regex_spans = [match.span() for match in compiled.finditer(text)]
        if any(start == end for start, end in regex_spans):
            raise RuleValidationError("受控正则不能产生空命中")
        return regex_spans
    if rule.match_mode is MatchMode.ALL_KEYWORDS and any(
        not _literal_spans(
            text,
            pattern,
            case_sensitive=rule.case_sensitive,
        )
        for pattern in rule.patterns
    ):
        return []
    spans: list[tuple[int, int]] = []
    for pattern in rule.patterns:
        spans.extend(
            _literal_spans(
                text,
                pattern,
                case_sensitive=rule.case_sensitive,
            )
        )
    return sorted(set(spans))


def _literal_spans(
    text: str,
    pattern: str,
    *,
    case_sensitive: bool,
) -> list[tuple[int, int]]:
    flags = 0 if case_sensitive else re.IGNORECASE
    return [match.span() for match in re.finditer(re.escape(pattern), text, flags)]


def _select_candidates(candidates: Sequence[_Candidate]) -> list[_Candidate]:
    ranked = sorted(
        candidates,
        key=lambda item: (
            -item.length,
            -item.rule.priority,
            item.start,
            item.end,
            item.rule.id,
        ),
    )
    selected: list[_Candidate] = []
    for candidate in ranked:
        overlaps = [
            existing
            for existing in selected
            if candidate.start < existing.end and candidate.end > existing.start
        ]
        if not overlaps:
            selected.append(candidate)
            continue
        strongest = max(
            overlaps,
            key=lambda item: (item.length, item.rule.priority),
        )
        candidate_rank = (candidate.length, candidate.rule.priority)
        strongest_rank = (strongest.length, strongest.rule.priority)
        if candidate_rank < strongest_rank:
            continue
        if candidate_rank > strongest_rank:
            selected = [
                existing
                for existing in selected
                if not (candidate.start < existing.end and candidate.end > existing.start)
            ]
            selected.append(candidate)
            continue
        if (
            candidate.start == strongest.start
            and candidate.end == strongest.end
            and candidate.rule.effect_signature() == strongest.rule.effect_signature()
        ):
            continue
        raise RuleConflictError(f"规则“{candidate.rule.name}”与“{strongest.rule.name}”命中范围冲突")
    return sorted(selected, key=lambda item: (item.start, item.end))


def _scope_applies(rule: RuleDefinition, scope: str) -> bool:
    return "all" in rule.applies_to or scope in rule.applies_to


def _mask_middle(value: str, *, keep_prefix: int, keep_suffix: int) -> str:
    if not value:
        return ""
    prefix_count = min(keep_prefix, max(0, len(value) - 1))
    suffix_count = min(keep_suffix, max(0, len(value) - prefix_count - 1))
    masked_count = len(value) - prefix_count - suffix_count
    suffix = value[len(value) - suffix_count :] if suffix_count else ""
    return f"{value[:prefix_count]}{'*' * masked_count}{suffix}"


def _compile_controlled_regex(pattern: str, case_sensitive: bool) -> re.Pattern[str]:
    """Compile a deliberately small, bounded regex subset.

    Supported expressions may use literals, character classes, anchors,
    selected character escapes, and exact `{m}` quantifiers where
    `1 <= m <= 64`. Groups, alternation, dot wildcards, zero-width escapes,
    optional or variable repetitions, backreferences and unbounded
    quantifiers are rejected to avoid uncontrolled backtracking.
    """

    if not pattern or len(pattern) > MAX_PATTERN_LENGTH:
        raise RuleValidationError("受控正则不能为空或过长")
    in_class = False
    escaped = False
    index = 0
    while index < len(pattern):
        char = pattern[index]
        if escaped:
            if char.isdigit() or char in {"g", "k", "b", "B"}:
                raise RuleValidationError("受控正则包含未允许的转义表达式")
            escaped = False
            index += 1
            continue
        if char == "\\":
            escaped = True
            index += 1
            continue
        if char == "[":
            if in_class:
                raise RuleValidationError("受控正则字符类不能嵌套")
            in_class = True
            index += 1
            continue
        if char == "]":
            if not in_class:
                raise RuleValidationError("受控正则字符类括号不匹配")
            in_class = False
            index += 1
            continue
        if not in_class and char in "().*+?|":
            raise RuleValidationError("受控正则包含未允许的表达式")
        if not in_class and char == "{":
            end = pattern.find("}", index + 1)
            if end < 0:
                raise RuleValidationError("受控正则的重复次数缺少右括号")
            token = pattern[index + 1 : end]
            match = re.fullmatch(r"\d{1,2}", token)
            if match is None:
                raise RuleValidationError("受控正则只允许固定重复次数")
            repetitions = int(token)
            if repetitions < 1 or repetitions > 64:
                raise RuleValidationError("受控正则重复次数必须在 1 到 64 之间")
            index = end + 1
            continue
        index += 1
    if escaped or in_class:
        raise RuleValidationError("受控正则表达式不完整")
    flags = 0 if case_sensitive else re.IGNORECASE
    try:
        compiled = re.compile(pattern, flags)
    except re.error as exc:
        raise RuleValidationError("受控正则表达式无效") from exc
    if compiled.search("") is not None:
        raise RuleValidationError("受控正则不能匹配空内容")
    return compiled


def _validate_source_conflicts(rules: Sequence[RuleDefinition]) -> None:
    effects: dict[tuple[object, ...], tuple[object, ...]] = {}
    for rule in rules:
        if rule.match_mode is MatchMode.MANUAL:
            # Manual criteria are selected by id and therefore do not compete
            # for one automatic source signature.
            continue
        source = rule.source_signature()
        previous = effects.get(source)
        if previous is not None and previous != rule.effect_signature():
            raise RuleConflictError(f"同一命中条件存在不同处理结果：{rule.name}")
        effects[source] = rule.effect_signature()


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, strict=True)
            if reader.fieldnames is None:
                raise RuleImportError("CSV 缺少表头")
            headers = tuple((header or "").strip() for header in reader.fieldnames)
            _validate_import_headers(headers)
            rows: list[dict[str, str]] = []
            for index, row in enumerate(reader, start=2):
                if index > MAX_IMPORT_ROWS + 1:
                    raise RuleImportError(f"规则导入不能超过 {MAX_IMPORT_ROWS} 行")
                normalized = {
                    str(key).strip(): "" if value is None else str(value).strip()
                    for key, value in row.items()
                    if key is not None
                }
                if not any(normalized.values()):
                    raise RuleImportError(f"第 {index} 行为空")
                _reject_formula_like_values(normalized.values(), index)
                rows.append(normalized)
            return rows
    except UnicodeDecodeError as exc:
        raise RuleImportError("CSV 必须使用 UTF-8 编码") from exc
    except csv.Error as exc:
        raise RuleImportError("CSV 格式无效") from exc
    except OSError as exc:
        raise RuleImportError("无法读取 CSV 规则文件") from exc


def _read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    _validate_xlsx_container(path)
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise RuleImportError("XLSX 规则文件无法安全打开") from exc
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration as exc:
            raise RuleImportError("XLSX 缺少表头") from exc
        headers = tuple(
            "" if cell.value is None else str(cell.value).strip() for cell in header_cells
        )
        _validate_import_headers(headers)
        rows: list[dict[str, str]] = []
        for row_index, cells in enumerate(iterator, start=2):
            if row_index > MAX_IMPORT_ROWS + 1:
                raise RuleImportError(f"规则导入不能超过 {MAX_IMPORT_ROWS} 行")
            if any(cell.data_type == "f" for cell in cells):
                raise RuleImportError(f"第 {row_index} 行包含公式")
            values = [
                "" if cell.value is None else str(cell.value).strip()
                for cell in cells[: len(headers)]
            ]
            if not any(values):
                raise RuleImportError(f"第 {row_index} 行为空")
            _reject_formula_like_values(values, row_index)
            rows.append(dict(zip(headers, values, strict=True)))
        return rows
    finally:
        workbook.close()


def _validate_xlsx_container(path: Path) -> None:
    forbidden_fragments = (
        "vbaproject",
        "/embeddings/",
        "/externallinks/",
        "/connections",
        "/customxml/",
        "oleobject",
    )
    try:
        with zipfile.ZipFile(path) as package:
            entries = package.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise RuleImportError("XLSX 包含过多文件部件")
            total_size = sum(item.file_size for item in entries)
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise RuleImportError("XLSX 解压后体积过大")
            for item in entries:
                normalized = item.filename.replace("\\", "/").casefold()
                if normalized.startswith("/") or ".." in Path(normalized).parts:
                    raise RuleImportError("XLSX 包含不安全路径")
                if any(fragment in normalized for fragment in forbidden_fragments):
                    raise RuleImportError("XLSX 包含不允许的活动或嵌入内容")
                if normalized.endswith(".rels"):
                    relationship_xml = package.read(item)
                    if re.search(
                        rb"\bTargetMode\s*=\s*['\"]External['\"]",
                        relationship_xml,
                        re.IGNORECASE,
                    ):
                        raise RuleImportError("XLSX 包含外部关系")
    except zipfile.BadZipFile as exc:
        raise RuleImportError("XLSX 文件结构无效") from exc
    except OSError as exc:
        raise RuleImportError("无法检查 XLSX 文件结构") from exc


def _validate_import_headers(headers: Sequence[str]) -> None:
    if not headers or any(not header for header in headers):
        raise RuleImportError("规则导入表头不能为空")
    if len(headers) != len(set(headers)):
        raise RuleImportError("规则导入表头重复")
    missing = [header for header in IMPORT_HEADERS if header not in headers]
    if missing:
        raise RuleImportError(f"规则导入缺少表头：{'、'.join(missing)}")
    allowed = set(IMPORT_HEADERS) | set(OPTIONAL_IMPORT_HEADERS)
    unknown = [header for header in headers if header not in allowed]
    if unknown:
        raise RuleImportError(f"规则导入包含未知表头：{'、'.join(unknown)}")


def _rule_from_import_row(row: Mapping[str, str], row_number: int) -> RuleDefinition:
    required_non_empty = ("规则类型", "规则名称", "匹配方式", "处理方式")
    missing = [header for header in required_non_empty if not row.get(header, "").strip()]
    match_mode_text = row.get("匹配方式", "").strip()
    if (
        match_mode_text not in {"仅人工", MatchMode.MANUAL.value}
        and not row.get("关键词", "").strip()
    ):
        missing.append("关键词")
    if missing:
        raise RuleImportError(f"第 {row_number} 行缺少必填值：{'、'.join(missing)}")

    try:
        kind = _enum_alias(
            row["规则类型"],
            {
                "固定替换": RuleKind.FIXED,
                "判断标准": RuleKind.STANDARD,
            },
            RuleKind,
        )
        match_mode = _enum_alias(
            row["匹配方式"],
            {
                "精确": MatchMode.EXACT,
                "包含": MatchMode.CONTAINS,
                "受控正则": MatchMode.REGEX,
                "任一关键词": MatchMode.ANY_KEYWORD,
                "全部关键词": MatchMode.ALL_KEYWORDS,
                "仅人工": MatchMode.MANUAL,
            },
            MatchMode,
        )
        action = _enum_alias(
            row["处理方式"],
            {
                "固定代号": ActionKind.FIXED_REPLACEMENT,
                "固定替换": ActionKind.FIXED_REPLACEMENT,
                "中间星号": ActionKind.MASK_MIDDLE,
                "任务内顺序代号": ActionKind.SEQUENCE_CODE,
                "删除": ActionKind.DELETE,
            },
            ActionKind,
        )
        patterns = () if match_mode is MatchMode.MANUAL else _split_values(row.get("关键词", ""))
        positive_examples = _split_values(row.get("正例", ""))
        negative_examples = _split_values(row.get("反例", ""))
        scopes = _split_values(row.get("适用范围", "")) or DEFAULT_SCOPES
        rule = RuleDefinition(
            id=row.get("规则编号", "").strip() or uuid4().hex,
            name=row["规则名称"],
            kind=kind,
            match_mode=match_mode,
            patterns=patterns,
            action=action,
            replacement=row.get("替换内容", ""),
            mandatory=_parse_bool(row.get("必须处理", ""), default=True),
            enabled=_parse_bool(row.get("启用", ""), default=True),
            priority=_parse_int(row.get("优先级", ""), default=0),
            positive_examples=positive_examples,
            negative_examples=negative_examples,
            applies_to=scopes,
            case_sensitive=_parse_bool(row.get("区分大小写", ""), default=True),
            keep_prefix=_parse_int(row.get("保留开头", ""), default=1),
            keep_suffix=_parse_int(row.get("保留结尾", ""), default=1),
            code_width=_parse_int(row.get("代号位数", ""), default=2),
            category=row.get("内容类别", ""),
        )
    except (KeyError, RuleValidationError, ValueError) as exc:
        if isinstance(exc, RuleImportError):
            raise
        raise RuleImportError(f"第 {row_number} 行规则无效：{exc}") from exc
    return rule


def _reject_formula_like_values(values: Iterable[str], row_number: int) -> None:
    for value in values:
        stripped = value.lstrip()
        if stripped.startswith(("=", "+", "-", "@")):
            raise RuleImportError(f"第 {row_number} 行包含公式或公式样式内容")


def _split_values(value: str) -> tuple[str, ...]:
    return tuple(item.strip() for item in re.split(r"[|｜]", value) if item.strip())


def _parse_bool(value: str, *, default: bool) -> bool:
    normalized = value.strip().casefold()
    if not normalized:
        return default
    if normalized in {"是", "true", "1", "yes", "启用"}:
        return True
    if normalized in {"否", "false", "0", "no", "停用"}:
        return False
    raise ValueError("布尔值应填写是或否")


def _parse_int(value: str, *, default: int) -> int:
    if not value.strip():
        return default
    try:
        return int(value)
    except ValueError as exc:
        raise ValueError("数字字段格式无效") from exc


def _enum_alias(
    value: str,
    aliases: Mapping[str, Any],
    enum_type: type[StrEnum],
) -> Any:
    normalized = value.strip()
    if normalized in aliases:
        return aliases[normalized]
    return enum_type(normalized)


def _mapping(value: object, label: str) -> Mapping[str, object]:
    if not isinstance(value, dict):
        raise RuleValidationError(f"{label} 必须是对象")
    if not all(isinstance(key, str) for key in value):
        raise RuleValidationError(f"{label} 的字段名必须是字符串")
    return value


def _string_tuple(value: object, label: str) -> tuple[str, ...]:
    if not isinstance(value, (list, tuple)) or not all(isinstance(item, str) for item in value):
        raise RuleValidationError(f"{label} 必须是字符串数组")
    return tuple(value)


def _strict_bool(value: object, label: str) -> bool:
    if not isinstance(value, bool):
        raise RuleValidationError(f"{label} 必须是布尔值")
    return value


def _strict_int(value: object, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise RuleValidationError(f"{label} 必须是整数")
    return value


def _strict_string(value: object, label: str) -> str:
    if not isinstance(value, str):
        raise RuleValidationError(f"{label} 必须是字符串")
    return value


__all__ = [
    "ActionKind",
    "MatchMode",
    "Protector",
    "RuleApplication",
    "RuleConflictError",
    "RuleDefinition",
    "RuleImportError",
    "RuleImportResult",
    "RuleKind",
    "RuleLibrary",
    "RuleLibraryError",
    "RuleMatch",
    "RuleProtectionError",
    "RuleSession",
    "RuleStore",
    "RuleStoreError",
    "RuleValidationError",
    "WindowsDpapiProtector",
    "default_industry_rules",
    "default_standard_rules",
    "default_rule_store_path",
    "import_rules",
    "is_preset_rule",
    "merge_imported_rules",
    "restore_default_rules",
]
