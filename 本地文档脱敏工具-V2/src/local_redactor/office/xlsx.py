from __future__ import annotations

import copy
import hashlib
import io
import re
import uuid
import zipfile
from collections.abc import Sequence
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from lxml import etree  # type: ignore[import-untyped]
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.drawing.image import Image as OpenpyxlImage  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import (  # type: ignore[import-untyped]
    column_index_from_string,
    get_column_letter,
)

from ..models import (
    DocumentKind,
    DocumentModel,
    Finding,
    ImageObject,
    ObjectDisposition,
    PackageItem,
    SourceLocation,
    TextBlock,
)
from .common import (
    ANALYSIS_COPY_NOTICE,
    ANALYSIS_COPY_TITLE,
    cleanup_temp,
    ensure_export_ready,
    finalize_export,
    forbidden_originals,
    image_replacement_legend,
    item_disposition,
    prepare_image_for_export,
    sanitize_image,
    secure_target_path,
    temporary_export_path,
    transformed_blocks,
    validate_source_image,
    verify_rebuilt_ooxml,
)
from .security import (
    ExportSafetyError,
    OoxmlPreflight,
    OoxmlSecurityError,
    preflight_ooxml,
    safe_xml_root,
    sha256_file,
)

S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
XDR_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS = {"s": S_NS, "r": R_NS, "a": A_NS, "xdr": XDR_NS}
_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_MAX_CELLS = 2_000_000


def _q(namespace: str, local_name: str) -> str:
    return f"{{{namespace}}}{local_name}"


def _new_block(
    model: DocumentModel,
    *,
    text: str,
    part: str,
    display: str,
    block_kind: str,
    hidden: bool = False,
    sheet: str | None = None,
    cell: str | None = None,
    style: dict[str, Any] | None = None,
) -> TextBlock:
    block_id = uuid.uuid4().hex
    block = TextBlock(
        id=block_id,
        text=text,
        location=SourceLocation(
            part=part,
            display=display,
            block_id=block_id,
            sheet=sheet,
            cell=cell,
        ),
        block_kind=block_kind,
        hidden=hidden,
        style=style or {},
    )
    model.blocks.append(block)
    return block


def _add_virtual_item(
    model: DocumentModel,
    *,
    part: str,
    kind: str,
    hidden: bool = True,
    requires_decision: bool = True,
) -> PackageItem:
    existing = next((item for item in model.inventory.parts if item.part == part), None)
    if existing is not None:
        existing.requires_decision = existing.requires_decision or requires_decision
        existing.hidden = existing.hidden or hidden
        return existing
    item = PackageItem(
        part=part,
        kind=kind,
        hidden=hidden,
        requires_decision=requires_decision,
        disposition=(ObjectDisposition.PENDING if requires_decision else ObjectDisposition.REMOVE),
    )
    model.inventory.parts.append(item)
    return item


def _workbook_sheet_parts(
    archive: zipfile.ZipFile,
    preflight: OoxmlPreflight,
) -> dict[str, str]:
    root = safe_xml_root(archive.read("xl/workbook.xml"), part_name="xl/workbook.xml")
    relationships = {
        record.relationship_id: record.resolved_target
        for record in preflight.relationships
        if record.source_part == "xl/workbook.xml"
        and record.resolved_target is not None
        and record.relationship_type.casefold().endswith("/worksheet")
    }
    result: dict[str, str] = {}
    sheets = root.find(_q(S_NS, "sheets"))
    if sheets is None:
        raise OoxmlSecurityError("missing-sheets", "Excel 工作簿缺少工作表清单")
    for sheet in sheets:
        name = sheet.get("name", "")
        relationship_id = sheet.get(_q(R_NS, "id"), "")
        target = relationships.get(relationship_id)
        if not name or target is None:
            raise OoxmlSecurityError(
                "invalid-sheet-relationship",
                "Excel 工作表名称或关系无法解析",
            )
        result[name] = target
    return result


def _property_blocks(
    model: DocumentModel,
    archive: zipfile.ZipFile,
    preflight: OoxmlPreflight,
) -> None:
    for part in sorted(
        name
        for name in preflight.part_names
        if name.startswith("docProps/") and name.endswith(".xml")
    ):
        root = safe_xml_root(archive.read(part), part_name=part)
        for element in root.iter():
            if len(element) or not element.text or not element.text.strip():
                continue
            key = etree.QName(element).localname
            value = element.text.strip()
            unique_key = key
            counter = 2
            while unique_key in model.document_properties:
                unique_key = f"{key}_{counter}"
                counter += 1
            model.document_properties[unique_key] = value
            _new_block(
                model,
                text=value,
                part=part,
                display=f"文件属性：{key}",
                block_kind="metadata",
                hidden=True,
            )


def _safe_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _cell_style(cell: Any) -> dict[str, Any]:
    return {
        "font": copy.copy(cell.font),
        "fill": copy.copy(cell.fill),
        "border": copy.copy(cell.border),
        "alignment": copy.copy(cell.alignment),
        "protection": copy.copy(cell.protection),
        "number_format": cell.number_format,
    }


def _scan_cells(
    model: DocumentModel,
    workbook: Any,
    values_workbook: Any,
    sheet_parts: dict[str, str],
) -> tuple[list[dict[str, Any]], dict[tuple[str, str], str]]:
    sheets_structure: list[dict[str, Any]] = []
    cell_blocks: dict[tuple[str, str], str] = {}
    total_cells = 0
    for worksheet in workbook.worksheets:
        sheet_part = sheet_parts.get(worksheet.title)
        if sheet_part is None:
            raise OoxmlSecurityError(
                "worksheet-model-mismatch",
                "Excel 高层模型与 OOXML 工作表关系不一致",
            )
        values_sheet = values_workbook[worksheet.title]
        name_block = _new_block(
            model,
            text=worksheet.title,
            part="xl/workbook.xml",
            display=f"工作表名：{worksheet.title}",
            block_kind="worksheet_name",
            hidden=worksheet.sheet_state != "visible",
            sheet=worksheet.title,
        )
        sheet_item_part = f"{sheet_part}#hidden-sheet"
        if worksheet.sheet_state != "visible":
            _add_virtual_item(
                model,
                part=sheet_item_part,
                kind="hidden_sheet",
            )
        cells: list[dict[str, Any]] = []
        source_cells = sorted(
            worksheet._cells.values(),
            key=lambda cell: (cell.row, cell.column),
        )
        total_cells += len(source_cells)
        if total_cells > _MAX_CELLS:
            raise OoxmlSecurityError(
                "too-many-cells",
                "Excel 实际单元格数量超过安全上限",
            )
        for cell in source_cells:
            if cell.__class__.__name__ == "MergedCell":
                continue
            value = cell.value
            cached_value = values_sheet[cell.coordinate].value
            is_formula = cell.data_type == "f" or (isinstance(value, str) and value.startswith("="))
            output_value = cached_value if is_formula else value
            block = _new_block(
                model,
                text=_safe_cell_text(output_value),
                part=sheet_part,
                display=f"{worksheet.title}!{cell.coordinate}",
                block_kind="cell",
                sheet=worksheet.title,
                cell=cell.coordinate,
                style={"data_type": cell.data_type},
            )
            cell_blocks[(worksheet.title, cell.coordinate)] = block.id
            formula_block_id = None
            formula_item_part = None
            if is_formula:
                formula_item_part = f"{sheet_part}#formula:{cell.coordinate}"
                _add_virtual_item(
                    model,
                    part=formula_item_part,
                    kind="formula",
                    requires_decision=cached_value is None,
                )
                formula_block = _new_block(
                    model,
                    text=_safe_cell_text(value),
                    part=formula_item_part,
                    display=f"{worksheet.title}!{cell.coordinate}（公式）",
                    block_kind="formula",
                    hidden=True,
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                )
                formula_block_id = formula_block.id

            comment_part = None
            comment_block_id = None
            if cell.comment is not None:
                comment_part = f"{sheet_part}#comment:{cell.coordinate}"
                _add_virtual_item(model, part=comment_part, kind="comment")
                comment_block = _new_block(
                    model,
                    text=cell.comment.text or "",
                    part=comment_part,
                    display=f"{worksheet.title}!{cell.coordinate}（批注）",
                    block_kind="comment",
                    hidden=True,
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                )
                comment_block_id = comment_block.id

            hyperlink_part = None
            hyperlink_block_id = None
            hyperlink = cell.hyperlink
            if hyperlink is not None:
                hyperlink_part = f"{sheet_part}#hyperlink:{cell.coordinate}"
                _add_virtual_item(model, part=hyperlink_part, kind="hyperlink")
                target = hyperlink.target or hyperlink.location or ""
                hyperlink_block = _new_block(
                    model,
                    text=target,
                    part=hyperlink_part,
                    display=f"{worksheet.title}!{cell.coordinate}（链接目标）",
                    block_kind="hyperlink_target",
                    hidden=True,
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                )
                hyperlink_block_id = hyperlink_block.id

            cells.append(
                {
                    "coordinate": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "block_id": block.id,
                    "source_value": value,
                    "cached_value": cached_value,
                    "is_formula": is_formula,
                    "formula_part": formula_item_part,
                    "formula_block_id": formula_block_id,
                    "comment_part": comment_part,
                    "comment_block_id": comment_block_id,
                    "hyperlink_part": hyperlink_part,
                    "hyperlink_block_id": hyperlink_block_id,
                    "style": _cell_style(cell),
                }
            )

        hidden_rows: list[dict[str, Any]] = []
        for row_index, dimension in sorted(worksheet.row_dimensions.items()):
            if dimension.hidden:
                part = f"{sheet_part}#hidden-row:{row_index}"
                _add_virtual_item(model, part=part, kind="hidden_row")
                hidden_rows.append({"index": int(row_index), "part": part})
        hidden_columns: list[dict[str, Any]] = []
        for key, dimension in sorted(worksheet.column_dimensions.items()):
            if not dimension.hidden:
                continue
            try:
                minimum = (
                    int(dimension.min)
                    if dimension.min is not None
                    else column_index_from_string(key)
                )
                maximum = (
                    int(dimension.max)
                    if dimension.max is not None
                    else column_index_from_string(key)
                )
            except (TypeError, ValueError):
                minimum = column_index_from_string(key.split(":", maxsplit=1)[0])
                maximum = column_index_from_string(key.split(":", maxsplit=1)[-1])
            part = f"{sheet_part}#hidden-column:{minimum}:{maximum}"
            _add_virtual_item(model, part=part, kind="hidden_column")
            hidden_columns.append({"minimum": minimum, "maximum": maximum, "part": part})

        headers_footers = _scan_headers_footers(
            model,
            worksheet,
            sheet_part,
        )
        row_dimensions = {
            int(index): {"height": dimension.height}
            for index, dimension in worksheet.row_dimensions.items()
            if dimension.height is not None
        }
        column_dimensions = {
            key: {"width": dimension.width}
            for key, dimension in worksheet.column_dimensions.items()
            if dimension.width is not None
        }
        sheets_structure.append(
            {
                "source_name": worksheet.title,
                "name_block_id": name_block.id,
                "part": sheet_part,
                "state": worksheet.sheet_state,
                "hidden_sheet_part": (
                    sheet_item_part if worksheet.sheet_state != "visible" else None
                ),
                "cells": cells,
                "merged_ranges": [str(cell_range) for cell_range in worksheet.merged_cells.ranges],
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_columns,
                "row_dimensions": row_dimensions,
                "column_dimensions": column_dimensions,
                "headers_footers": headers_footers,
                "freeze_panes": (str(worksheet.freeze_panes) if worksheet.freeze_panes else None),
            }
        )
    return sheets_structure, cell_blocks


def _scan_headers_footers(
    model: DocumentModel,
    worksheet: Any,
    sheet_part: str,
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for group_name in (
        "oddHeader",
        "oddFooter",
        "evenHeader",
        "evenFooter",
        "firstHeader",
        "firstFooter",
    ):
        group = getattr(worksheet, group_name)
        for position in ("left", "center", "right"):
            item = getattr(group, position)
            text = item.text or ""
            if not text:
                continue
            part = f"{sheet_part}#{group_name}:{position}"
            _add_virtual_item(model, part=part, kind="header_footer")
            block = _new_block(
                model,
                text=text,
                part=part,
                display=f"{worksheet.title} {group_name}.{position}",
                block_kind="header_footer",
                hidden=True,
                sheet=worksheet.title,
            )
            result.append(
                {
                    "part": part,
                    "group": group_name,
                    "position": position,
                    "block_id": block.id,
                }
            )
    return result


def _scan_defined_names(model: DocumentModel, workbook: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    values = (
        workbook.defined_names.values()
        if hasattr(workbook.defined_names, "values")
        else workbook.defined_names.definedName
    )
    for index, defined_name in enumerate(values, start=1):
        name = str(getattr(defined_name, "name", ""))
        value = str(getattr(defined_name, "attr_text", "") or "")
        part = f"xl/workbook.xml#defined-name:{index}"
        _add_virtual_item(model, part=part, kind="defined_name")
        block = _new_block(
            model,
            text=f"{name} {value}".strip(),
            part=part,
            display=f"定义名称：{name}",
            block_kind="defined_name",
            hidden=True,
        )
        result.append({"part": part, "block_id": block.id})
    return result


def _scan_images(
    model: DocumentModel,
    archive: zipfile.ZipFile,
    preflight: OoxmlPreflight,
    sheet_parts: dict[str, str],
    cell_blocks: dict[tuple[str, str], str],
) -> dict[str, str]:
    sheet_name_by_part = {part: name for name, part in sheet_parts.items()}
    drawing_to_sheet = {
        record.resolved_target: sheet_name_by_part[record.source_part]
        for record in preflight.relationships
        if record.source_part in sheet_name_by_part
        and record.resolved_target is not None
        and record.relationship_type.casefold().endswith("/drawing")
    }
    image_relations = {
        (record.source_part, record.relationship_id): record
        for record in preflight.relationships
        if record.source_part in drawing_to_sheet
        and record.resolved_target is not None
        and record.relationship_type.casefold().endswith("/image")
    }
    represented_parts: set[str] = set()
    alt_blocks: dict[str, str] = {}
    for drawing_part, sheet_name in sorted(drawing_to_sheet.items()):
        if drawing_part not in preflight.part_names:
            continue
        root = safe_xml_root(archive.read(drawing_part), part_name=drawing_part)
        for blip in root.iter(_q(A_NS, "blip")):
            relationship_id = blip.get(_q(R_NS, "embed"), "")
            relationship = image_relations.get((drawing_part, relationship_id))
            if relationship is None or relationship.resolved_target is None:
                continue
            media_part = relationship.resolved_target
            represented_parts.add(media_part)
            anchor = next(
                (
                    ancestor
                    for ancestor in blip.iterancestors()
                    if etree.QName(ancestor).localname
                    in {"oneCellAnchor", "twoCellAnchor", "absoluteAnchor"}
                ),
                None,
            )
            coordinate = "A1"
            if anchor is not None:
                from_node = anchor.find(_q(XDR_NS, "from"))
                if from_node is not None:
                    row_node = from_node.find(_q(XDR_NS, "row"))
                    column_node = from_node.find(_q(XDR_NS, "col"))
                    if row_node is not None and column_node is not None:
                        try:
                            row = int(row_node.text or "0") + 1
                            column = int(column_node.text or "0") + 1
                            coordinate = f"{get_column_letter(column)}{row}"
                        except ValueError:
                            coordinate = "A1"
            properties = anchor.find(f".//{_q(XDR_NS, 'cNvPr')}") if anchor is not None else None
            alt_text = ""
            if properties is not None:
                alt_text = " ".join(
                    value
                    for value in (
                        properties.get("name", ""),
                        properties.get("title", ""),
                        properties.get("descr", ""),
                    )
                    if value
                )
            image_id = uuid.uuid4().hex
            content = archive.read(media_part)
            width, height = validate_source_image(
                content,
                source_part=media_part,
            )
            model.images.append(
                ImageObject(
                    id=image_id,
                    source_part=media_part,
                    location=SourceLocation(
                        part=drawing_part,
                        display=f"{sheet_name}!{coordinate} 图片",
                        block_id=cell_blocks.get((sheet_name, coordinate)),
                        sheet=sheet_name,
                        cell=coordinate,
                        image_id=image_id,
                    ),
                    content=content,
                    media_type=preflight.content_types.get(
                        media_part,
                        "application/octet-stream",
                    ),
                    width=width,
                    height=height,
                    alt_text=alt_text,
                )
            )
            if alt_text:
                block = _new_block(
                    model,
                    text=alt_text,
                    part=f"{drawing_part}#image-alt:{image_id}",
                    display=f"{sheet_name}!{coordinate} 图片替代文字",
                    block_kind="alt_text",
                    hidden=True,
                    sheet=sheet_name,
                    cell=coordinate,
                )
                block.location.image_id = image_id
                alt_blocks[image_id] = block.id

    for media_part in sorted(
        name
        for name in preflight.part_names
        if preflight.content_types.get(name, "").startswith("image/")
        and "/media/" in name
        and name not in represented_parts
    ):
        image_id = uuid.uuid4().hex
        content = archive.read(media_part)
        width, height = validate_source_image(
            content,
            source_part=media_part,
        )
        model.images.append(
            ImageObject(
                id=image_id,
                source_part=media_part,
                location=SourceLocation(
                    part=media_part,
                    display=f"未引用图片：{media_part}",
                    image_id=image_id,
                ),
                content=content,
                media_type=preflight.content_types.get(
                    media_part,
                    "application/octet-stream",
                ),
                width=width,
                height=height,
            )
        )
        model.warnings.append("发现未被工作表绘图关系引用的图片，仍需逐图复核")
    return alt_blocks


def _scan_external_relationships(
    model: DocumentModel,
    preflight: OoxmlPreflight,
) -> None:
    for relationship in preflight.relationships:
        if not relationship.external:
            continue
        _new_block(
            model,
            text=relationship.target,
            part=(f"{relationship.relationship_part}#{relationship.relationship_id}"),
            display="外部关系真实目标",
            block_kind="external_relationship",
            hidden=True,
        )


class XlsxAdapter:
    kind = DocumentKind.XLSX

    def scan(self, path: Path) -> DocumentModel:
        source_path = Path(path)
        before_hash = sha256_file(source_path)
        preflight = preflight_ooxml(source_path, self.kind)
        model = DocumentModel(
            kind=self.kind,
            source_path=source_path,
            source_sha256=preflight.source_sha256,
            display_name=source_path.name,
            inventory=preflight.inventory,
            structure={"format": "xlsx"},
        )
        try:
            workbook = load_workbook(
                source_path,
                data_only=False,
                read_only=False,
                keep_links=False,
            )
            values_workbook = load_workbook(
                source_path,
                data_only=True,
                read_only=False,
                keep_links=False,
            )
            with zipfile.ZipFile(source_path, mode="r") as archive:
                sheet_parts = _workbook_sheet_parts(archive, preflight)
                _property_blocks(model, archive, preflight)
                sheets, cell_blocks = _scan_cells(
                    model,
                    workbook,
                    values_workbook,
                    sheet_parts,
                )
                defined_names = _scan_defined_names(model, workbook)
                alt_blocks = _scan_images(
                    model,
                    archive,
                    preflight,
                    sheet_parts,
                    cell_blocks,
                )
                _scan_external_relationships(model, preflight)
                model.structure.update(
                    {
                        "sheets": sheets,
                        "defined_names": defined_names,
                        "image_alt_blocks": alt_blocks,
                    }
                )
            workbook.close()
            values_workbook.close()
        except OoxmlSecurityError:
            raise
        except (KeyError, ValueError, TypeError, zipfile.BadZipFile) as exc:
            raise OoxmlSecurityError(
                "changed-or-damaged-source",
                "Excel 文件无法建立安全工作簿模型",
            ) from exc
        after_hash = sha256_file(source_path)
        if before_hash != after_hash or after_hash != preflight.source_sha256:
            raise OoxmlSecurityError(
                "source-changed",
                "Excel 原文件在扫描期间发生变化，请重新选择",
            )
        return model

    def export(
        self,
        document: DocumentModel,
        findings: Sequence[Finding],
        target_path: Path,
    ) -> None:
        if document.kind is not self.kind:
            raise ExportSafetyError("文档模型类型与 Excel 适配器不一致")
        if sha256_file(document.source_path) != document.source_sha256:
            raise ExportSafetyError("原文件自扫描后发生变化，不能继续导出")
        ensure_export_ready(document, findings)
        target = secure_target_path(document.source_path, target_path, ".xlsx")
        temporary = temporary_export_path(target)
        texts = transformed_blocks(document, findings)
        original_image_hashes = {
            hashlib.sha256(image.content).hexdigest() for image in document.images
        }
        prepared_images = [prepare_image_for_export(image, findings) for image in document.images]
        try:
            rebuilt = Workbook()
            default_sheet = rebuilt.active
            rebuilt.remove(default_sheet)
            _reset_workbook_properties(rebuilt)
            notice_sheet = rebuilt.create_sheet("脱敏说明")
            _populate_analysis_notice_sheet(notice_sheet)
            used_names: set[str] = {"脱敏说明"}
            notes: list[tuple[str, str]] = []
            source_to_output_sheet: dict[str, Any] = {}
            for sheet_data in document.structure.get("sheets", []):
                hidden_part = sheet_data.get("hidden_sheet_part")
                if (
                    hidden_part
                    and item_disposition(
                        document,
                        str(hidden_part),
                    )
                    is not ObjectDisposition.VISIBLE_NOTE
                ):
                    continue
                proposed_name = texts.get(
                    str(sheet_data["name_block_id"]),
                    str(sheet_data["source_name"]),
                )
                sheet_name = _safe_unique_sheet_name(proposed_name, used_names)
                worksheet = rebuilt.create_sheet(sheet_name)
                worksheet.sheet_state = "visible"
                source_to_output_sheet[str(sheet_data["source_name"])] = worksheet
                removed_rows = {
                    int(item["index"])
                    for item in sheet_data.get("hidden_rows", [])
                    if item_disposition(document, str(item["part"]))
                    is not ObjectDisposition.VISIBLE_NOTE
                }
                removed_columns: set[int] = set()
                for item in sheet_data.get("hidden_columns", []):
                    if (
                        item_disposition(
                            document,
                            str(item["part"]),
                        )
                        is ObjectDisposition.VISIBLE_NOTE
                    ):
                        continue
                    removed_columns.update(range(int(item["minimum"]), int(item["maximum"]) + 1))
                for cell_data in sheet_data.get("cells", []):
                    row = int(cell_data["row"])
                    column = int(cell_data["column"])
                    if row in removed_rows or column in removed_columns:
                        continue
                    target_cell = worksheet.cell(row=row, column=column)
                    target_cell.value = _rebuilt_cell_value(
                        document,
                        cell_data,
                        texts,
                    )
                    _copy_cell_style(target_cell, cell_data.get("style", {}))
                    comment_part = cell_data.get("comment_part")
                    comment_block_id = cell_data.get("comment_block_id")
                    if (
                        comment_part
                        and comment_block_id
                        and item_disposition(document, str(comment_part))
                        is ObjectDisposition.VISIBLE_NOTE
                    ):
                        notes.append(
                            (
                                f"{sheet_name}!{cell_data['coordinate']} 原批注",
                                texts.get(str(comment_block_id), ""),
                            )
                        )
                    hyperlink_part = cell_data.get("hyperlink_part")
                    hyperlink_block_id = cell_data.get("hyperlink_block_id")
                    if (
                        hyperlink_part
                        and hyperlink_block_id
                        and item_disposition(document, str(hyperlink_part))
                        is ObjectDisposition.VISIBLE_NOTE
                    ):
                        notes.append(
                            (
                                f"{sheet_name}!{cell_data['coordinate']} 原链接用途",
                                texts.get(str(hyperlink_block_id), ""),
                            )
                        )
                for merged_range in sheet_data.get("merged_ranges", []):
                    try:
                        worksheet.merge_cells(str(merged_range))
                    except ValueError:
                        notes.append(
                            (
                                f"{sheet_name} 合并区域",
                                "原合并区域因内容移除未能保留",
                            )
                        )
                for row_index, dimension in sheet_data.get("row_dimensions", {}).items():
                    worksheet.row_dimensions[int(row_index)].height = dimension.get("height")
                    worksheet.row_dimensions[int(row_index)].hidden = False
                for key, dimension in sheet_data.get("column_dimensions", {}).items():
                    worksheet.column_dimensions[str(key)].width = dimension.get("width")
                    worksheet.column_dimensions[str(key)].hidden = False
                freeze_panes = sheet_data.get("freeze_panes")
                if freeze_panes:
                    worksheet.freeze_panes = str(freeze_panes)
                for entry in sheet_data.get("headers_footers", []):
                    if (
                        item_disposition(
                            document,
                            str(entry["part"]),
                        )
                        is ObjectDisposition.VISIBLE_NOTE
                    ):
                        notes.append(
                            (
                                f"{sheet_name} 页眉页脚",
                                texts.get(str(entry["block_id"]), ""),
                            )
                        )

            for entry in document.structure.get("defined_names", []):
                if (
                    item_disposition(
                        document,
                        str(entry["part"]),
                    )
                    is ObjectDisposition.VISIBLE_NOTE
                ):
                    notes.append(
                        (
                            "定义名称",
                            texts.get(str(entry["block_id"]), ""),
                        )
                    )

            if not rebuilt.worksheets:
                worksheet = rebuilt.create_sheet("已移除内容")
                worksheet["A1"] = "原工作表内容已按复核决定移除。"

            image_streams: list[io.BytesIO] = []
            for image in prepared_images:
                sanitized = sanitize_image(image)
                if sanitized is None:
                    continue
                content, _media_type = sanitized
                source_sheet = image.location.sheet
                worksheet = (
                    source_to_output_sheet.get(source_sheet)
                    if source_sheet is not None
                    else next(iter(source_to_output_sheet.values()), notice_sheet)
                )
                if worksheet is None:
                    continue
                stream = io.BytesIO(content)
                image_streams.append(stream)
                exported_image = OpenpyxlImage(stream)
                exported_image.anchor = image.location.cell or "A1"
                worksheet.add_image(exported_image)
                legend = image_replacement_legend(image)
                if legend:
                    location = f"{worksheet.title}!{image.location.cell or 'A1'}"
                    notes.append((f"{location} 图片替代图例", legend))

            if notes:
                notes_name = _safe_unique_sheet_name("可见附注", used_names)
                notes_sheet = rebuilt.create_sheet(notes_name)
                notes_sheet.append(["来源", "经复核内容"])
                for label, value in notes:
                    notes_sheet.append([_safe_formula_text(label), _safe_formula_text(value)])
                notes_sheet.sheet_state = "visible"
                notes_sheet.column_dimensions["A"].width = 36
                notes_sheet.column_dimensions["B"].width = 80

            for worksheet in rebuilt.worksheets:
                worksheet.sheet_state = "visible"
                for dimension in worksheet.row_dimensions.values():
                    dimension.hidden = False
                for dimension in worksheet.column_dimensions.values():
                    dimension.hidden = False
            rebuilt.save(temporary)
            rebuilt.close()
            verify_rebuilt_ooxml(
                temporary,
                self.kind,
                forbidden_values=forbidden_originals(findings),
                original_image_hashes=original_image_hashes,
            )
            reopened = load_workbook(
                temporary,
                data_only=False,
                read_only=False,
                keep_links=False,
            )
            if any(sheet.sheet_state != "visible" for sheet in reopened.worksheets):
                reopened.close()
                raise ExportSafetyError("重建 Excel 仍包含隐藏工作表")
            reopened.close()
            finalize_export(temporary, target)
        finally:
            cleanup_temp(temporary)


def _reset_workbook_properties(workbook: Any) -> None:
    properties = workbook.properties
    properties.creator = ""
    properties.lastModifiedBy = ""
    properties.title = ""
    properties.subject = ""
    properties.description = ""
    properties.keywords = ""
    properties.category = ""
    properties.contentStatus = ""
    properties.identifier = ""
    properties.language = ""
    properties.version = ""
    fixed_time = datetime(2000, 1, 1)
    properties.created = fixed_time
    properties.modified = fixed_time
    properties.lastPrinted = fixed_time
    properties.revision = "1"
    workbook.security.lockStructure = False
    workbook.security.lockWindows = False


def _populate_analysis_notice_sheet(worksheet: Any) -> None:
    worksheet.merge_cells("A1:H1")
    title = worksheet["A1"]
    title.value = ANALYSIS_COPY_TITLE
    title.font = Font(bold=True, color="FFFFFF", size=16)
    title.fill = PatternFill("solid", fgColor="176B62")
    title.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells("A3:H5")
    notice = worksheet["A3"]
    notice.value = ANALYSIS_COPY_NOTICE
    notice.font = Font(bold=True, color="8A3B0A", size=12)
    notice.fill = PatternFill("solid", fgColor="FFF7ED")
    notice.alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )
    for column in range(1, 9):
        worksheet.column_dimensions[get_column_letter(column)].width = 14
    worksheet.sheet_view.showGridLines = False


def _safe_unique_sheet_name(proposed: str, used_names: set[str]) -> str:
    clean = _INVALID_SHEET_CHARS.sub("_", proposed).strip().strip("'")
    if not clean:
        clean = "工作表"
    clean = clean[:31]
    base = clean
    suffix = 2
    while clean.casefold() in used_names:
        ending = f"_{suffix}"
        clean = f"{base[: 31 - len(ending)]}{ending}"
        suffix += 1
    used_names.add(clean.casefold())
    return clean


def _rebuilt_cell_value(
    document: DocumentModel,
    cell_data: dict[str, Any],
    texts: dict[str, str],
) -> Any:
    if cell_data.get("is_formula"):
        cached_value = cell_data.get("cached_value")
        if cached_value is None:
            formula_part = str(cell_data.get("formula_part", ""))
            disposition = item_disposition(document, formula_part)
            if disposition is ObjectDisposition.VISIBLE_NOTE:
                return "[公式无可用缓存值，已移除公式]"
            return None
        if isinstance(cached_value, str):
            return _safe_formula_text(cached_value)
        return cached_value
    source_value = cell_data.get("source_value")
    transformed = texts.get(str(cell_data["block_id"]), _safe_cell_text(source_value))
    if isinstance(source_value, str):
        return _safe_formula_text(transformed)
    if transformed == _safe_cell_text(source_value):
        return source_value
    if isinstance(source_value, bool):
        return transformed.casefold() in {"true", "1", "是"}
    if isinstance(source_value, int):
        try:
            return int(transformed)
        except ValueError:
            return _safe_formula_text(transformed)
    if isinstance(source_value, (float, Decimal)):
        try:
            return float(transformed)
        except ValueError:
            return _safe_formula_text(transformed)
    if isinstance(source_value, (datetime, date, time)):
        return _safe_formula_text(transformed)
    return _safe_formula_text(transformed)


def _safe_formula_text(value: str) -> str:
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _copy_cell_style(cell: Any, style: dict[str, Any]) -> None:
    for attribute in ("font", "fill", "border", "alignment", "protection"):
        if attribute in style:
            setattr(cell, attribute, copy.copy(style[attribute]))
    if "number_format" in style:
        cell.number_format = str(style["number_format"])
